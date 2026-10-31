"""Current-export property/element/device audit. Identifiers remain distinct."""

__release__ = '134'
from OpenSees_Timing_V130 import timed_outputs
from pathlib import Path
import json
import math

OVERRIDES='OpenSees_Element_Overrides_V125.json'
STEEL_REF='https://opensees.github.io/OpenSeesDocumentation/user/manual/material/uniaxialMaterials/Steel02.html'
FVD_REF='https://opensees.berkeley.edu/wiki/index.php/ViscousDamper_Material'
CONCRETE_REF='https://opensees.github.io/OpenSeesDocumentation/user/manual/material/uniaxialMaterials/Concrete01.html'
FVD_EXAMPLE='https://opensees.ist.berkeley.edu/wiki/index.php?title=Dynamic_Analyses_of_1-Story_Moment_Frame_with_Viscous_Dampers'


def read(folder,name,default=None):
    path=Path(folder)/name
    return json.loads(path.read_text(encoding='utf-8-sig')) if path.is_file() else ({} if default is None else default)


def parameter_sources():
    """Source of the formulation does not establish a project's parameters."""
    rows=[]
    def add(family,parameter,value,unit,status,source):
        rows.append(dict(family=family,parameter=parameter,value=value,unit=unit,status=status,source=source))
    import Define_Section_V109 as definitions
    import ast
    source=Path(definitions._component_source)
    configs=next(ast.literal_eval(n.value) for n in ast.parse(source.read_text(encoding='utf-8-sig')).body if isinstance(n,ast.Assign) and any(isinstance(t,ast.Name) and t.id=='RC_SECTION_CONFIG' for t in n.targets))
    for section,config in configs.items():
        for key,value in config.items():
            unit='MPa' if key.endswith('_mpa') else 'mm' if key.endswith('_mm') else 'kN/m3' if 'unit_weight' in key else 'strain' if 'eps' in key.lower() else 'dimensionless / tag / count'
            add(section,key,value,unit,'Current editable template value; no project calibration evidence supplied','No project parameter source supplied; formulation references only: '+CONCRETE_REF+' ; '+STEEL_REF)
    for key,value in [('R0',20.),('cR1',.925),('cR2',.15)]:
        add('Steel02',key,value,'dimensionless','OpenSees recommended transition value; not a device calibration',STEEL_REF)
    for key,unit in [('K','kN/m'),('Cd','kN*(s/m)^alpha'),('alpha','dimensionless')]:
        add('FVD',key,None,unit,'Required manufacturer/test/calibration input',FVD_REF)
    for key,unit in [('core_area','m²'),('Fy','kN/m²'),('E0','kN/m²'),('b','ratio'),('a1','ratio'),('a2','ratio'),('a3','ratio'),('a4','ratio')]:
        add('BRB',key,None,unit,'Required project/test input; SHS area is not BRB core area',STEEL_REF)
    for key,value,unit in [('K_original',25.,'kN/mm'),('Cd_original',20.74,'kN*(s/mm)^0.35'),('alpha',.35,'dimensionless'),('K_converted',25000.,'kN/m'),('Cd_converted',20.74*1000**.35,'kN*(s/m)^0.35')]:
        add('Published FVD example',key,value,unit,'Reference example only; not a calibration for this building',FVD_EXAMPLE)
    from Define_Section_V109 import WALL_REFERENCE_PROFILE
    peer='https://peer.berkeley.edu/sites/default/files/webpeer-2015-12-kristijan_kolozvari_kutay_orakcal_john_wallace.pdf'
    def walk(prefix,value):
        if isinstance(value,dict):
            for key,item in value.items():walk(prefix+'.'+str(key),item)
        elif isinstance(value,(int,float,str)):
            unit='ksi' if prefix.endswith('_ksi') else 'strain' if 'eps' in prefix or prefix.endswith('.et') else 'see source parameter definition'
            add('PEER wall benchmark',prefix,value,unit,'Published specimen input; source-checked card; NOT calibrated to this building',peer+' | PEER 2015/12 Appendix B.3')
    walk('RW-A15-P10-S78',WALL_REFERENCE_PROFILE)
    for family,parameters,ref,status in [
        ('ConcreteCM','fpcc,epcc,Ec,rc,xcrn,ft,et,rt,xcrp,GapClose','https://openseespydoc.readthedocs.io/en/latest/src/ConcreteCM.html','Constitutive formulation; project backbone/calibration required'),
        ('SteelMPF','fyp,fyn,E0,bp,bn,R0,cR1,cR2,a1,a2,a3,a4','https://openseespydoc.readthedocs.io/en/stable/src/SteelMPF.html','Rebar cyclic formulation; tensile/compressive coupon/hardening calibration required'),
        ('FSAM','rho,sX,sY,conc,rhoX,rhoY,nu,alfadow','https://peer.berkeley.edu/publications/2015-12','RC panel model; reinforcement, concrete and dowel/friction calibration required'),
        ('Steel01 exploratory SHS','Fy,E0,b','https://opensees.github.io/OpenSeesDocumentation/user/manual/material/uniaxialMaterials/Steel01.html','Illustrative corotTruss brace route; does not model buckling/fatigue/fracture'),
        ('ElasticMembranePlateSection','E,nu,h,rho','https://opensees.berkeley.edu/wiki/index.php/Elastic_Membrane_Plate_Section','Elastic slab/wall formulation; not nonlinear concrete damage'),
        ('ShellMITC4','nodes,section','https://openseespydoc.readthedocs.io/en/latest/src/ShellMITC4.html','Shell interpolation formulation; mesh convergence still required'),
        ('Elastic torsion/shear','G,J,A','https://openseespydoc.readthedocs.io/en/latest/src/elasticBeamColumn.html','Source elastic stiffness; fiber axial/flexural response does not establish shear failure capacity')]:
        add(family,'Required parameters',parameters,'per command definition',status,ref)
    return rows


def library():
    return [
        dict(name='KEEP_SOURCE_ELASTIC',kind='ANY',description='Retain exported source geometry and elastic properties',source='Current ETABS audit'),
        dict(name='RC_BEAM_600x800_FIBER',kind='BEAM',description='Existing RC template; reinforcement/material calibration required',source=CONCRETE_REF),
        dict(name='RC_COLUMN_1100x1100_FIBER',kind='COLUMN',description='Existing RC template, 1100 mm; review source material',source=CONCRETE_REF),
        dict(name='RC_COLUMN_1250x1250_FIBER',kind='COLUMN',description='Existing RC template, 1250 mm; review source material',source=CONCRETE_REF),
        dict(name='SHS_350x350x16_STEEL_FIBER',kind='BRACE',description='SHS brace section; source area/E audited separately; section may be reused for other frame roles',source=STEEL_REF),
        dict(name='RC_WALL_600_E_SFI_MVLEM3D',kind='SHEAR WALL',description='E-SFI-MVLEM-3D RC wall; 600 mm template; FSAM calibration required',source='https://kkolozvari.github.io/E-SFI-MVLEM-3D/'),
        dict(name='FVD_MAXWELL',kind='BRACE',description='ViscousDamper + twoNodeLink; K, Cd, alpha required',source=FVD_REF),
        dict(name='BRB_STEEL02',kind='BRACE',description='Steel02 + corotTruss; core area and calibrated parameters required',source=STEEL_REF)]


def source_parameter_inventory(folder,payload=None):
    payload=read(folder,'OpenSees_Complete_Model_Data_V118.json') if payload is None else payload;rows=[];seen=set()
    for e in payload.get('frames',[])+payload.get('shells',[]):
        prop=e.get('properties',{});family=prop.get('section',prop.get('property','UNKNOWN'))
        for key,value in prop.items():
            ident=(family,key,str(value))
            if ident in seen:continue
            seen.add(ident);unit={'E':'kN/m2','G':'kN/m2','weight':'kN/m3','thickness':'m','A':'m2','Iy':'m4','Iz':'m4','J':'m4','nu':'dimensionless'}.get(key,'source metadata')
            rows.append(dict(family=family,parameter=key,value=value,unit=unit,status='CURRENT EXPORTED SOURCE VALUE; source identity must be checked; no nonlinear calibration inferred',source=str(Path(folder)/'OpenSees_Complete_Model_Data_V118.json')))
    return rows or [dict(family='Source inventory',status='Rebuild a selected ETABS export to load its actual parameters')]


def validation_inventory():
    return [dict(subsystem='Concrete01 / Steel02 RC templates',evidence='Existing illustrative M35/Fe550 inputs, published command formulations',validation='Project M45/rebar calibration NOT established'),
        dict(subsystem='PEER FSAM wall material card',evidence='PEER 2015/12 Appendix B.3; single-element capability checks in worked examples',validation='Published input provenance; full-building five-fiber wall trial failed gravity; not approved'),
        dict(subsystem='Exploratory building history',evidence='Actual short recorded trials; nonlinear frames and ELASTIC walls',validation='Numerical startup only; full earthquake/design performance NOT established'),
        dict(subsystem='FVD / BRB',evidence='Command formulations and separately labelled published examples',validation='Project device K/Cd/alpha or core area/backbone/test data still required'),
        dict(subsystem='Slab vertical transfer',evidence='ShellMITC4 + shared refined beam/slab nodes; Q=q*A equivalent nodal vectors',validation='V141 regression benchmark and mesh study; project mesh convergence still required')]


def collect(folder):
    folder=Path(folder)
    topology=read(folder,'OpenSees_Element_Connectivity.json')
    sections=read(folder,'OpenSees_Section_Assignment_Audit.json')
    area=read(folder,'OpenSees_Area_Section_Assignment_Audit.json')
    # Actual exported keys are area_assignments / assignments, not sections.
    frames={str(r.get('etabs_object',r.get('etabs_frame_object',''))):r for r in sections.get('frame_assignments',sections.get('sections',[]))}
    areas={str(r.get('etabs_area_object','')):r for r in sections.get('area_assignments',area.get('assignments',[]))}
    payload=read(folder,'OpenSees_Complete_Model_Data_V118.json')
    report=read(folder,'OpenSees_Complete_Export_Report_V118.json')
    if report.get('assembly_errors'):payload={}
    element_map={};solver_rows=[]
    for tag,e in enumerate(payload.get('frames',[]),100001):
        key='FRAME:'+str(e['object']);element_map.setdefault(key,[]).append(tag)
        solver_rows.append(dict(key=key,solver_tag=tag,node_tags=e['nodes'],category=e['category'],properties=e['properties']))
    for tag,e in enumerate(payload.get('shells',[]),100001+len(payload.get('frames',[]))):
        key='AREA:'+str(e['object']);element_map.setdefault(key,[]).append(tag)
        solver_rows.append(dict(key=key,solver_tag=tag,node_tags=e['nodes'],category=e['category'],properties=e['properties']))
    selections=read(folder,OVERRIDES,dict(active=False,assignments={}))
    nodes=[dict(**n,origin='ETABS source point') for n in topology.get('node_map',[])]
    known={str(n['opensees_node_tag']) for n in nodes}
    for tag,xyz in payload.get('nodes',{}).items():
        if str(tag) not in known:nodes.append(dict(etabs_point_name=None,opensees_node_tag=int(tag),etabs_storey_number=None,coordinates_m=xyz,origin='Generated mesh node'))
    for g in payload.get('diaphragms',[]):
        nodes.append(dict(etabs_point_name=None,opensees_node_tag=g['master_tag'],etabs_storey_number=g['storey_number'],coordinates_m=g['master_xyz'],origin='Generated diaphragm master'))
    rows=[];properties={}
    for e in topology.get('elements',[]):
        is_frame=e.get('dimension')=='line_2_node'
        if e.get('dimension') not in ('line_2_node','area_polygon'):continue
        obj=str(e['etabs_object']);key=('FRAME:' if is_frame else 'AREA:')+obj
        a=(frames if is_frame else areas).get(obj,{})
        section=e.get('assigned_section') if is_frame else e.get('assigned_property')
        section=section or a.get('assigned_section' if is_frame else 'assigned_property') or 'UNASSIGNED'
        category=e.get('category','UNKNOWN');group=category+' / '+section
        target=selections.get('assignments',{}).get(key,{}).get('target','KEEP_SOURCE_ELASTIC')
        record=dict(key=key,property=section,category=category,dimension=e['dimension'],etabs_object=obj,
                    topology_tag=e.get('element_tag'),solver_tags=element_map.get(key,[]),node_tags=e.get('node_tags',[]),
                    etabs_points=e.get('node_names',[]),storey=e.get('storey_number'),target=target)
        rows.append(record)
        if group not in properties:
            physical=next((r['properties'] for r in solver_rows if r['key']==key),{})
            properties[group]=dict(group=group,property=section,category=category,material=a.get('material_name',physical.get('material')),
                thickness_m=a.get('thickness_m',physical.get('thickness')),width_m=a.get('width_m'),depth_m=a.get('depth_m'),
                E_kN_m2=physical.get('E',a.get('E_kN_m2')),area_m2=physical.get('A',a.get('A_m2')),count=0,
                source='OpenSees_Section_Assignment_Audit.json / topology / assembled payload')
        properties[group]['count']+=1
    return dict(folder=str(folder),source_edb=topology.get('source_edb'),properties=list(properties.values()),elements=rows,
                solver_elements=solver_rows,nodes=nodes,selections=selections,library=library(),parameters=parameter_sources(),
                load_audit=read(folder,'OpenSees_Load_Conversion_Audit.json'),assembly=read(folder,'OpenSees_Complete_Export_Report_V118.json'))


def save_assignment(folder, keys, target, parameters=None, citation='', active=False):
    data=collect(folder);bykey={r['key']:r for r in data['elements']}
    definitions={r['name']:r for r in library()}
    if target not in definitions:raise ValueError('Unknown library target')
    assignments=data['selections'].get('assignments',{})
    for key in keys:
        owner=key if not key.startswith('SOLVER:') else next((r['key'] for r in data['solver_elements'] if r['solver_tag']==int(key.split(':')[1])),None)
        if owner not in bykey:raise ValueError('No current element for '+key)
        r=bykey[owner];kind=definitions[target]['kind']
        if kind not in ('ANY','FRAME') and r['category']!=kind:
            raise ValueError(f"{r['category']} cannot map to {target}; device assignments require a BRACE.")
        if target=='KEEP_SOURCE_ELASTIC':assignments.pop(key,None);continue
        if kind=='FRAME' and r['dimension']!='line_2_node':raise ValueError('Frame target requires a frame object')
        assignments[key]=dict(target=target,parameters=parameters or {},citation=citation,source_property=r['property'],
            status='candidate; not calibrated' if target in ('FVD_MAXWELL','BRB_STEEL02') else 'reviewed mapping required')
    if active:
        for item in assignments.values():
            if item.get('target') in ('FVD_MAXWELL','BRB_STEEL02'):validate_device(item)
    result=dict(schema_version=125,active=bool(active),assignments=assignments,
                note='Inactive selections preserve the elastic baseline. Active device scenarios require project calibration, physical-weight review and a nonlinear analysis model.')
    Path(folder,OVERRIDES).write_text(json.dumps(result,indent=2),encoding='utf-8')
    return result


def add_device(ops, tag, e, selection):
    """Called only inside a separately enabled nonlinear scenario."""
    from OpenSees_FVD_BRB_Devices_V115 import define_fvd_material,add_fvd_link,define_brb_core_material,add_brb_corot_truss
    validate_device(selection)
    p=selection.get('parameters',{});target=selection['target']
    if not selection.get('citation') or not p.get('weight_reviewed'):
        raise ValueError('Record the device parameter citation and physical-weight review before execution.')
    mat=900000+tag;ni,nj=e['nodes']
    if target=='FVD_MAXWELL':
        define_fvd_material(ops,material_tag=mat,K=p['K'],Cd=p['Cd'],alpha=p['alpha'])
        return add_fvd_link(ops,element_tag=tag,i_node=ni,j_node=nj,material_tag=mat)
    if target=='BRB_STEEL02':
        keys=('Fy','E0','b','R0','cR1','cR2','a1','a2','a3','a4')
        define_brb_core_material(ops,material_tag=mat,**{k:p[k] for k in keys})
        return add_brb_corot_truss(ops,element_tag=tag,i_node=ni,j_node=nj,material_tag=mat,core_area=p['core_area'])
    raise ValueError('Not an executable device target')


def validate_device(selection):
    p=selection.get('parameters',{});target=selection.get('target')
    if p.get('reference_only'):raise ValueError('The published benchmark is reference-only. Supply project-calibrated parameters before activating a device scenario.')
    if not selection.get('citation') or not p.get('weight_reviewed'):
        raise ValueError('Device activation requires a project parameter citation and physical-weight review.')
    keys=('K','Cd','alpha') if target=='FVD_MAXWELL' else ('core_area','Fy','E0','b','R0','cR1','cR2','a1','a2','a3','a4')
    missing=[k for k in keys if not isinstance(p.get(k),(int,float)) or not math.isfinite(p[k])]
    if missing:raise ValueError('Missing/invalid calibrated device parameters: '+', '.join(missing))
    if target=='FVD_MAXWELL' and (p['K']<=0 or p['Cd']<=0 or not 0<p['alpha']<=1):raise ValueError('FVD requires K>0, Cd>0 and 0<alpha<=1')
    if target=='BRB_STEEL02' and (min(p['core_area'],p['Fy'],p['E0'],p['R0'],p['cR1'],p['cR2'])<=0 or not 0<=p['b']<1):raise ValueError('Invalid BRB core area, strengths, stiffness or hardening parameters')


@timed_outputs('Mapping and load audit workbook','folder')
def export_excel(folder,path=None):
    """Application export, using its existing installed Python Excel backend."""
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    data=collect(folder);path=Path(path or Path(folder)/'OpenSees_Property_Element_Load_Audit_V125.xlsx')
    wb=Workbook();wb.remove(wb.active)
    def sheet(name,records):
        ws=wb.create_sheet(name)
        if not records:ws.append(['Status']);ws.append(['No records in current export']);return
        headers=list(records[0]);ws.append(headers)
        for r in records:ws.append([json.dumps(r.get(k),ensure_ascii=False) if isinstance(r.get(k),(list,dict)) else r.get(k) for k in headers])
        ws.freeze_panes='B2';ws.auto_filter.ref=ws.dimensions
        for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='17365D')
        for column in ws.columns:
            ws.column_dimensions[column[0].column_letter].width=min(55,max(17,max(len(str(c.value or '')) for c in list(column)[:80])+2))
            for c in column:c.alignment=Alignment(vertical='top');c.number_format='0.000000' if isinstance(c.value,float) else 'General'
    sheet('Properties',data['properties']);sheet('Elements and mappings',data['elements'])
    sheet('Solver elements',data['solver_elements']);sheet('Nodes',data['nodes'])
    loads=data['load_audit']
    for title,key in [('Frame loads','raw_frame_api'),('Area loads','raw_area_api')]:
        records=loads.get(key,[])
        sheet(title,records if isinstance(records,list) else [dict(object=k,record=v) for k,v in records.items()])
    sheet('Pattern ledger',[dict(pattern=k,gravity_kN=v) for k,v in data['assembly'].get('gravity_by_pattern_kN',{}).items()])
    sheet('Library parameters',data['parameters'])
    sheet('Device selections',[dict(key=k,**v) for k,v in data['selections'].get('assignments',{}).items()])
    sheet('Scope',[dict(item='Source EDB',value=data['source_edb']),dict(item='Device activation',value=data['selections'].get('active',False)),dict(item='Material assumptions',value='RC template values are not calibrated project values. See Library parameters.'),dict(item='Assembly warnings',value=data['assembly'].get('warnings',[]))])
    wb.save(path)
    Path(folder,'OpenSees_Device_and_Parameter_Library_V125.json').write_text(json.dumps(dict(library=library(),parameter_sources=parameter_sources()),indent=2),encoding='utf-8')
    return path
