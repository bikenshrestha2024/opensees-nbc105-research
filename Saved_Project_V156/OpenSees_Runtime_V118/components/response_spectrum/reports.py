"""Saved MRSM reports; generation requires no new model/eigen solve."""
from pathlib import Path
import json,math
from OpenSees_Timing_V130 import timed_outputs


def saved_cases(folder):
    folder=Path(folder);paths=[folder/'MRSM_Results_V132.json'] if (folder/'MRSM_Results_V132.json').is_file() else sorted(folder.rglob('MRSM_Results_V132.json'))
    return [(path,key,case) for path in paths for key,case in json.loads(path.read_text(encoding='utf-8-sig')).items()]


def figures(cases):
    from matplotlib.figure import Figure
    items=[]
    for state in ('ULS','SLS'):
        rows=[(key,r) for _,key,r in cases if r['limit_state']==state]
        if not rows:continue
        fig=Figure(figsize=(11,7),layout='constrained');axes=fig.subplots(1,2)
        for key,r in rows:
            z=[v['height_m'] for v in r['stories']];label=r['axis']+' | '+('centered' if r.get('accidental_eccentricity_sign',0)==0 else 'e=+5% b' if r['accidental_eccentricity_sign']>0 else 'e=-5% b');style={0:'-',1:'--',-1:':'}[r.get('accidental_eccentricity_sign',0)]
            axes[0].plot([v['scaled_story_shear_kN'] for v in r['stories']],z,ls=style,label=label);axes[1].plot([v['design_drift_ratio'] for v in r['stories']],z,ls=style,label=label)
        axes[0].set(xlabel='Scaled storey shear (kN)',ylabel='Height above base (m)',title=state+' MRSM actions');axes[1].set(xlabel='Design edge drift ratio',ylabel='Height above base (m)',title=state+' drift; force scale is not applied');axes[1].axvline(.006 if state=='SLS' else .025,color='black',ls='--',label='NBC drift limit')
        for ax in axes:ax.grid(alpha=.25);ax.legend(fontsize=8)
        items.append((state+' MRSM: shear and drift',fig))
    comparison=Figure(figsize=(12,10),layout='constrained');grid=comparison.subplots(3,2,sharey='row',sharex='row')
    for col,state in enumerate(('ULS','SLS')):
        rows=[(path,key,r) for path,key,r in cases if r['limit_state']==state]
        for path,key,r in rows:
            z=[v['height_m'] for v in r['stories']];label=r['axis']+' '+key
            grid[0,col].plot([v['scaled_story_shear_kN'] for v in r['stories']],z,label=label)
            grid[1,col].plot([v['design_drift_ratio'] for v in r['stories']],z,label=label)
            deformation=path.parent/('Deformation_MRSM_'+state+'_'+key+'_V132.json')
            if deformation.is_file():
                d=json.loads(deformation.read_text(encoding='utf-8-sig'));levels={};i=0 if r['axis']=='X' else 1
                for tag,xyz in d['nodes'].items():
                    u=d['displacements'].get(str(tag),[])
                    if len(u)>i:levels[float(xyz[2])]=max(levels.get(float(xyz[2]),0.),abs(float(u[i]))*1000)
                heights=sorted(levels);grid[2,col].plot([levels[z] for z in heights],heights,label=label)
        for row,title,unit in [(0,'Storey shear','kN'),(1,'Design edge drift','ratio'),(2,'Nodal deflection envelope','mm')]:
            grid[row,col].set(title=state+' '+title,xlabel=unit,ylabel='Elevation / height (m)');grid[row,col].grid(alpha=.25)
            if grid[row,col].lines:grid[row,col].legend(fontsize=7)
        grid[1,col].axvline(.025 if state=='ULS' else .006,color='black',ls='--')
    comparison.suptitle('ULS left / SLS right | Deflection: saved non-concurrent CQC envelopes; force-only scale excluded')
    items.append(('ULS-SLS comparison: shear drift deflection',comparison))
    return items


@timed_outputs('MRSM charts and Excel report','folder')
def export_reports(folder):
    folder=Path(folder);cases=saved_cases(folder)
    if not cases:raise ValueError('No saved MRSM_Results_V132.json in '+str(folder))
    # Application runtime uses its existing openpyxl dependency, as the gravity and history exporters do.
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.chart import ScatterChart,Series,Reference
    wb=Workbook();summary=wb.active;summary.title='Case summary';summary.append(['MRSM saved analysis results — NBC 105:2025']);summary.append(['E0 = centered; Eplus/Eminus = signed 5% perpendicular breadth eccentric inertia. CQC envelopes are not concurrent physical states.']);summary.append(['Displacement/drift are not multiplied by the force-only section 7.5 base-shear scale.']);summary.append([])
    headers=['State','Axis','Case','Eccentricity sign','Raw V (kN)','ESM reference V (kN)','Force scale','Scaled V (kN)','Max design drift','Drift limit','Modal mass (%)','Analysis status','Source JSON']
    summary.append(headers);storey=wb.create_sheet('Storey response');storey.append(['State','Case','Storey','Height (m)','Scaled shear (kN)','Elastic master drift','Design edge drift','Design master drift','Drift limit','Drift status'])
    modal=wb.create_sheet('Modal response');modal.append(['State','Case','Mode','Period (s)','Frequency (Hz)','Signed modal base shear (kN)','UX mode (%)','UY mode (%)','RZ mode (%)','UX cumulative (%)','UY cumulative (%)'])
    nodes=wb.create_sheet('Nodal 6DOF');nodes.append(['State','Case','OpenSees node','X (m)','Y (m)','Z (m)','Ux (m)','Uy (m)','Uz (m)','Rx (rad)','Ry (rad)','Rz (rad)','Scope'])
    chart_index=0
    for path,key,r in cases:
        limit=.006 if r['limit_state']=='SLS' else .025;maxdrift=max(v['design_drift_ratio'] for v in r['stories']);summary.append([r['limit_state'],r['axis'],key,r.get('accidental_eccentricity_sign',0),r['unscaled_modal_base_shear_kN'],r['reference_ESM_base_shear_kN'],r['force_scale'],r['scaled_modal_base_shear_kN'],maxdrift,limit,r['mass_participation_percent'],r['status'],str(path)])
        first=storey.max_row+1
        for v in r['stories']:storey.append([r['limit_state'],key,v['story_number'],v['height_m'],v['scaled_story_shear_kN'],v['elastic_drift_ratio'],v['design_drift_ratio'],v.get('master_design_drift_ratio'),limit,'PASS' if v['design_drift_ratio']<=limit else 'EXCEEDS'])
        if len(r['stories'])>1 and not r.get('accidental_eccentricity_sign',0):
            chart=ScatterChart();chart.title=r['limit_state']+' '+key+' storey shear';chart.x_axis.title='kN';chart.y_axis.title='Height (m)';chart.width=16;chart.height=10;chart.series.append(Series(Reference(storey,min_col=4,min_row=first,max_row=storey.max_row),Reference(storey,min_col=5,min_row=first,max_row=storey.max_row),title=key));summary.add_chart(chart,'O'+str(5+chart_index*20));chart_index+=1
        for i,v in enumerate(r['modal_rows']):modal.append([r['limit_state'],key,v['mode'],v['period_s'],v.get('frequency_Hz'),r['modal_base_shears_kN'][i],v.get('UX_participation_percent'),v.get('UY_participation_percent'),v.get('RZ_participation_percent'),v.get('UX_cumulative_percent'),v.get('UY_cumulative_percent')])
        deformation=path.parent/('Deformation_MRSM_'+r['limit_state']+'_'+key+'_V132.json')
        if deformation.is_file():
            d=json.loads(deformation.read_text(encoding='utf-8-sig'));mask=d.get('dof_availability',[True]*6)
            for tag,xyz in d['nodes'].items():
                u=d['displacements'].get(tag,[]);nodes.append([r['limit_state'],key,int(tag),*xyz,*[u[i] if i<len(u) and i<len(mask) and mask[i] else None for i in range(6)],'Non-concurrent positive CQC design component envelopes; missing legacy rotations left blank'])
    sources=wb.create_sheet('Sources and units');sources.append(['Source','Meaning']);sources.append(['https://ebps.kathmandu.gov.np/UploadFiles/NBC_105_2025_Sesmic_Design.pdf','NBC sections 4, 5.5.3 and 7']);sources.append(['https://opensees.github.io/OpenSeesDocumentation/user/manual/analysis/responseSpectrumAnalysis.html','Native modal responses combined with CQC']);sources.append(['kN-m-s; rotations rad','Dimensionless Cd converted to m/s2 using standard gravity 9.80665; displayed node translation envelopes are in metres.']);sources.append(['Case summary / Source JSON','Per-case numerical provenance; full member action and support diagnostics remain in sibling CSV files.'])
    for ws in wb:
        header=5 if ws==summary else 1;ws.freeze_panes='A'+str(header+1);ws.auto_filter.ref=f'A{header}:{ws.cell(ws.max_row,ws.max_column).coordinate}'
        for c in ws[header]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='16334C')
        for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=22
        if ws==summary:ws.column_dimensions['M'].width=65
    wb.save(folder/'MRSM_Results_V134.xlsx')
    for title,fig in figures(cases):
        stem=title.split()[0]+'_MRSM_Profiles_V134'
        for ext in ('png','pdf'):fig.savefig(folder/(stem+'.'+ext),dpi=180)
    return str(folder/'MRSM_Results_V134.xlsx')


def launch(parent,folder):
    import tkinter as tk
    from tkinter import ttk,messagebox
    from OpenSees_Research_UI_V123 import embed_figures
    cases=saved_cases(folder);win=tk.Toplevel(parent);win.title('Saved MRSM graphs — ULS / SLS');win.geometry('1220x850')
    ttk.Label(win,text='E0: centered | Eplus: +5% breadth | Eminus: -5% breadth. X/Y denote force axis. Similar base-shear magnitudes can have different torsion and edge drift. This viewer does not rerun analysis.',wraplength=1150).pack(anchor='w')
    def export():
        try:messagebox.showinfo('MRSM reports',export_reports(folder),parent=win)
        except Exception as e:messagebox.showerror('MRSM reports',str(e),parent=win)
    actions=ttk.Frame(win);actions.pack(fill='x');ttk.Button(actions,text='Export saved data to Excel / PNG / PDF',command=export).pack(side='left')
    def open_book(state):
        books=[p for p in Path(folder).rglob('MRSM_Results_V134.xlsx') if p.parent.name=='MRSM_'+state]
        if not books and (Path(folder)/'MRSM_Results_V134.xlsx').is_file():books=[Path(folder)/'MRSM_Results_V134.xlsx']
        if books:__import__('os').startfile(str(books[0]))
        else:messagebox.showinfo('MRSM Excel','Export '+state+' saved results first.',parent=win)
    for state in ('ULS','SLS'):ttk.Button(actions,text='Open '+state+' Excel',command=lambda state=state:open_book(state)).pack(side='left')
    body=ttk.Frame(win);body.pack(fill='both',expand=True)
    if cases:embed_figures(body,figures(cases))
    else:ttk.Label(body,text='Run MRSM or select the folder holding saved MRSM results.').pack()
    return win
