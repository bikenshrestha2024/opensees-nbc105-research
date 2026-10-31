"""Storey gravity ledger from exact assembled source nodal contributions."""
from pathlib import Path
import csv,json,math
from OpenSees_Timing_V130 import timed_outputs

COLUMNS=['beam_self_weight_kN','column_self_weight_kN','brace_self_weight_kN','shear_wall_self_weight_kN','slab_self_weight_kN','other_self_weight_kN','floor_finish_kN','partition_kN','other_dead_kN','live_kN','roof_live_kN','unclassified_kN']

def classify(component,pattern,pattern_type=None):
    if component.startswith('self_weight_'):
        kind=component[len('self_weight_'):].replace(' ','_')
        return {'beam':'beam','column':'column','brace':'brace','shear_wall':'shear_wall','slab':'slab','floor':'slab'}.get(kind,'other')+'_self_weight_kN','DEAD'
    name=pattern.casefold().replace('_',' ').replace('-',' ')
    if 'finish' in name:return 'floor_finish_kN','DEAD'
    if 'partition' in name:return 'partition_kN','DEAD'
    if 'roof' in name and ('live' in name or pattern_type==11):return 'roof_live_kN','LIVE'
    if 'live' in name or pattern_type in (3,4,11):return 'live_kN','LIVE'
    if pattern_type in (1,2) or name in ('dead','dl','sdl','super dead','superdead'):return 'other_dead_kN','DEAD'
    return 'unclassified_kN','UNCLASSIFIED'

@timed_outputs('Storey gravity load classification and workbook','folder')
def export_storey_gravity(folder,data,config,pattern_types=None):
    folder=Path(folder);folder.mkdir(parents=True,exist_ok=True);pattern_types=pattern_types or {}
    groups=sorted(data.get('diaphragms',[]),key=lambda g:g['master_xyz'][2]);nodes=data['nodes'];base=float(data['base_z_m'])
    rows={};details={}
    def row(key,z):
        if key not in rows:rows[key]=dict(storey=key,elevation_m=z,**{c:0. for c in COLUMNS},dead_total_kN=0.,live_total_kN=0.,unfactored_total_kN=0.,applied_dead_kN=0.,applied_live_kN=0.,applied_unclassified_kN=0.,applied_total_kN=0.)
        return rows[key]
    for item in data.get('load_provenance',[]):
        xyz=nodes.get(item['node'],nodes.get(str(item['node'])));z=float(xyz[2])
        g=next((g for g in groups if g['master_xyz'][2]>=z-1e-7),None)
        label='BASE' if z<=base+1e-7 else ('Storey '+str(g['storey_number']) if g else 'UNASSIGNED ABOVE FLOORS')
        level=base if label=='BASE' else (g['master_xyz'][2] if g else z)
        out=row(label,level);column,kind=classify(item.get('component','unknown'),item['pattern'],pattern_types.get(item['pattern']))
        value=item['downward_kN'];factor=float(config['gravity_factors'][item['pattern']]);applied=value*factor
        out[column]+=value;out['unfactored_total_kN']+=value;out['applied_total_kN']+=applied
        if kind=='DEAD':out['dead_total_kN']+=value;out['applied_dead_kN']+=applied
        elif kind=='LIVE':out['live_total_kN']+=value;out['applied_live_kN']+=applied
        else:out['applied_unclassified_kN']+=applied
        key=(label,item.get('source','unknown'),item.get('component','unknown'),item['pattern'],column,kind,factor)
        details[key]=details.get(key,0.)+value
    table=sorted(rows.values(),key=lambda r:r['elevation_m'])
    actual=sum(r['applied_total_kN'] for r in table);expected=float(data['total_downward_load_kN'])
    if not math.isclose(actual,expected,rel_tol=1e-10,abs_tol=1e-6):raise ValueError(f'Storey gravity ledger mismatch: {actual} vs {expected}')
    detail=[dict(storey=k[0],source=k[1],component=k[2],pattern=k[3],classification=k[4],load_class=k[5],gravity_factor=k[6],unfactored_kN=v,applied_kN=v*k[6]) for k,v in sorted(details.items())]
    note='Equivalent source nodal weights assigned to next diaphragm level above each node; base-node shares retained in BASE. Member weights spanning floors are apportioned by their consistent nodal forces. This is applied load by level, not storey axial force or cumulative storey reaction. Signed downward positive; upward negative. Unknown patterns remain unclassified.'
    report=dict(rows=table,details=detail,total_applied_kN=actual,assembled_total_kN=expected,reconciliation_error_kN=actual-expected,convention=note)
    (folder/'Gravity_Storey_Load_Breakdown_V131.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
    for name,values in [('Gravity_Storey_Load_Breakdown_V131.csv',table),('Gravity_Source_Load_Detail_V131.csv',detail)]:
        if values:
            with (folder/name).open('w',newline='',encoding='utf-8-sig') as f:
                w=csv.DictWriter(f,fieldnames=list(values[0]));w.writeheader();w.writerows(values)
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill
    wb=Workbook();wb.remove(wb.active)
    for title,values in [('Storey totals',table),('Source detail',detail)]:
        ws=wb.create_sheet(title)
        if values:
            ws.append(list(values[0]))
            for item in values:ws.append(list(item.values()))
            ws.freeze_panes='B2';ws.auto_filter.ref=ws.dimensions
            for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='1E3A5F')
            for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=min(42,max(16,len(str(col[0].value))+2))
    wb.create_sheet('Read me').append([note])
    from OpenSees_Load_Path_V141 import add_formula_audit
    add_formula_audit(wb,folder,data,config)
    wb.save(folder/'Gravity_Storey_Load_Breakdown_V131.xlsx')
    print(f'[STOREY GRAVITY] {len(table)} levels | applied={actual:.6f} kN | reconciliation={actual-expected:.6g} kN',flush=True)
    return report
