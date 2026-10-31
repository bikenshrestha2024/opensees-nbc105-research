"""Formula trace of applied shell loads; does not invent equivalent beam UDLs."""
from pathlib import Path
import json

def add_formula_audit(wb,folder,data=None,cfg=None):
    from openpyxl.styles import Font,PatternFill
    from openpyxl.workbook.properties import CalcProperties
    folder=Path(folder)
    def read(name):
        path=folder/name
        return json.loads(path.read_text(encoding='utf-8-sig')) if path.is_file() else {}
    data=read('OpenSees_Complete_Model_Data_V118.json') if data is None else data;cfg=read('OpenSees_Complete_Model_Config_V118.json') if cfg is None else cfg;audit=read('OpenSees_Load_Conversion_Audit.json')
    if not data:return
    # Formula arithmetic replaces only derived cells, never recorded solver DOFs.
    for title in ('Source detail','Source load detail'):
        if title in wb:
            ws=wb[title];cols={c.value:c.column for c in ws[1]}
            if all(k in cols for k in ('gravity_factor','unfactored_kN','applied_kN')):
                from openpyxl.utils import get_column_letter as letter
                for row in range(2,ws.max_row+1):ws.cell(row,cols['applied_kN'],f'={letter(cols["gravity_factor"])}{row}*{letter(cols["unfactored_kN"])}{row}')
    for title in ('Surface load formulas','Loaded node path','Member weight formulas','Vertical load path'):
        if title in wb:del wb[title]
    ws=wb.create_sheet('Surface load formulas');ws.append(['Source area','Component','Pattern','Mesh cell','Area m2','Downward q kN per m2','Gravity factor','Mass factor','Q unfactored kN','Applied kN','Mass tonnes','Node share kN','Nodes','Origin','Physical thickness m','Unit weight kN per m3','Self-weight multiplier'])
    links=wb.create_sheet('Loaded node path');links.append(['Source area','Pattern','Cell row','Node','X m','Y m','Z m','Applied nodal Fz kN downward','Connected beam solver tags','Connected wall solver tags','Role'])
    adjacency={};walls={}
    for tag,e in enumerate(data.get('frames',[]),100001):
        if e['category']=='BEAM':
            for n in e['nodes']:adjacency.setdefault(int(n),[]).append(tag)
    for tag,e in enumerate(data.get('shells',[]),100001+len(data.get('frames',[]))):
        if e['category']=='SHEAR WALL':
            for n in e['nodes']:walls.setdefault(int(n),[]).append(tag)
    pressure={}
    for item in audit.get('applied_area_loads',[]):pressure.setdefault(str(item['object']),[]).append(item)
    multipliers={r['pattern']:float(r['multiplier']) for r in audit.get('etabs_load_pattern_self_weight_multipliers',[]) if r.get('multiplier') is not None}
    surface_row=1
    for index,cell in enumerate(data.get('shells',[]),1):
        if cell.get('category')!='FLOOR':continue
        entries=[]
        for item in pressure.get(str(cell['object']),[]):entries.append((item['pattern'],'area_load',float(item['value'])*(1 if int(item['direction_code'])==10 else -1),item.get('load_set') or item.get('api','source')))
        for pattern,mult in multipliers.items():
            if mult:entries.append((pattern,'self_weight_floor',cell['properties']['thickness']*cell['properties']['weight']*mult,'physical thickness * source unit weight * ETABS self-weight multiplier'))
        for pattern,component,q,origin in entries:
            surface_row+=1;row=surface_row;own=component=='self_weight_floor';ws.append([cell['object'],component,pattern,index,cell['area'],f'=O{row}*P{row}*Q{row}' if own else q,cfg['gravity_factors'].get(pattern,0),cfg['mass_factors'].get(pattern,0),f'=E{row}*F{row}',f'=I{row}*G{row}',f'=I{row}*H{row}/9.80665',f'=J{row}/4',','.join(map(str,cell['nodes'])),origin,cell['properties']['thickness'] if own else None,cell['properties']['weight'] if own else None,multipliers.get(pattern) if own else None])
            for node in cell['nodes']:
                xyz=data['nodes'].get(int(node),data['nodes'].get(str(node))); beams=adjacency.get(int(node),[]);wall=walls.get(int(node),[])
                links.append([cell['object'],pattern,row,node,*xyz,f"='Surface load formulas'!L{row}",','.join(map(str,beams)),','.join(map(str,wall)),'beam joint' if beams else 'wall joint' if wall else 'slab interior / other joint'])
    weights=wb.create_sheet('Member weight formulas');weights.append(['Source frame','Solver tag','Category','Area m2','Length m','Unit weight kN per m3','Volume m3','Gross physical weight kN','Note'])
    for tag,e in enumerate(data.get('frames',[]),100001):
        r=tag-100001+2;weights.append([e['object'],tag,e['category'],e['properties']['A'],e['length'],e['properties']['weight'],f'=D{r}*E{r}',f'=G{r}*F{r}','Before self-weight pattern multiplier and gravity factor; overlap conventions require review'])
    floor_objects={str(e['object']) for e in data.get('shells',[]) if e.get('category')=='FLOOR'}
    for title in ('Source detail','Source load detail'):
        if title not in wb:continue
        sheet=wb[title];cols={c.value:c.column for c in sheet[1]}
        if not all(key in cols for key in ('source','pattern','component','unfactored_kN')):continue
        from openpyxl.utils import get_column_letter as letter
        end=ws.max_row
        for row in range(2,sheet.max_row+1):
            obj=str(sheet.cell(row,cols['source']).value).removeprefix('AREA:')
            comp=sheet.cell(row,cols['component']).value
            if obj in floor_objects and comp in ('area_load','self_weight_floor'):
                pattern=f'{letter(cols["pattern"])}{row}';component=f'{letter(cols["component"])}{row}'
                sheet.cell(row,cols['unfactored_kN'],f'=SUMIFS(\'Surface load formulas\'!I$2:I${end},\'Surface load formulas\'!A$2:A${end},"{obj}",\'Surface load formulas\'!C$2:C${end},{pattern},\'Surface load formulas\'!B$2:B${end},{component})')
    notes=wb.create_sheet('Vertical load path')
    for line in [
        'Area pressure -> shell mesh equivalent nodal loads -> shell bending/shear and connected beam/wall/column stiffness -> supports.',
        'rigidDiaphragm(3) constrains UX, UY and RZ. It does not constrain UZ or carry vertical pressure.',
        'A uniform four-node parallelogram shell has Q=q*A and a Q/4 vertical equivalent load at each mesh node. This is a finite-element load vector, not tributary beam UDL.',
        'No second q*tributary-width beam load is added: that would double-count the same area pressure.',
        'Loaded node path lists the actual connected beam and wall solver tags. Interior shell nodes transfer through shell stiffness.',
        'A corner-only panel at column joints cannot resolve beam-span bending. V141 adds floor subdivisions (default 4 per source-panel direction) and splits connected beams at shared mesh nodes.',
        'Refinement is a numerical choice, not an NBC-prescribed mesh. Compare 2/4/8 subdivisions and benchmark beam moments, slab deflection and reactions before design use.',
        'Gravity_Beam_End_Actions_ULS_V141.csv / SLS equivalent are written by complete-model preload. Gravity GUI action tabs also show local shears and bending moments. Beam end actions and shell forces are solver outputs; they cannot be replaced by q*A arithmetic or treated as a unique tributary UDL. Use gravity action reports for actual response.',
        'Gravity factor is preload; it is not automatically the complete design live-load combination. Mass factor is separately shown. Roof-live preload may be zero while the source roof-live pressure remains visible.',
        'https://openseespydoc.readthedocs.io/en/latest/src/ShellMITC4.html',
        'https://opensees.berkeley.edu/OpenSees/api/doxygen2/html/RigidDiaphragm_8cpp-source.html',
        'https://opensees.berkeley.edu/OpenSees/api/doxygen2/html/ShellMITC4_8cpp-source.html'] :notes.append([line])
    import csv
    for state in ('ULS','SLS'):
        path=folder/('Gravity_Beam_End_Actions_'+state+'_V141.csv')
        if path.is_file():
            title='Beam end actions '+state
            if title in wb:del wb[title]
            output=wb.create_sheet(title)
            with path.open(encoding='utf-8-sig',newline='') as stream:
                for row in csv.reader(stream):output.append(row)
            output.freeze_panes='A2';output.auto_filter.ref=output.dimensions
    notes.column_dimensions['A'].width=140
    for sheet in (ws,links,weights):
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for c in sheet[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='1E3A5F')
        for col in sheet.columns:sheet.column_dimensions[col[0].column_letter].width=min(55,max(16,len(str(col[0].value))+2))
    from openpyxl.utils import get_column_letter as letter
    source_title='Source load detail' if 'Source load detail' in wb else 'Source detail' if 'Source detail' in wb else None
    if source_title:
        detail=wb[source_title];dc={c.value:letter(c.column) for c in detail[1]};last=detail.max_row
        def rng(key):return "'"+source_title+"'!"+dc[key]+"$2:"+dc[key]+"$"+str(last)
        for title in ('Storey totals','Storey load breakdown'):
            if title not in wb:continue
            sheet=wb[title];sc={c.value:c.column for c in sheet[1]}
            for row in range(2,sheet.max_row+1):
                st=letter(sc['storey'])+str(row)
                for key,col in sc.items():
                    if key in ('storey','elevation_m'):continue
                    applied=key.startswith('applied_');amount='applied_kN' if applied else 'unfactored_kN'
                    formula='=SUMIFS('+rng(amount)+','+rng('storey')+','+st
                    if key in ('dead_total_kN','applied_dead_kN'):formula+=','+rng('load_class')+',"DEAD"'
                    elif key in ('live_total_kN','applied_live_kN'):formula+=','+rng('load_class')+',"LIVE"'
                    elif key=='applied_unclassified_kN':formula+=','+rng('load_class')+',"UNCLASSIFIED"'
                    elif key not in ('unfactored_total_kN','applied_total_kN'):formula+=','+rng('classification')+',"'+key+'"'
                    sheet.cell(row,col,formula+')')
            if 'Summary' in wb and title=='Storey load breakdown':
                col=letter(sc['applied_total_kN']);wb['Summary']['B4']=f"=SUM('{title}'!{col}2:{col}{sheet.max_row})"
    if 'Summary' in wb and 'Support Reactions' in wb:
        wb['Summary']['B5']=f"=SUM('Support Reactions'!D2:D{wb['Support Reactions'].max_row})"
        wb['Summary']['B6']='=IF(B4=0,"REVIEW ZERO LOAD",ABS(B5-B4)/ABS(B4)*100)'
    from OpenSees_Mapping_Audit_V125 import parameter_sources,source_parameter_inventory
    title='Material unit formulas'
    if title in wb:del wb[title]
    units=wb.create_sheet(title);units.append(['Family','Parameter','Source value','Source unit','Target unit','Converted value','Validation status','Source'])
    for item in parameter_sources()+source_parameter_inventory(folder,data):
        value=item.get('value');unit=item.get('unit')
        if not isinstance(value,(float,int)) or unit not in ('MPa','ksi','kN/m2','mm'):continue
        row=units.max_row+1;target='m' if unit=='mm' else 'kN/m2' if unit=='MPa' else 'MPa'
        expr=f'=C{row}*1000' if unit=='MPa' else f'=C{row}*6.894757293' if unit=='ksi' else f'=C{row}/1000'
        units.append([item.get('family'),item.get('parameter'),value,unit,target,expr,item.get('status'),item.get('source')])
    units.freeze_panes='A2';units.auto_filter.ref=units.dimensions
    wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True)


def export(folder):
    from openpyxl import Workbook
    folder=Path(folder);wb=Workbook();wb.remove(wb.active);add_formula_audit(wb,folder)
    if not wb.worksheets:raise ValueError('Rebuild the assembled model first')
    path=folder/'Slab_Load_Path_Formulas_V141.xlsx';wb.save(path);return path


def launch(parent,folder):
    import tkinter as tk
    from tkinter import ttk,messagebox
    folder=Path(folder);config=folder/'OpenSees_Complete_Model_Config_V118.json'
    cfg=json.loads(config.read_text(encoding='utf-8-sig')) if config.is_file() else {}
    win=tk.Toplevel(parent);win.title('Slab vertical load path / mesh and formula audit');win.geometry('850x360')
    ttk.Label(win,text='Vertical loads act on ShellMITC4 mesh nodes. Rigid diaphragm UX/UY/RZ constraints do not carry gravity. Shared mesh nodes connect slab bending to beam spans and walls. Beam force distributions come from the solver, not an extra tributary UDL.',wraplength=800).pack(padx=12,pady=10)
    choice=tk.StringVar(value=str(cfg.get('floor_subdivisions',4)))
    ttk.Label(win,text='Subdivisions in each source slab-panel direction (4 starting mesh; compare 8 and 16 for convergence)').pack(anchor='w',padx=12)
    ttk.Combobox(win,textvariable=choice,values=['2','4','8','16'],state='readonly').pack(anchor='w',padx=12)
    def save():
        if not config.is_file():messagebox.showerror('Mesh','Export the model first.',parent=win);return
        cfg['floor_subdivisions']=int(choice.get());config.write_text(json.dumps(cfg,indent=2),encoding='utf-8')
        messagebox.showinfo('Mesh setting saved','Rebuild the assembled export, then rerun gravity and other analyses. Old results are not updated by this setting.',parent=win)
    def workbook():
        try:__import__('os').startfile(str(export(folder)))
        except Exception as exc:messagebox.showerror('Load path',str(exc),parent=win)
    ttk.Button(win,text='Save mesh choice (rebuild required)',command=save).pack(padx=12,pady=8,anchor='w')
    ttk.Button(win,text='Export / open current load-path formulas',command=workbook).pack(padx=12,anchor='w')
    return win
