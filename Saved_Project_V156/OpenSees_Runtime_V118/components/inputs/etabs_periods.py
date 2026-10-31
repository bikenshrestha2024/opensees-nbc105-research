"""ETABS session opening and explicitly attributed X/Y period references."""

__release__ = '134'
from pathlib import Path
import json
import math

_sessions=[]


def source_edb(folder):
    folder=Path(folder)
    for name in ('OpenSees_Element_Connectivity.json','OpenSees_Section_Assignment_Audit.json'):
        file=folder/name
        if file.is_file():
            value=json.loads(file.read_text(encoding='utf-8-sig')).get('source_edb')
            if value and Path(value).is_file():return str(Path(value).resolve())
    return None


def attach_or_open(folder,open_if_needed=False):
    from OpenSees_Project_Inputs_V123 import connect_etabs
    desired=source_edb(folder)
    try:
        app=connect_etabs()
        current=str(app.SapModel.GetModelFilename())
        if not open_if_needed or not desired or Path(current).resolve()==Path(desired).resolve():
            _sessions.append(app);return app
    except Exception:
        if not open_if_needed:raise
    if not desired:raise ValueError('The source EDB path is unavailable. Select/export the intended EDB in the main window first.')
    import comtypes.client
    helper=comtypes.client.CreateObject('ETABSv1.Helper')
    try:
        from comtypes.gen import ETABSv1
        helper=helper.QueryInterface(ETABSv1.cHelper)
    except ImportError:pass
    app=helper.CreateObjectProgID('CSI.ETABS.API.ETABSObject')
    rc=app.ApplicationStart()
    if rc not in (None,0):raise RuntimeError('ETABS ApplicationStart returned '+str(rc))
    rc=app.SapModel.File.OpenFile(desired)
    if rc!=0:raise RuntimeError('ETABS could not open the source EDB; return code '+str(rc))
    _sessions.append(app)
    return app


def workbook_references(root,folder):
    from OpenSees_Project_Inputs_V123 import source_workbooks
    from openpyxl import load_workbook
    result={}
    for path in source_workbooks(root,folder):
        if path is None:continue
        wb=load_workbook(path,data_only=True,read_only=True)
        if 'NBC Calculator' in wb:
            ws=wb['NBC Calculator']
            for axis,cell in [('X','B16'),('Y','B17')]:
                value=ws[cell].value
                if isinstance(value,(int,float)) and math.isfinite(value) and value>0:
                    result['optional_'+axis]=dict(value=value,kind='Workbook ETABS/Rayleigh reference (ambiguous label)',source=str(path),sheet=ws.title,cell=cell)
        if 'SCALING SUMMARY' in wb:
            for axis,cell in [('X','C14'),('Y','C15')]:
                value=wb['SCALING SUMMARY'][cell].value
                if isinstance(value,(int,float)) and math.isfinite(value) and value>0:
                    result['adopted_'+axis]=dict(value=value,kind='Workbook NBC adopted reference, not eigenperiod',source=str(path),sheet='SCALING SUMMARY',cell=cell)
        # A clearly named modal export has a mode/period/UX/UY table.
        for ws in wb:
            if 'modal' not in ws.title.lower() and 'participating' not in ws.title.lower():continue
            records=list(ws.iter_rows(max_row=250,values_only=True))
            for index,row in enumerate(records[:12]):
                headers=[str(c or '').strip().lower() for c in row]
                period=next((i for i,h in enumerate(headers) if h in ('period','period_s','period (s)','period (sec)')),None)
                ux=next((i for i,h in enumerate(headers) if h in ('ux','ux_ratio','ux participation')),None)
                uy=next((i for i,h in enumerate(headers) if h in ('uy','uy_ratio','uy participation')),None)
                if None in (period,ux,uy):continue
                candidates=[]
                for n,r in enumerate(records[index+1:],index+2):
                    try:
                        p,x,y=map(float,(r[period],r[ux],r[uy]))
                        if p>0 and all(math.isfinite(v) for v in (p,x,y)):candidates.append((p,x,y,n))
                    except (TypeError,ValueError,IndexError):continue
                for axis,col in [('X',1),('Y',2)]:
                    if candidates:
                        selected=max(candidates,key=lambda r:r[col])
                        result['modal_'+axis]=dict(value=selected[0],kind='ETABS modal table in workbook',source=str(path),sheet=ws.title,cell=f'row {selected[3]}, period column {period+1}')
                break
        wb.close()
    return result


def populate(gui,root,folder):
    refs=workbook_references(root,folder)
    for axis in ('X','Y'):
        r=refs.get('adopted_'+axis)
        if r:getattr(gui,'excel_t'+axis.lower()).set(f"{r['value']:.9g} s — {r['sheet']}!{r['cell']}")
        r=refs.get('modal_'+axis) or refs.get('optional_'+axis) or refs.get('adopted_'+axis)
        if r:getattr(gui,'etabs_t'+axis.lower()).set(f"{r['value']:.9g}")
    cache=Path(folder)/'ETABS_Modal_Periods_V125.json'
    if cache.is_file():
        saved=json.loads(cache.read_text(encoding='utf-8-sig'))
        actual=source_edb(folder)
        if actual and saved.get('source_edb')==actual and saved.get('source_edb_mtime_ns')==Path(actual).stat().st_mtime_ns:
            for axis in ('X','Y'):
                value=saved.get(axis,{}).get('period_s')
                if isinstance(value,(int,float)) and value>0:getattr(gui,'etabs_t'+axis.lower()).set(str(value))
            refs['ETABS_cache']=saved
    gui.excel_period_references=refs
    return refs


def save_live(folder,tx,ty,metadata):
    file=source_edb(folder)
    result=dict(source_edb=file,source_edb_mtime_ns=Path(file).stat().st_mtime_ns if file else None,
                X=dict(period_s=float(tx)),Y=dict(period_s=float(ty)),metadata=metadata)
    Path(folder,'ETABS_Modal_Periods_V125.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
    return result
