"""Grade selection is metadata, never automatic experimental calibration."""
from pathlib import Path
import json,csv,subprocess,sys,datetime,math,webbrowser
SOURCES={
 'Single-coupon deterministic / Bayesian calibration':'https://nheri-simcenter.github.io/quoFEM-Documentation/common/user_manual/examples/desktop/qfem-0019/README.html',
 'RC column surrogate calibration':'https://nheri-simcenter.github.io/quoFEM-Documentation/common/user_manual/examples/desktop/qfem-0034/README.html',
 'Concrete01':'https://opensees.github.io/OpenSeesDocumentation/user/manual/material/uniaxialMaterials/Concrete01.html',
 'Steel02':'https://opensees.github.io/OpenSeesDocumentation/user/manual/material/uniaxialMaterials/Steel02.html',
 'Experimental Steel02 calibration':'https://nheri-simcenter.github.io/quoFEM-Documentation/common/user_manual/examples/desktop/qfem-0028/README.html',
 'CarreÃ±o et al. 2020':'https://doi.org/10.1061/(ASCE)ST.1943-541X.0002505',
 'Confined concrete experiments, Mander et al. 1988':'https://doi.org/10.1061/(ASCE)0733-9445(1988)114:8(1827)',
 'FVD':'https://opensees.berkeley.edu/wiki/index.php?title=ViscousDamper_Material',
 'BRB tests':'https://peer.berkeley.edu/publications/2002-08',
 'E-SFI wall formulation':'https://kkolozvari.github.io/E-SFI-MVLEM-3D/' }
def card(model,grade):
 params={'fpc_MPa':None,'epsc0':None,'fpcu_MPa':None,'epsU':None} if model=='Concrete01' else dict(Fy_MPa=None,E_MPa=None,b=None,R0=20.,cR1=.925,cR2=.15,a1=None,a2=1.,a3=None,a4=1.)
 return dict(model=model,grade=grade,status='UNCALIBRATED INPUT CARD',parameters=params,experimental_source='',specimen_id='',strength_basis='Measured cylinder compression required; M label alone is cube grade' if model=='Concrete01' else 'Measured cyclic coupon; Fe grade is not a Steel02 calibration',protocol_csv='',units='strain, MPa',sources=SOURCES,notes='Null values require measured or explicitly justified parameters. R0/cR values are documentation recommendations, not Fe-specific experiments. UCSD Grade60 ASTM A615/A706 is not equivalent to Fe415/500/550 without evidence. FVD/BRB/walls require independent component parameters.')
def simulate(data,out):
 from components.validation.visual_materials import uniaxial,ramp,write
 pars=data['parameters'];keys=['fpc_MPa','epsc0','fpcu_MPa','epsU'] if data['model']=='Concrete01' else ['Fy_MPa','E_MPa','b','R0','cR1','cR2','a1','a2','a3','a4'];values=[float(pars[k]) for k in keys]
 if not all(math.isfinite(v) for v in values):raise ValueError('Finite parameters required')
 if data['model']=='Concrete01':
  f,e,u,eu=values
  if not(f<0 and e<0 and u<=0 and abs(u)<=abs(f) and eu<e):raise ValueError('Use negative compression; epsU < epsc0 < 0 and residual magnitude <= peak')
  protocol=ramp([e*.25,0,e*.5,0,e,0,eu,0,.001])
 else:
  if values[0]<=0 or values[1]<=0 or not 0<=values[2]<1:raise ValueError('Positive Fy/E and 0<=b<1 required')
  protocol=ramp([.0005,-.0005,.002,-.002,.005,-.005,.01,-.01,0])
 if data.get('protocol_csv'):
  with Path(data['protocol_csv']).open(encoding='utf-8-sig',newline='') as stream:
   protocol=[float(r['strain']) for r in csv.DictReader(stream)]
  if len(protocol)<2 or not all(math.isfinite(e) for e in protocol):raise ValueError('Protocol needs at least two finite strain samples in original order')
 out=Path(out);out.mkdir(parents=True,exist_ok=True);rows=uniaxial(data['model'],values,protocol);write(out/'cyclic.csv',rows)
 from matplotlib.figure import Figure
 from matplotlib.backends.backend_agg import FigureCanvasAgg
 fig=Figure(figsize=(9,5),layout='constrained');FigureCanvasAgg(fig);ax=fig.add_subplot();ax.plot([r['x'] for r in rows],[r['y'] for r in rows]);ax.set(xlabel='Strain',ylabel='Stress (MPa)',title=data['grade']+' '+data['model']+' - supplied parameters, NOT experimental acceptance');ax.grid(alpha=.25);fig.savefig(out/'cyclic.png',dpi=150)
 # Runtime exporter, consistent with existing desktop workbook exporters.
 from openpyxl import Workbook
 from openpyxl.chart import ScatterChart,Series,Reference
 wb=Workbook();ws=wb.active;ws.title='Cyclic curve';ws.append(['Step','Strain','Stress MPa','Tangent MPa','Stress kN per m2'])
 for i,r in enumerate(rows,2):ws.append([r['step'],r['x'],r['y'],r['tangent'],'=C'+str(i)+'*1000'])
 ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
 chart=ScatterChart();chart.title='Supplied parameter cyclic response';chart.x_axis.title='Strain';chart.y_axis.title='Stress MPa';chart.series.append(Series(Reference(ws,min_col=3,min_row=2,max_row=len(rows)+1),Reference(ws,min_col=2,min_row=2,max_row=len(rows)+1)));ws.add_chart(chart,'G2')
 ps=wb.create_sheet('Parameters and sources');ps.append(['Item','Value','Basis'])
 for k in keys:ps.append([k,pars[k],'User supplied; calibration not established'])
 for k,v in SOURCES.items():ps.append([k,v,'Reference scope must match specimen'])
 ps.append(['Grade',data['grade'],'Label is not measured cyclic behaviour']);ps.append(['Initial modulus formula', '=2*ABS(B2)/ABS(B3)' if data['model']=='Concrete01' else '=B3','Concrete01 initial tangent or Steel02 supplied E, MPa'])
 for sheet in wb:
  for col in sheet.columns:sheet.column_dimensions[col[0].column_letter].width=28
 wb.save(out/'Material_Curve_V151.xlsx')
 data=dict(data,finished_at=datetime.datetime.now().astimezone().isoformat(),status='NATIVE CURVE GENERATED - CALIBRATION NOT ESTABLISHED');(out/'card.json').write_text(json.dumps(data,indent=2),encoding='utf8');return rows

def launch(parent):
 import tkinter as tk
 from tkinter import ttk,filedialog,messagebox
 from components.validation.visual_workbench import playback
 win=tk.Toplevel(parent);win.title('Grade cards, experimental sources and native cyclic curves - V151');win.geometry('1100x760');bar=ttk.Frame(win);bar.pack(fill='x');model=tk.StringVar(value='Concrete01');grade=tk.StringVar(value='M30')
 ttk.Combobox(bar,textvariable=model,values=['Concrete01','Steel02'],state='readonly',width=15).pack(side='left');choice=ttk.Combobox(bar,textvariable=grade,values=['M30','M35','M40','M45'],state='readonly',width=10);choice.pack(side='left')
 ttk.Label(win,text='Select grade, enter specimen-specific parameters and provenance. No verified experimental card for these exact grades is bundled. Published UCSD Grade60 calibration data is linked below. MPa coupon units; structural stress conversion is x1000 to kN/m2.',wraplength=1050).pack(fill='x')
 edit=tk.Text(win,wrap='word',height=18);edit.pack(fill='both',expand=True);latest={}
 def reset(*args):
  options=['M30','M35','M40','M45'] if model.get()=='Concrete01' else ['Fe415','Fe500','Fe550'];choice.configure(values=options)
  if grade.get() not in options:grade.set(options[0])
  edit.delete('1.0','end');edit.insert('1.0',json.dumps(card(model.get(),grade.get()),indent=2))
 model.trace_add('write',reset);grade.trace_add('write',lambda *a:None)
 def published():
  model.set('Steel02');data=card('Steel02','Terzic-Stojadinovic circular-column specimen')
  data.update(status='PUBLISHED ANALYTICAL CARD - NOT PROJECT CALIBRATION',experimental_source='10.1061/(ASCE)BE.1943-5592.0000702, Table 3 and Fig 4',specimen_id='Analytical model 1, circular bridge-column series')
  data['parameters']=dict(Fy_MPa=490.,E_MPa=200000.,b=.025,R0=15.,cR1=.925,cR2=.15,a1=0.,a2=1.,a3=0.,a4=1.)
  data['notes']='Six Table 3 parameters transcribed. a1/a3=0,a2/a4=1 are explicit implementation defaults, not additional fitted paper values. Table 2 measured steel487MPa differs from Table3 analytical490MPa. Not Fe500 or wall calibration; generated default strain history is illustrative, not the experimental protocol. Supply matching measured strain CSV for reproduction.'
  edit.delete('1.0','end');edit.insert('1.0',json.dumps(data,indent=2))
 def references():
  import os
  os.startfile(str(Path(__file__).resolve().parents[2]/'Calibration_References_V152'))
 ttk.Button(bar,text='Published Steel02 specimen card',command=published).pack(side='left')
 ttk.Button(win,text='Open supplied calibration papers and source cards',command=references).pack(anchor='w')
 def run():
  try:
   data=json.loads(edit.get('1.0','end'));out=filedialog.askdirectory(parent=win,title='New material curve output folder')
   if not out:return
   # Native material tester runs separately, so it cannot wipe a building domain.
   out=Path(out)/('GradeCurve_'+datetime.datetime.now().strftime('%Y%m%d_%H%M%S'));out.mkdir(parents=True);cfg=out/'input.json';cfg.write_text(json.dumps(data),encoding='utf8')
   script='from components.validation.grade_cards import simulate;import json;simulate(json.load(open('+repr(str(cfg))+')), '+repr(str(out))+')'
   proc=subprocess.run([sys.executable,'-c',script],cwd=str(Path(__file__).resolve().parents[2]),capture_output=True,text=True)
   (out/'run.log').write_text(proc.stdout+proc.stderr,encoding='utf8')
   if proc.returncode:raise ValueError(proc.stderr)
   latest['out']=out;playback(win,out/'cyclic.csv','Strain','Stress (MPa)',data['grade']+' supplied-parameter cyclic response')
  except Exception as e:messagebox.showerror('Material card inputs',str(e),parent=win)
 ttk.Button(bar,text='Populate blank grade card',command=reset).pack(side='left');ttk.Button(bar,text='Run / graph / animate supplied card',command=run).pack(side='left')
 for label,url in SOURCES.items():ttk.Button(win,text=label,command=lambda u=url:webbrowser.open(u)).pack(anchor='w')
 reset();return win
