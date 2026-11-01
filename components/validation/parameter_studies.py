"""Individual native parameter studies. Examples are not project-calibrated materials."""
from pathlib import Path
import json,csv,math,time,sys,argparse,datetime
DOC='https://openseespydoc.readthedocs.io/en/latest/src/'
STEEL=dict(Fy_MPa=355.,E_MPa=200000.,b=.01,R0=20.,cR1=.925,cR2=.15,a1=0.,a2=1.,a3=0.,a4=1.)
CONC=dict(fpc_MPa=-30.,epsc0=-.002,fpcu_MPa=-6.,epsU=-.006)
CM=dict(fpcc=-30.,epcc=-.002,Ec=30000.,rc=7.,xcrn=1.035,ft=2.,et=.00008,rt=1.2,xcrp=10000.)
MPF=dict(fyp=355.,fyn=355.,E0=200000.,bp=.01,bn=.01,R0=20.,cR1=.925,cR2=.15,a1=0.,a2=1.,a3=0.,a4=1.)
SCHEMAS={
'Concrete01':(CONC,'Concrete01.html'),
'Concrete02':(dict(CONC,unloading_ratio=.1,ft_MPa=2.,Ets_MPa=200.),'Concrete02.html'),
'ConcreteCM':(CM,'ConcreteCM.html'),
'Steel02':(STEEL,'steel02.html'),
'SteelMPF':(MPF,'SteelMPF.html'),
'FVD':(dict(K_kN_m=300000.,Cd=280.3*1000**.3,alpha=.3),'ViscousDamper.html'),
'BRB':(dict(STEEL,core_area_m2=.002,yielding_length_m=3.),'corotTruss.html'),
'SFI_MVLEM_3D':(dict(width_m=1.2,height_m=.3,thickness_m=.15,fibers=5,axial_kN=60.,rho_x=.01,rho_y=.01,nu=.25,CoR=.4,ThickMod=.63,steel_Fy_MPa=355.,steel_E_MPa=200000.,concrete_fc_MPa=-30.),'SFI_MVLEM_3D.html')}
def card(model):
 data=dict(model=model,parameters=dict(SCHEMAS[model][0]),parameter_basis='ILLUSTRATIVE STARTING VALUES - replace with specimen-specific measurements; not a verified grade or wall specimen',specimen_id='',citation='',source=DOC+SCHEMAS[model][1],protocol_csv='',measured_csv='',fit_bounds={},train_fraction=.75,scope='Ordered same-specimen fit is not independent validation or project design approval',wall_formulation='SFI_MVLEM_3D',targets=([.0005,-.0005,.002,-.002,.005,-.005,0.] if model.startswith('Steel') else [.001,-.001,.003,-.003,0.] if model=='BRB' else [.00001,-.00001,.0001,-.0001,0.] if model=='SFI_MVLEM_3D' else [-.0005,0.,-.001,0.,-.002,0.,-.004,0.,-.006,0.,.001]),samples_per_leg=30,dt_s=.005)
 if model=='SFI_MVLEM_3D':data.update(wall_concrete=dict(CM),wall_steel=dict(MPF),fsam=dict(nu=.1,alpha=.01),boundary_scope='Base fixed; top out-of-plane and torsional DOFs restrained; equal top UX. Numerical planar patch, not a nonplanar core specimen.')
 return data
def check(d):
 p=d['parameters'];model=d['model']
 for k in SCHEMAS[model][0]:
  if k not in p or not math.isfinite(float(p[k])):raise ValueError('Finite parameter required: '+k)
 if int(d.get('samples_per_leg',30))<1:raise ValueError('samples_per_leg must be positive')
 if model.startswith('Concrete'):
  f=p.get('fpc_MPa',p.get('fpcc'));e=p.get('epsc0',p.get('epcc'))
  if f>=0 or e>=0:raise ValueError('Concrete compression strength and peak strain must be negative')
  if 'epsU' in p and not(p['epsU']<e and f<=p['fpcu_MPa']<=0):raise ValueError('Require epsU < epsc0 and fpc <= fpcu <= 0')
  if model=='Concrete02' and not(0<=p['unloading_ratio']<=1 and p['ft_MPa']>=0 and p['Ets_MPa']>0):
   raise ValueError('Concrete02 requires 0<=unloading_ratio<=1, ft_MPa>=0 and Ets_MPa>0')
  if model=='ConcreteCM' and not(p['Ec']>0 and p['rc']>0 and p['xcrn']>0 and p['ft']>=0 and p['et']>0 and p['rt']>0 and p['xcrp']>0):
   raise ValueError('ConcreteCM requires positive Ec/shape/transition inputs and nonnegative tensile strength')
 if model in ('Steel02','BRB') and not(p['Fy_MPa']>0 and p['E_MPa']>0 and 0<=p['b']<1 and all(p[k]>0 for k in ('R0','cR1','cR2'))):raise ValueError('Require positive Fy/E/transition values and 0<=b<1')
 if model=='SteelMPF' and not(all(p[k]>0 for k in ('fyp','fyn','E0','R0','cR1','cR2')) and all(0<=p[k]<1 for k in ('bp','bn'))):raise ValueError('Invalid SteelMPF strength/modulus/hardening/transition values')
 if model=='FVD' and not(p['K_kN_m']>0 and p['Cd']>0 and 0<p['alpha']<=1 and float(d['dt_s'])>0):raise ValueError('FVD requires K>0, Cd>0, 0<alpha<=1 and dt>0')
 if model=='BRB' and not(p['core_area_m2']>0 and p['yielding_length_m']>0):raise ValueError('BRB area and yielding length must be positive')
 if model=='SFI_MVLEM_3D' and not(all(p[k]>0 for k in ('width_m','height_m','thickness_m','steel_Fy_MPa','steel_E_MPa')) and int(p['fibers'])==p['fibers'] and p['fibers']>=2 and 0<p['rho_x']<1 and 0<p['rho_y']<1 and p['axial_kN']>=0 and p['concrete_fc_MPa']<0):raise ValueError('Invalid panel geometry, reinforcement or material values')
def protocol(d):
 from components.validation.visual_materials import ramp
 if d.get('protocol_csv'):
  with Path(d['protocol_csv']).open(encoding='utf-8-sig') as f:r=list(csv.DictReader(f))
  return [float(a['x']) for a in r]
 return ramp(d['targets'],int(d.get('samples_per_leg',30)))
def native(d,x):
 check(d);p=d['parameters'];m=d['model']
 from components.validation.visual_materials import uniaxial,damper_history
 if m=='FVD':
  if d.get('protocol_csv'):raise ValueError('FVD uses the documented 12/24/36 mm, 0.5 Hz imposed history; custom displacement-time protocol is not implemented here')
  return damper_history(p['K_kN_m'],p['Cd'],p['alpha'],float(d['dt_s']))
 if m=='SFI_MVLEM_3D':return wall(d,x)
 keys=list(SCHEMAS[m][0])
 if m=='BRB':keys=list(STEEL);x=[e/p['yielding_length_m'] for e in x]
 rows=uniaxial('Steel02' if m=='BRB' else m,[p[k] for k in keys],x)
 if m=='BRB':
  for r in rows:r['x']*=p['yielding_length_m'];r['y']*=1000*p['core_area_m2']
 return rows
def wall(d,x):
 import openseespy.opensees as o
 partial_dir=Path(d.get('_output','.'))
 partial_dir.mkdir(parents=True,exist_ok=True)
 p=d['parameters'];w=p['width_m'];h=p['height_m'];n=int(p['fibers']);o.wipe();o.model('basic','-ndm',3,'-ndf',6)
 for tag,xyz in [(1,(0,0,0)),(2,(w,0,0)),(3,(w,0,h)),(4,(0,0,h))]:o.node(tag,*xyz)
 for tag in (1,2):o.fix(tag,1,1,1,1,1,1)
 for tag in (3,4):o.fix(tag,0,1,0,1,0,1)
 o.equalDOF(3,4,1)
 steel=dict(d.get('wall_steel',MPF));steel.update(fyp=p['steel_Fy_MPa'],fyn=p['steel_Fy_MPa'],E0=p['steel_E_MPa']);o.uniaxialMaterial('SteelMPF',1,*[steel[k]*(1000 if k in ('fyp','fyn','E0') else 1) for k in MPF])
 c=d.get('wall_concrete',CM);v=dict(c);v['fpcc']=p['concrete_fc_MPa'];args=[v[k]*(1000 if k in ('fpcc','Ec','ft') else 1) for k in CM];o.uniaxialMaterial('ConcreteCM',2,*args)
 o.nDMaterial('FSAM',3,0.,1,1,2,p['rho_x'],p['rho_y'],float(d.get('fsam',{}).get('nu',.1)),float(d.get('fsam',{}).get('alpha',.01)))
 element=d.get('wall_formulation','SFI_MVLEM_3D')
 if element not in ('SFI_MVLEM_3D','E_SFI_MVLEM_3D'):raise ValueError('Unsupported wall formulation')
 o.element(element,1,1,2,3,4,n,'-thick',*[p['thickness_m']]*n,'-width',*[w/n]*n,'-mat',*[3]*n,'-CoR',p['CoR'],'-ThickMod',p['ThickMod'],'-Poisson',p['nu'],'-Density',0.)
 o.timeSeries('Linear',1);o.pattern('Plain',1,1)
 for tag in (3,4):o.load(tag,0,0,-p['axial_kN']/2,0,0,0)
 o.constraints('Transformation');o.numberer('RCM');o.system('UmfPack');o.test('NormDispIncr',1e-9,80);o.algorithm('Newton');o.integrator('LoadControl',.1);o.analysis('Static')
 rc=o.analyze(10)
 if rc:raise RuntimeError(element+' gravity failed: '+str(rc))
 o.loadConst('-time',0);o.timeSeries('Linear',2);o.pattern('Plain',2,2);o.load(3,1,0,0,0,0,0)
 rows=[]
 for i,u in enumerate(x):
  du=u-o.nodeDisp(3,1)
  if abs(du)>1e-14:
   o.integrator('DisplacementControl',3,1,du);rc=o.analyze(1)
   if rc:
    import csv
    partial=partial_dir/'partial_wall_response.csv'
    if rows:
     with partial.open('w',newline='') as f:
      writer=csv.DictWriter(f,fieldnames=list(rows[0]));writer.writeheader();writer.writerows(rows)
    raise RuntimeError(element+' cyclic step '+str(i)+' failed: '+str(rc)+'; partial response retained')
  o.reactions();rows.append(dict(step=i,x=o.nodeDisp(3,1),y=-sum(o.nodeReaction(tag,1) for tag in (1,2))))
 o.wipe();return rows

def run(d,out):
 import numpy as np
 from scipy.optimize import least_squares
 from components.validation.visual_materials import write
 start=time.perf_counter();out=Path(out);out.mkdir(parents=True,exist_ok=True);check(d);d['_output']=str(out);x=protocol(d);measured=None;fit={}
 if d.get('measured_csv'):
  if not d.get('citation') or not d.get('specimen_id'):raise ValueError('Measured comparison requires citation and specimen ID')
  with Path(d['measured_csv']).open(encoding='utf-8-sig') as f:ref=list(csv.DictReader(f))
  measured=np.array([float(r['y']) for r in ref]);xr=[float(r['x']) for r in ref]
  if not np.isfinite(measured).all() or not np.isfinite(xr).all():raise ValueError('Nonfinite measured data')
  if d['model']=='FVD':raise ValueError('Custom FVD experimental protocol alignment is not yet implemented; use numerical response only')
  x=xr
 bounds=d.get('fit_bounds',{})
 if bounds:
  if measured is None:raise ValueError('Tuning requires measured x,y CSV with provenance')
  if d['model'] in ('FVD','SFI_MVLEM_3D'):raise ValueError('Automatic optimizer is for uniaxial materials and idealized BRB core only; use manual component studies')
  keys=list(bounds);lo=[];hi=[];initial=[]
  for k in keys:
   if k not in d['parameters']:raise ValueError('Unknown fitted parameter '+k)
   a,b=map(float,bounds[k]);v=d['parameters'][k]
   if not(math.isfinite(a) and math.isfinite(b) and a<v<b):raise ValueError('Bounds must be finite and enclose starting value: '+k)
   lo.append(a);hi.append(b);initial.append(v)
  fraction=float(d.get('train_fraction',.75));cut=int(len(x)*fraction)
  if not 0<fraction<1 or cut<max(3,len(keys)+1) or len(x)-cut<2:raise ValueError('Need adequate training and withheld samples')
  def residual(values):
   trial=dict(d,parameters=dict(d['parameters']));trial['parameters'].update(zip(keys,map(float,values)))
   return np.array([r['y'] for r in native(trial,x)])[:cut]-measured[:cut]
  opt=least_squares(residual,initial,bounds=(lo,hi),x_scale='jac',max_nfev=100);d['parameters'].update(zip(keys,map(float,opt.x)));fit=dict(success=bool(opt.success),message=opt.message,nfev=opt.nfev,training_samples=cut,withheld_samples=len(x)-cut,scope='Same ordered specimen temporal holdout; not independent validation')
 rows=native(d,x);y=np.array([r['y'] for r in rows]);metrics={}
 if measured is not None:
  if len(y)!=len(measured):raise ValueError('Protocol lengths differ')
  for r,v in zip(rows,measured):r['measured_y']=float(v);r['residual']=r['y']-float(v)
  cut=fit.get('training_samples',len(rows))
  for label,a,b in [('full',0,len(y)),('training',0,cut),('withheld_tail',cut,len(y))]:
   if b>a:
    e=y[a:b]-measured[a:b];metrics[label]=dict(RMSE=float(np.sqrt(np.mean(e**2))),MAE=float(np.mean(abs(e))),max_error=float(max(abs(e))))
 write(out/'response.csv',rows)
 from matplotlib.figure import Figure
 fig=Figure(figsize=(11,7),layout='constrained');a=fig.add_subplot(221);b=fig.add_subplot(222);c=fig.add_subplot(223);e=fig.add_subplot(224)
 a.plot([r['x'] for r in rows],y,label='Native response');b.plot(range(len(y)),y,label='Native response')
 if measured is not None:a.plot(x,measured,label='Measured');b.plot(range(len(y)),measured,label='Measured');c.plot(y-measured);e.scatter(measured,y,s=5)
 else:c.text(.1,.5,'No experimental data supplied');e.text(.1,.5,'Numerical parameter study only')
 unit='kN' if d['model'] in ('FVD','BRB','SFI_MVLEM_3D') else 'MPa';xu='Extension / displacement (m)' if unit=='kN' else 'Strain'
 a.set(xlabel=xu,ylabel=unit);b.set(xlabel='Ordered sample',ylabel=unit);c.set(title='Residual (native - measured)',ylabel=unit);e.set(xlabel='Measured '+unit,ylabel='Native '+unit)
 for ax in (a,b,c,e):ax.grid(alpha=.2)
 a.legend();b.legend();fig.suptitle(d['model']+' - specimen-specific study; no automatic project approval');fig.savefig(out/'response.png',dpi=150)
 from openpyxl import Workbook
 wb=Workbook();ws=wb.active;ws.title='Response';ws.append(list(rows[0]))
 for row in rows:ws.append(list(row.values()))
 if unit=='MPa':
  col=len(rows[0])+1;ws.cell(1,col,'Stress kN/m2')
  for i in range(2,len(rows)+2):ws.cell(i,col,'=C'+str(i)+'*1000')
 ps=wb.create_sheet('Parameters sources');ps.append(['Parameter','Value'])
 for k,v in d['parameters'].items():ps.append([k,v])
 for k in ('parameter_basis','specimen_id','citation','source'):ps.append([k,d.get(k,'')])
 if d.get('profile_id'):
  from components.validation.material_catalog import parameter_rows,commands,FORMULAS
  audit=wb.create_sheet('Exact provenance');audit.append(['Group','Parameter','Value','Unit','Basis','Source','Locator'])
  for r in parameter_rows(d):audit.append(list(r.values()))
  command_sheet=wb.create_sheet('Native commands');command_sheet.append(['kN-m native commands; MPa multiplied by 1000 once'])
  for line in commands(d).splitlines():command_sheet.append([line])
  command_sheet.append([FORMULAS[d['model']]])
  (out/'Native_Command_Trace.txt').write_text(commands(d),encoding='utf8')
 wb.save(out/'Parameter_Study.xlsx')
 report=dict(input=d,optimizer=fit,metrics=metrics,status='EXPERIMENTAL_COMPARISON' if measured is not None else 'NUMERICAL_PARAMETER_STUDY',elapsed_s=time.perf_counter()-start,finished_at=datetime.datetime.now().astimezone().isoformat(),project_transfer_approved=False)
 (out/'report.json').write_text(json.dumps(report,indent=2),encoding='utf8');print(json.dumps(report,indent=2));return report
if __name__=='__main__':
 parser=argparse.ArgumentParser();parser.add_argument('input');parser.add_argument('output');a=parser.parse_args()
 try:run(json.loads(Path(a.input).read_text(encoding='utf-8-sig')),a.output)
 except Exception as exc:
  out=Path(a.output);out.mkdir(parents=True,exist_ok=True);(out/'failure.json').write_text(json.dumps(dict(status='FAILED',error=str(exc),project_transfer_approved=False),indent=2));raise
