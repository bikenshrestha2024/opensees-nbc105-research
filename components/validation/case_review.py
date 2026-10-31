"""Compare saved responses and expose combination/preload coverage without inventing runs."""
from pathlib import Path
import json,csv,os

def collect(folder):
    from OpenSees_Design_Checks_V135 import scan_drift,combinations
    folder=Path(folder);rows,issues=scan_drift(folder/'OpenSees_Analysis_Results')
    profiles=[]
    for path in folder.rglob('Gravity_Profile_*_V142.json'):
        item=json.loads(path.read_text(encoding='utf-8-sig'));profiles.append(dict(file=str(path),limit_state=item.get('limit_state'),**item.get('gravity_context',{})))
    deflections=[]
    for path in (folder/'OpenSees_Analysis_Results').rglob('Floor_Edge_Drift_*.csv'):
        with path.open(encoding='utf-8-sig',newline='') as stream:
            for row in csv.DictReader(stream):
                if row.get('floor_displacement_m') not in (None,''):deflections.append(dict(file=str(path),case=path.stem,limit_state=row.get('limit_state',''),storey=row.get('story_number'),axis=row.get('axis',''),displacement_m=float(row['floor_displacement_m']),scope='Saved floor-edge response; verify source model and combination before comparison'))
    for path in (folder/'OpenSees_Analysis_Results').rglob('Deformation_*.json'):
        try:
            item=json.loads(path.read_text(encoding='utf-8-sig')); state=item.get('limit_state','')
            if state not in ('ULS','SLS'):continue
            levels={}
            for tag,xyz in item.get('nodes',{}).items():
                u=item.get('displacements',{}).get(str(tag),[])
                if len(u)<2:continue
                z=round(float(xyz[2]),7)
                for i,axis in enumerate(('X','Y')):levels[z,axis]=max(levels.get((z,axis),0.),abs(float(u[i])))
            for (z,axis),value in levels.items():deflections.append(dict(file=str(path),case=path.stem,limit_state=state,storey=z,axis=axis,displacement_m=value,scope='Maximum absolute saved nodal component at elevation (m); CQC envelopes are non-concurrent; no force-only scale applied'))
        except (ValueError,KeyError,TypeError) as exc:issues.append(dict(file=str(path),error=str(exc)))
    coverage=[]
    for case in combinations(nonparallel=True):
        for method in ('ESM','MRSM','LTH'):
            coverage.append(dict(case=case['case'],limit_state=case['limit_state'],method=method,factors=json.dumps(case['factors']),status='REQUIRES CASE-SPECIFIC RESULT / IDENTITY REVIEW',note='Existing similar state/direction files do not prove this full combination was solved; CQC envelopes cannot be treated as concurrent signed actions.'))
    return dict(drift=rows,deflection=deflections,issues=issues,preload_profiles=profiles,combination_coverage=coverage)


def export(folder):
    from openpyxl import Workbook
    folder=Path(folder);result=collect(folder);out=folder/'OpenSees_Analysis_Results/Case_Review_V142';out.mkdir(parents=True,exist_ok=True)
    wb=Workbook();wb.remove(wb.active)
    for key,title in [('drift','Saved drift'),('deflection','Saved deflection'),('preload_profiles','PDelta preload provenance'),('combination_coverage','Combination coverage'),('issues','Read errors')]:
        ws=wb.create_sheet(title);items=result[key]
        if items:
            keys=list(dict.fromkeys(k for r in items for k in r));ws.append(keys)
            for item in items:ws.append([json.dumps(item.get(k)) if isinstance(item.get(k),(dict,list)) else item.get(k) for k in keys])
            ws.auto_filter.ref=ws.dimensions;ws.freeze_panes='A2'
    note=wb.create_sheet('Read first');note.append(['Missing combination results are not PASS. Profiles select gravity preloads for future fresh model runs, with original seismic mass retained.'])
    note.append(['AS_CONFIGURED = source configured preload; GRAVITY_1.2D_1.5L; UPLIFT_0.9D. Snow/liquid/earth/wind/vertical earthquake applicability still requires project action data.'])
    note.append(['Linear time history is a linear transient model about a selected preloaded state. It is not a full nonlinear PDelta time-history solution. Positive CQC component envelopes are not concurrent deformed shapes.'])
    wb.save(out/'All_Case_Review_V142.xlsx');(out/'All_Case_Review_V142.json').write_text(json.dumps(result,indent=2),encoding='utf-8');return result,out


def launch(parent,folder):
    import tkinter as tk
    from tkinter import ttk,messagebox
    from matplotlib.figure import Figure
    from OpenSees_Research_UI_V123 import embed_figures
    win=tk.Toplevel(parent);win.title('V142 saved drift / deflection and combination coverage');win.geometry('1200x800');controls=ttk.Frame(win);controls.pack(fill='x')
    profile=tk.StringVar(value=os.environ.get('OPENSEES_GRAVITY_PROFILE','AS_CONFIGURED'));ttk.Label(controls,text='Gravity preload for subsequent analysis runs:').pack(side='left');ttk.Combobox(controls,textvariable=profile,values=['AS_CONFIGURED','GRAVITY_1.2D_1.5L','UPLIFT_0.9D'],state='readonly',width=28).pack(side='left')
    def apply():os.environ['OPENSEES_GRAVITY_PROFILE']=profile.get();messagebox.showinfo('Preload selected','Subsequent fresh analysis processes inherit this preload. Existing results are not recalculated. Mass factors remain unchanged.',parent=win)
    ttk.Button(controls,text='Apply to subsequent runs',command=apply).pack(side='left')
    text=tk.StringVar();ttk.Label(win,textvariable=text,wraplength=1150).pack(fill='x');body=ttk.Frame(win);body.pack(fill='both',expand=True)
    def refresh():
        try:
            result,out=export(folder);fig=Figure(figsize=(11,7),layout='constrained');axes=fig.subplots(2,2)
            for col,state in enumerate(('ULS','SLS')):
                grouped={}
                for row in result['drift']:
                    if row['limit_state']==state:grouped.setdefault(row['case']+' '+row.get('axis',''),[]).append(row)
                for case,items in grouped.items():
                    valid=[r for r in items if str(r['storey']).replace('.','',1).isdigit()]
                    valid.sort(key=lambda r:float(r['storey']))
                    if valid:axes[0,col].plot([r['demand'] for r in valid],[float(r['storey']) for r in valid],label=case)
                axes[0,col].axvline(.025 if state=='ULS' else .006,color='red',ls='--');axes[0,col].set(title=state+' saved drift',xlabel='Design drift ratio',ylabel='Storey')
                grouped={}
                for row in result['deflection']:
                    if row['limit_state']==state:grouped.setdefault(row['case']+' '+row.get('axis',''),[]).append(row)
                for case,items in grouped.items():
                    valid=[r for r in items if str(r['storey']).replace('.','',1).isdigit()];valid.sort(key=lambda r:float(r['storey']))
                    if valid:axes[1,col].plot([abs(r['displacement_m'])*1000 for r in valid],[float(r['storey']) for r in valid],label=case)
                axes[1,col].set(title=state+' saved deflection (case envelope)',xlabel='Absolute displacement (mm)',ylabel='Storey / elevation (see source workbook)')
            for ax in axes.flat:ax.grid(alpha=.25)
            embed_figures(body,[('ULS / SLS saved responses',fig)]);fig.savefig(out/'Drift_Deflection_Comparison_V142.png',dpi=150)
            text.set(str(len(result['drift']))+' saved drift rows; '+str(len(result['deflection']))+' saved deflection rows. Full combination coverage remains subject to explicit case/identity verification. Workbook: '+str(out/'All_Case_Review_V142.xlsx'))
        except Exception as exc:messagebox.showerror('Case review',str(exc),parent=win)
    ttk.Button(controls,text='Refresh / export comparison',command=refresh).pack(side='left');refresh();return win
