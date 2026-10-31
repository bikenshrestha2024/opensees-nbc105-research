"""Auditable combination coverage, serviceability and reviewed section checks.
No capacity is inferred from elastic stiffness or analysis convergence.
"""
from pathlib import Path
import csv,json,math,time,datetime,hashlib,itertools
SOURCE='https://giwmscdnone.gov.np/media/files/NBC_105_2025_Seismic%20Design%20of%20Buildings%20in%20Nepal_lqqvx6a.pdf'
COMPONENTS=('N_kN','Vy_kN','Vz_kN','T_kNm','My_kNm','Mz_kNm')
REQUIRED=('axial_biaxial_interaction','shear','torsion','stability','detailing_capacity_design')

def read_rows(path):
    with Path(path).open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))

def write_rows(path,rows):
    if not rows:return
    with Path(path).open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=list(dict.fromkeys(k for r in rows for k in r)));w.writeheader();w.writerows(rows)

def number(x):
    v=float(x)
    if not math.isfinite(v):raise ValueError('Non-finite numerical input')
    return v

def combinations(storage=False,nonparallel=False):
    lam=.6 if storage else .3;out=[]
    for state in ('ULS','SLS'):
        def add(name,factors,clause='3.6.1',note=''):
            out.append(dict(case=state+'_'+name,limit_state=state,factors=factors,clause='NBC 105:2025 '+clause,note=note))
        add('GRAVITY',{'DL':1.2,'LL':1.5})
        for action,factor in [('SNOW',1.2),('LIQUID',1.2),('EARTH',1.5)]:
            add(action,{'DL':1.2,'LL':.5,action:factor},note='Conditional action; confirm applicability, not silently omitted.')
        directions=[(1.,0.),(0.,1.)]
        if nonparallel:directions=[(1.,.3),(.3,1.)]
        for dx,dy in directions:
            for sx,sy in itertools.product((-1,1),repeat=2):
                if (not dx and sx<0) or (not dy and sy<0):continue
                label=f'EX{sx*dx:+g}_EY{sy*dy:+g}'
                for ecc in ('E0','EPLUS','EMINUS'):
                    for base,gravity in [('SEISMIC',{'DL':1.,'LL':lam}),('UPLIFT',{'DL':.9})]:
                        factors=dict(gravity)
                        if dx:factors[state+'_EX_'+ecc]=sx*dx
                        if dy:factors[state+'_EY_'+ecc]=sy*dy
                        add(base+'_'+label+'_'+ecc,factors,'3.6.2' if nonparallel else '3.6.1',
                            'E0 is a centered reference; eccentricity signs use separately solved actions. State-specific seismic action; gravity factors do not multiply ground motion.')
    return out

def combine_actions(rows,cases,model_hash):
    """Signed elastic action superposition only. Same end/station/axes required.
    A second-order/nonlinear total response must be solved as a combination,
    never linearly superposed here. CQC component maxima are not accepted.
    """
    basis={};issues=[]
    for r in rows:
        if r.get('model_sha256')!=model_hash:raise ValueError('Action model SHA256 does not match selected assembly')
        if r.get('response_kind')!='signed_first_order_linear_increment':
            raise ValueError('Combination basis must be signed_first_order_linear_increment; CQC envelopes and total/preloaded responses cannot be superposed.')
        key=(r['limit_state'],r['member'],r['station'],r['axes']);slot=(key,r['basis_case'])
        if slot in basis:raise ValueError('Duplicate action basis '+str(slot))
        if r['axes']!='local':raise ValueError('Only consistently oriented local member actions supported')
        basis[slot]=[number(r[k]) for k in COMPONENTS]
    keys=sorted(set(k for k,b in basis));out=[]
    for c in cases:
        for key in keys:
            if key[0]!=c['limit_state']:continue
            missing=[b for b in c['factors'] if (key,b) not in basis]
            if missing:
                issues.append(dict(case=c['case'],member=key[1],station=key[2],status='NOT EVALUATED',reason='Missing basis: '+', '.join(missing)));continue
            v=[sum(f*basis[key,b][i] for b,f in c['factors'].items()) for i in range(6)]
            out.append(dict(case=c['case'],limit_state=key[0],member=key[1],station=key[2],axes=key[3],**dict(zip(COMPONENTS,v)),scope='FIRST ORDER ACTION COMBINATION ONLY; second-order and capacity checks separate'))
    return out,issues

def evaluate_checks(rows,model_hash,kind):
    out=[]
    for r in rows:
        q=dict(r);reason='';ratio=None
        try:
            if r.get('model_sha256')!=model_hash:raise ValueError('Model fingerprint mismatch')
            if r.get('reviewed','').lower()!='true':raise ValueError('Evidence not reviewed')
            if not r.get('source') or not r.get('clause') or not r.get('result_file'):raise ValueError('Missing source, clause or result file')
            if r.get('limit_state') not in ('ULS','SLS'):raise ValueError('Missing ULS/SLS context')
            if not r.get('unit'):raise ValueError('Missing compatible demand/capacity unit')
            if not Path(r['result_file']).is_file():raise ValueError('Result evidence file does not exist')
            if kind=='nonlinear' and (not r.get('performance_objective') or r.get('control_type') not in ('force','deformation') or not r.get('suite_statistic')):raise ValueError('Nonlinear check requires performance objective, force/deformation control and suite statistic')
            demand=abs(number(r['demand']));limit=number(r['limit'])
            if limit<=0:raise ValueError('Limit must be positive')
            if kind in ('section','nonlinear') and not r.get('concurrent_action_id'):raise ValueError('Missing concurrent action/station/time or approved envelope-method reference')
            ratio=demand/limit;status='PASS' if ratio<=1 else 'FAIL'
        except (ValueError,KeyError,TypeError) as e:status='NOT EVALUATED';reason=str(e)
        q.update(utilization=ratio,status=status,reason=reason,scope='Reviewed supplied criterion only; not whole-section certification' if kind=='section' else 'Reviewed nonlinear performance criterion; not whole-building acceptance' if kind=='nonlinear' else 'Reviewed project deflection criterion')
        out.append(q)
    return out

def scan_drift(folder):
    out=[];issues=[]
    for p in sorted(Path(folder).rglob('*.csv')):
        if 'Design_Checks_V135' in p.parts:continue
        if not (p.name.startswith('Drift_Check_') or p.name.startswith('Floor_Edge_Drift_')):continue
        try:
            for r in read_rows(p):
                state=r.get('limit_state','').upper()
                if state not in ('ULS','SLS'):raise ValueError('Drift result lacks limit_state')
                d=abs(number(r['design_drift_ratio']));limit=.025 if state=='ULS' else .006
                out.append(dict(file=str(p),case=p.stem,limit_state=state,storey=r.get('story_number',r.get('story','')),axis=r.get('axis',''),edge=r.get('edge','master'),demand=d,limit=limit,utilization=d/limit,status='PASS' if d<=limit else 'FAIL',clause='NBC 105:2025 5.5.3',scope='Saved-result check; verify selected current model and completed case. Not a section strength result.'))
        except Exception as e:issues.append(dict(file=str(p),status='NOT EVALUATED',reason=str(e)))
    # MRSM stores story checks in result JSON, without separate drift CSV.
    for p in sorted(Path(folder).rglob('MRSM_Results_V132.json')):
        try:
            results=json.loads(p.read_text(encoding='utf-8-sig'))
            for case,r in results.items():
                state=r['limit_state'];limit={'ULS':.025,'SLS':.006}[state]
                for s in r['stories']:
                    d=abs(number(s['design_drift_ratio']))
                    out.append(dict(file=str(p),case=case,limit_state=state,storey=s.get('story_number',''),axis=r.get('axis',''),edge='CQC design drift',demand=d,limit=limit,utilization=d/limit,status='PASS' if d<=limit else 'FAIL',clause='NBC 105:2025 5.5.3',scope='Saved CQC drift; no force-only scale applied to displacement. Verify source model.'))
        except Exception as e:issues.append(dict(file=str(p),status='NOT EVALUATED',reason=str(e)))
    return out,issues

def templates(folder):
    folder=Path(folder);folder.mkdir(parents=True,exist_ok=True)
    definitions={
      'Nonlinear_Performance_V135.csv':['performance_objective','control_type','suite_statistic','limit_state','case','member','check','demand','limit','unit','concurrent_action_id','model_sha256','reviewed','source','clause','result_file'],
      'Section_Criteria_V135.csv':['limit_state','case','member','check','demand','limit','unit','concurrent_action_id','model_sha256','reviewed','source','clause','result_file'],
      'Deflection_Criteria_V135.csv':['limit_state','case','member','check','demand','limit','unit','model_sha256','reviewed','source','clause','result_file'],
      'Action_Basis_V135.csv':['limit_state','basis_case','member','station','axes','response_kind','model_sha256',*COMPONENTS]}
    for name,fields in definitions.items():
        path=folder/name
        if not path.exists():
            with path.open('w',encoding='utf-8-sig',newline='') as f:csv.writer(f).writerow(fields)
    return folder

def run(folder,storage=False,nonparallel=False):
    start=time.perf_counter();started=datetime.datetime.now().astimezone().isoformat();folder=Path(folder)
    model=folder/'OpenSees_Complete_Model_Data_V118.json'
    if not model.is_file():raise ValueError('Select export folder containing OpenSees_Complete_Model_Data_V118.json')
    digest=hashlib.sha256(model.read_bytes()).hexdigest();data=json.loads(model.read_text(encoding='utf-8-sig'))
    output=templates(folder/'OpenSees_Analysis_Results'/'Design_Checks_V135');cases=combinations(storage,nonparallel)
    drift,issues=scan_drift(folder/'OpenSees_Analysis_Results')
    criteria=evaluate_checks(read_rows(output/'Section_Criteria_V135.csv'),digest,'section')
    nonlinear=evaluate_checks(read_rows(output/'Nonlinear_Performance_V135.csv'),digest,'nonlinear')
    deflections=evaluate_checks(read_rows(output/'Deflection_Criteria_V135.csv'),digest,'deflection')
    from components.validation.response_inventory import available_response_rows
    recorded_deflections,recorded_actions=available_response_rows(folder)
    deflections.extend(recorded_deflections)
    if not criteria:
        criteria=recorded_actions or [dict(status='CAPACITY REQUIRED',reason='No section capacities or recorded member actions found. Run analysis and supply actual reinforcement/design criteria.')]
    combined,missing=combine_actions(read_rows(output/'Action_Basis_V135.csv'),cases,digest)
    inventory=[]
    for kind,items in [('FRAME',data.get('frames',[])),('AREA',data.get('shells',[]))]:
        for item in items:
            member=kind+':'+str(item['object'])
            if any(r['member']==member for r in inventory):continue
            found=[r for r in criteria if r.get('member')==member]
            failed=any(r['status']=='FAIL' for r in found)
            # Per-member summary never promotes a partial list to whole-section PASS.
            inventory.append(dict(member=member,category=item.get('category',''),section=item.get('properties',{}).get('section',item.get('properties',{}).get('property','')),status='FAIL' if failed else 'NOT EVALUATED',checks_supplied=len(found),required_checks=', '.join(REQUIRED),reason='A supplied criterion fails' if failed else 'Whole-section design requires all governing combinations, stations, interactions, stability and detailing evidence; individual checks listed separately.'))
    coverage=[]
    for c in cases:
        n=sum(r['case']==c['case'] for r in combined)
        coverage.append(dict(case=c['case'],limit_state=c['limit_state'],expression=' + '.join(f'{f:g} {b}' for b,f in c['factors'].items()),clause=c['clause'],action_rows=n,status='PARTIAL FIRST ORDER ACTIONS' if n else 'NOT EVALUATED',note=c['note']))
    for name,note in [('Vertical earthquake','NBC 3.6.3 / 4.3 applicability and signed vertical actions must be reviewed'),('Wind','Applicable wind code and combinations require separate project basis'),('Long-term vertical deflection','Creep, shrinkage, construction stages, finishes/partitions and support-chord reference required'),('Nonlinear component acceptance','Force-controlled brittle actions and deformation-controlled ductile actions require calibrated acceptance limits, suite coverage and performance objective; no automatic IO/LS/CP thresholds'),('Linear reinforcement demand','Requires approved RC material strengths, actual bar layout, cover, transverse steel, concurrent axial/biaxial actions and material-code design engine; illustrative fiber templates are insufficient'),('ULS stability and capacity design','Buckling, P-Delta, joints, anchorage, ductile detailing and capacity-design checks are separate from drift')]:
        issues.append(dict(check=name,status='NOT EVALUATED',reason=note))
    if not drift:issues.append(dict(check='Seismic drift',status='NOT EVALUATED',reason='No saved drift results found; run ULS and SLS analysis cases'))
    for state in ('ULS','SLS'):
        if not any(r['limit_state']==state for r in drift):issues.append(dict(check=state+' drift coverage',status='NOT EVALUATED',reason='No saved results in this limit state'))
    from components.validation.settings_audit import audit
    settings=audit(folder)
    issues.extend(dict(check=r['item'],status=r['status'],reason=r['basis']) for r in settings['rows'] if r['status'] not in ('CODE LIMIT REFERENCE','IMPLEMENTATION REFERENCE'))
    tables={'Combination coverage':coverage,'Seismic drift':drift,'Project deflection':deflections,'Section checks':criteria,'Nonlinear performance':nonlinear,'Section inventory':inventory,'Combined actions':combined,'Missing action basis':missing,'Open items':issues}
    for name,rows in tables.items():write_rows(output/(name.replace(' ','_')+'.csv'),rows)
    elapsed=time.perf_counter()-start
    report=dict(model_sha256=digest,started=started,finished=datetime.datetime.now().astimezone().isoformat(),elapsed_s=elapsed,overall='INCOMPLETE' if not any(r.get('status')=='FAIL' for rows in tables.values() for r in rows) else 'FAIL / INCOMPLETE',sources=[SOURCE],tables=tables,scope='No whole-section PASS is inferred. Existing analysis outputs are checked without rerunning the solver. MRSM force envelopes and nonlinear totals are never superposed as linear increments.')
    (output/'Design_Check_Report_V135.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
    from openpyxl import Workbook
    wb=Workbook();wb.remove(wb.active)
    for name,rows in tables.items():
        ws=wb.create_sheet(name[:31]);fields=list(dict.fromkeys(k for r in rows for k in r))
        if fields:
            ws.append(fields)
            for r in rows:ws.append([r.get(k) for k in fields])
            ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
            for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=min(55,max(14,max(len(str(c.value or '')) for c in list(col)[:100])+2))
        else:ws.append(['NOT EVALUATED - no supplied/current results'])
    ws=wb.create_sheet('Basis and timing');ws.append(['Source',SOURCE]);ws.append(['Model SHA256',digest]);ws.append(['Started',started]);ws.append(['Elapsed before workbook (s)',elapsed]);ws.append(['Scope',report['scope']])
    wb.save(output/'Design_Check_Report_V135.xlsx')
    elapsed=time.perf_counter()-start;report['elapsed_s']=elapsed
    report['finished']=datetime.datetime.now().astimezone().isoformat()
    (output/'Design_Check_Report_V135.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
    (output/'Design_Check_V135.log').write_text(f"{started}\nModel SHA256: {digest}\n{report['overall']}\nElapsed including workbook: {elapsed:.3f} s\nFinished: {report['finished']}\n",encoding='utf-8')
    return report,output

def launch(parent,folder):
    import tkinter as tk
    from tkinter import ttk,messagebox
    import os
    win=tk.Toplevel(parent);win.title('ULS / SLS combination, serviceability and section checks - V135');win.geometry('1250x780')
    path=tk.StringVar(value=str(folder));storage=tk.BooleanVar();np=tk.BooleanVar();summary=tk.StringVar(value='A section passes only after its governing strength, stability and detailing checks are completed. Drift PASS is separate.')
    ttk.Entry(win,textvariable=path).pack(fill='x',padx=8,pady=5)
    bar=ttk.Frame(win);bar.pack(fill='x');ttk.Checkbutton(bar,text='Storage occupancy (lambda 0.6; otherwise 0.3)',variable=storage).pack(side='left');ttk.Checkbutton(bar,text='Nonparallel lateral systems (100% + 30%)',variable=np).pack(side='left')
    ttk.Label(win,textvariable=summary,wraplength=1200).pack(fill='x',padx=8,pady=8)
    book=ttk.Notebook(win);book.pack(fill='both',expand=True);state={'output':None}
    def refresh():
        try:
            report,output=run(path.get(),storage.get(),np.get());state['output']=output
            for tab in book.tabs():book.forget(tab)
            for name,rows in report['tables'].items():
                if not rows:rows=[dict(status='NOT EVALUATED',reason='No supplied criteria or matching saved results. Open report/input templates for required inputs.')]
                frame=ttk.Frame(book);book.add(frame,text=name);fields=list(dict.fromkeys(k for r in rows for k in r)) or ['status']
                tree=ttk.Treeview(frame,columns=fields,show='headings');xs=ttk.Scrollbar(frame,orient='horizontal',command=tree.xview);ys=ttk.Scrollbar(frame,command=tree.yview);tree.configure(xscrollcommand=xs.set,yscrollcommand=ys.set)
                tree.grid(row=0,column=0,sticky='nsew');ys.grid(row=0,column=1,sticky='ns');xs.grid(row=1,column=0,sticky='ew');frame.rowconfigure(0,weight=1);frame.columnconfigure(0,weight=1)
                for f in fields:tree.heading(f,text=f);tree.column(f,width=180,stretch=False)
                for r in rows:tree.insert('', 'end',values=[r.get(f,'') for f in fields])
            summary.set(f"{report['overall']} | {len(report['tables']['Seismic drift'])} saved drift checks | {len(report['tables']['Section checks'])} section/action review rows | elapsed {report['elapsed_s']:.2f} s | Model SHA: {report['model_sha256'][:16]}")
        except Exception as e:messagebox.showerror('Check inputs',str(e),parent=win)
    def open_output():
        if state['output']:os.startfile(str(state['output']))
    ttk.Button(bar,text='Refresh checks + export Excel',command=refresh).pack(side='left',padx=8);ttk.Button(bar,text='Open report / input templates',command=open_output).pack(side='left')
    win.after(100,refresh);return win
