# -*- coding: utf-8 -*-
"""NBC 105:2025 equivalent-static seismic weight and floor force calculator.

Input: ETABS converter story-weight CSV or verified user CSV. Geometry/master-node
metadata are linked to the ETABS export folder. Modal periods can be read from
the active ETABS API session after running a modal case.
The dead/live totals must be independently verified against the project model; this GUI never guesses or extracts them.
Outputs: calculation JSON, storey force CSV, and an OpenSeesPy lateral-load
runner. The runner needs a separately verified structural model already built.
"""
from __future__ import annotations

__release__ = '134'
from OpenSees_Timing_V130 import timed_outputs
import csv, json, math, os, sys, time, traceback, subprocess, threading
from pathlib import Path
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

HERE=Path(__file__).resolve().parent
try:
    from NBC105_Analysis_Settings_V109 import (
        NBC105_ESM_SYSTEM_FACTORS, calculate_esm,
    )
except Exception as exc:
    raise RuntimeError("Keep NBC105_Analysis_Settings_V109.py beside this GUI.") from exc


def read_story_csv(path):
    rows=[]
    with open(path, newline="", encoding="utf-8-sig") as f:
        for line, r in enumerate(csv.DictReader(f),2):
            if not str(r.get("story_number","")).strip():
                continue
            for key in ("height_m","dead_load_kN","storage_live_load_kN","other_live_load_kN"):
                value=str(r.get(key,"")).strip()
                if not value:
                    raise ValueError(f"CSV row {line}: {key} is blank; enter verified ETABS data.")
                try: r[key]=float(value)
                except ValueError as e: raise ValueError(f"CSV row {line}: {key} must be numeric.") from e
            tag=str(r.get("master_node_tag","")).strip()
            r["master_node_tag"]=int(tag) if tag else None
            if not r["master_node_tag"]:
                raise ValueError(f"CSV row {line}: no unique rigid-diaphragm master node tag; resolve the mapping before force export.")
            r["story_number"]=int(r["story_number"])
            r["story_name"]=str(r.get("story_name",r["story_number"]))
            # Table 5-1 makes roof live load nil in seismic weight regardless
            # of the entered audit field. Preserve it only as a visible note.
            roof=str(r.get("roof_live_load_kN_ignored","")).strip()
            r["roof_live_load_kN_ignored"]=float(roof) if roof else 0.0
            rows.append(r)
    if not rows: raise ValueError("No storey rows were found in the selected CSV.")
    tags=[r["master_node_tag"] for r in rows]
    if len(set(tags)) != len(tags): raise ValueError("Each storey must have a unique diaphragm master-node tag.")
    return rows


def safe_write(path, text):
    Path(path).write_text(text,encoding="utf-8",newline="")


@timed_outputs('Equivalent-static calculations and case export','outdir')
def write_outputs(template_path, rows, result, direction, polarity, accidental, breadth, outdir, model_script=""):
    out=Path(outdir); out.mkdir(parents=True,exist_ok=True)
    ecc_case=(f"E{'+' if int(breadth[1])>0 else '-'}" if accidental and int(breadth[1])!=0 else "E0")
    suffix=f"{result['limit_state']}_{direction}{'+' if polarity>0 else '-'}_{ecc_case}"
    result["response_case"]=suffix
    report=dict(result)
    report["comparison_remark"] = result["esm_eligibility"] if not result.get("code_eligible",True) else "Eligible ESM case"
    report["direction_case"]=suffix
    report["accidental_torsion_included"]=bool(accidental and int(breadth[1])!=0)
    report["accidental_eccentricity_sign"]=int(breadth[1]) if accidental else 0
    report["accidental_eccentricity_m"]=0.05*float(breadth[0])*abs(int(breadth[1])) if accidental else 0.0
    report["note"]="Story gravity totals are user-verified ETABS inputs. Site subsoil status is reported separately; this calculation is not an independent code-compliance certification."
    out_json=out/f"NBC105_ESM_Calculation_{suffix}_V111.json"
    safe_write(out_json,json.dumps(report,indent=2,ensure_ascii=False))
    out_csv=out/f"NBC105_ESM_Storey_Forces_{suffix}_V111.csv"
    out_png=out/f"NBC105_ESM_Force_Profile_{suffix}_V111.png"
    out_xlsx=out/f"NBC105_ESM_Calculation_{suffix}_V111.xlsx"
    fields=["limit_state","location","z_source","soil_type","Vs30_m_s","soil_type_from_Vs30","N_SPT_blow_30cm","soil_type_from_N_SPT","site_subsoil_status","Z","story_number","story_name","master_node_tag","height_m","dead_load_kN","storage_live_load_kN","storage_factor","other_live_load_kN","other_live_factor","roof_live_load_kN_ignored","seismic_weight_kN","k","floor_force_kN","storey_shear_kN","lateral_force_kN","torsional_moment_kNm"]
    with open(out_csv,"w",newline="",encoding="utf-8-sig") as f:
        w=csv.DictWriter(f,fieldnames=fields); w.writeheader()
        for story_index,r in enumerate(result["stories"]):
            fx=r["lateral_force_kN"]*polarity
            e=0.05*float(breadth[0])*int(breadth[1]) if accidental else 0.0
            mz=(-fx*e if direction=="X" else fx*e) if accidental else 0.0
            w.writerow({"limit_state":result["limit_state"],"location":result.get("location_name",""),"z_source":result.get("zoning_factor_source",""),"soil_type":result["soil_type"],"Vs30_m_s":result.get("vs30_m_s"),"soil_type_from_Vs30":result.get("soil_type_from_vs30"),"N_SPT_blow_30cm":result.get("n_spt_blow_30cm"),"soil_type_from_N_SPT":result.get("soil_type_from_n_spt"),"site_subsoil_status":result["site_subsoil_status"],"Z":result["zoning_factor_Z"],"story_number":r["story_number"],"story_name":r["story_name"],"master_node_tag":r["master_node_tag"],"height_m":r["height_m"],"dead_load_kN":r["dead_load_kN"],"storage_live_load_kN":r["storage_live_load_kN"],"storage_factor":0.60,"other_live_load_kN":r["other_live_load_kN"],"other_live_factor":0.30,"roof_live_load_kN_ignored":r["roof_live_load_kN_ignored"],"seismic_weight_kN":r["seismic_weight_kN"],"k":r["distribution_exponent_k"],"floor_force_kN":r["lateral_force_kN"],"storey_shear_kN":sum(x["lateral_force_kN"] for x in result["stories"][story_index:]),"lateral_force_kN":fx,"torsional_moment_kNm":mz})
    # Immediate visualization and workbook for the NBC calculation, before
    # launching the OpenSees structural response runner.
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        floors=result["stories"]
        fig,ax=plt.subplots(1,2,figsize=(10,6),constrained_layout=True)
        ax[0].barh([r["height_m"] for r in floors],[r["lateral_force_kN"] for r in floors],height=max(0.15,0.35*min(r["height_m"] for r in floors)))
        ax[0].set(title="NBC 105:2025 floor forces",xlabel="Fi (kN)",ylabel="Height (m)"); ax[0].grid(axis="x",alpha=.25)
        ax[1].step([sum(x["lateral_force_kN"] for x in floors[i:]) for i in range(len(floors))],[r["height_m"] for r in floors],where="mid")
        ax[1].set(title="Storey shear",xlabel="Vi (kN)",ylabel="Height (m)"); ax[1].grid(alpha=.25)
        fig.suptitle(f"{'COMPARISON ONLY | ' if not result.get('code_eligible',True) else ''}{result['limit_state']} {direction} | Z={result['zoning_factor_Z']:g} | soil {result['soil_type']} | V={result['base_shear_kN']:.3f} kN")
        fig.savefig(out_png,dpi=180); plt.close(fig)
    except ImportError:
        report.setdefault("optional_outputs",[]).append("Matplotlib unavailable: force-profile PNG not generated.")
    try:
        from openpyxl import Workbook
        wb=Workbook(); ws=wb.active; ws.title="NBC ESM Summary"
        for row in (("Item","Value"),("Standard",result["standard"]),("Limit state",result["limit_state"]),("Location",result.get("location_name","")),("Z source",result.get("zoning_factor_source","")),("Z",result["zoning_factor_Z"]),("Soil type",result["soil_type"]),("Vs,30 (m/s)",result.get("vs30_m_s")),("Soil type from Vs,30",result.get("soil_type_from_vs30")),("N-SPT (blows/30 cm)",result.get("n_spt_blow_30cm")),("Soil type from N-SPT",result.get("soil_type_from_n_spt")),("Site subsoil status",result["site_subsoil_status"]),("T1 (s)",result["period_s"]),("Ch(T1)",result["spectral_shape_Ch"]),("Cd(T1)",result["design_coefficient_Cd"]),("Seismic weight W (kN)",result["total_seismic_weight_kN"]),("Base shear V (kN)",result["base_shear_kN"]),("Exponent k",result["period_exponent_k"]),("ESM eligibility",result["esm_eligibility"])): ws.append(row)
        wf=wb.create_sheet("Storey Forces")
        with open(out_csv,newline="",encoding="utf-8-sig") as f:
            for row in csv.reader(f): wf.append(row)
        wb.save(out_xlsx)
    except ImportError:
        report.setdefault("optional_outputs",[]).append("openpyxl unavailable: calculation workbook not generated.")
        safe_write(out_json,json.dumps(report,indent=2,ensure_ascii=False))
    load_rows=[]
    e=0.05*float(breadth[0])*int(breadth[1]) if accidental else 0.0
    for r in result["stories"]:
        horizontal=polarity*r["lateral_force_kN"]
        fx,fy=(horizontal,0.0) if direction=="X" else (0.0,horizontal)
        mz=(-fx*e if direction=="X" else fy*e) if accidental else 0.0
        load_rows.append({"story_number":r["story_number"],"story_name":r["story_name"],"height_m":r["height_m"],"master_node_tag":int(r["master_node_tag"]),"Fx_kN":fx,"Fy_kN":fy,"Mz_kNm":mz})
    # This script assumes the verified structural domain is already in OpenSees,
    # and gravity has already been applied/committed if required by the model.
    script=[
        "# NBC 105:2025 equivalent-static lateral load case (generated by V111 earthquake-analysis GUI).",
        "# NBC 105:2025 source clauses: §§3.2.1, 4.1-4.2, 5.1-5.2, 5.5.1, 5.5.3, 5.6, 6.1-6.5; Tables 4-1, 5-1, 5-2, 6-1; Annex C/Table C-1.",
        "# REQUIRED: build the validated structural model first, with matching diaphragm master tags.",
        "# Freeze the verified gravity stage before running this lateral case. This file does not create slabs/walls/elements/materials.",
        "import json, openseespy.opensees as ops",
        "import os, csv, time, traceback, sys",
        "OUTPUT_DIR = "+repr(str(out.resolve())),
        "sys.path.insert(0,os.path.join(os.path.dirname("+repr(str(model_script))+"),'OpenSees_Runtime_V118'))",
        "sys.path.insert(0,"+repr(str(HERE))+")",
        "from OpenSees_Timing_V130 import StageTimer",
        "import atexit",
        "_processing_timer=StageTimer(OUTPUT_DIR,'Equivalent-static solver and response export')",
        "atexit.register(_processing_timer.finish)",
        "os.makedirs(OUTPUT_DIR, exist_ok=True)",
        "LOG_PATH = os.path.join(OUTPUT_DIR, 'NBC105_ESM_OpenSees_Run_'+"+repr(suffix)+"+'.log')",
        "_log_file = open(LOG_PATH,'w',encoding='utf-8')",
        "class _Tee:",
        "    def __init__(self,*streams): self.streams=streams",
        "    def write(self,text):",
        "        for s in self.streams: s.write(text); s.flush()",
        "    def flush(self): [s.flush() for s in self.streams]",
        "sys.stdout = _Tee(sys.__stdout__,_log_file); sys.stderr = _Tee(sys.__stderr__,_log_file)",
        "MODEL_BUILDER_SCRIPT = "+repr(str(Path(model_script).resolve()) if model_script else ""),
        "if not MODEL_BUILDER_SCRIPT: raise RuntimeError('Select a complete OpenSeesPy model builder in the earthquake-analysis GUI first.')",
        "if not os.path.isfile(MODEL_BUILDER_SCRIPT): raise FileNotFoundError(MODEL_BUILDER_SCRIPT)",
        "import runpy",
        "sys.path.insert(0, "+repr(str(HERE))+ ")",
        "sys.path.insert(0, os.path.join(os.path.dirname(MODEL_BUILDER_SCRIPT), 'OpenSees_Runtime_V118'))",
        "from OpenSees_Step13_Analysis_Basis_V115 import validate_initial_pdelta_results",
        "print('Building OpenSees model from:', MODEL_BUILDER_SCRIPT)",
        "MODEL_WORKDIR=os.path.dirname(MODEL_BUILDER_SCRIPT)",
        "os.chdir(MODEL_WORKDIR)",
        "from OpenSees_Initial_PDelta_V124 import ensure_initial_checks",
        "ensure_initial_checks(MODEL_WORKDIR)",
        "os.environ['OPENSEES_LIMIT_STATE'] = "+repr(result["limit_state"]),
        "_MODEL_NAMESPACE = runpy.run_path(MODEL_BUILDER_SCRIPT, run_name='__main__')",
        "IS16700_PDELTA_AUDIT = validate_initial_pdelta_results(_MODEL_NAMESPACE.get('IS16700_INITIAL_PDELTA_RESULTS'), lateral_actions=('EL',), directions=('X','Y'))",
        "print('Verified IS 16700:2023 §7.2(d) initial P-Delta evidence:', IS16700_PDELTA_AUDIT)",
        "os.chdir(OUTPUT_DIR)",
        "available_nodes=set(ops.getNodeTags())",
        "missing_masters=[row['master_node_tag'] for row in "+repr(load_rows)+" if row['master_node_tag'] not in available_nodes]",
        "if missing_masters: raise RuntimeError('Diaphragm master node(s) missing from model: '+str(missing_masters))",
        "ops.loadConst('-time', 0.0)  # after the required initial P-Delta stage, freeze the reviewed gravity state before this NBC ESM load case",
        "ESM_CASE = "+repr(suffix),
        "ESM_DIRECTION = "+repr(direction),
        "ESM_LOADS = "+repr(load_rows),
        "_run_started = time.perf_counter()",
        "ESM_LOAD_PATTERN_TAG = 98001",
        "ESM_TIME_SERIES_TAG = 98001",
        "# Model construction and gravity freeze are completed above, before creating this seismic pattern.",
        "DOF = 1 if ESM_DIRECTION == 'X' else 2",
        "MODEL_DATA_PATH=os.path.join(os.path.dirname(MODEL_BUILDER_SCRIPT),'OpenSees_Complete_Model_Data_V118.json')",
        "if not os.path.isfile(MODEL_DATA_PATH): raise FileNotFoundError('ESM reaction recording requires '+MODEL_DATA_PATH)",
        "with open(MODEL_DATA_PATH,encoding='utf-8-sig') as _stream: MODEL_DATA=json.load(_stream)",
        "ESM_SUPPORT_NODES=[int(n) for n,v in MODEL_DATA['supports'].items() if any(bool(x) for x in v[:3])]",
        "if not ESM_SUPPORT_NODES: raise ValueError('No support nodes found for ESM reaction recording')",
        "ops.reactions()",
        "BASE_REACTION_BEFORE_KN = sum(ops.nodeReaction(n, DOF) for n in ESM_SUPPORT_NODES)",
        "BASE_REACTIONS_BEFORE = {n:[float(ops.nodeReaction(n,j)) for j in range(1,7)] for n in ESM_SUPPORT_NODES}",
        "ALL_GRAVITY_U={n:list(ops.nodeDisp(n)) for n in ops.getNodeTags()}",
        "GRAVITY_BASELINE_U = {row['master_node_tag']: ops.nodeDisp(row['master_node_tag'], DOF) for row in ESM_LOADS}",
        "ops.timeSeries('Linear', ESM_TIME_SERIES_TAG)",
        "ops.pattern('Plain', ESM_LOAD_PATTERN_TAG, ESM_TIME_SERIES_TAG)",
        "for row in ESM_LOADS:",
        "    ops.load(row['master_node_tag'], row['Fx_kN'], row['Fy_kN'], 0.0, 0.0, 0.0, row['Mz_kNm'])",
        "ops.wipeAnalysis()",
        "ops.constraints('Transformation')  # required for rigidDiaphragm MP constraints",
        "ops.numberer('RCM')",
        "ops.system('UmfPack')",
        "ops.test('NormDispIncr', 1.0e-8, 50, 0)",
        "ops.algorithm('Newton')",
        "ops.integrator('LoadControl', 0.1)",
        "ops.analysis('Static')",
        "ESM_BASE_CURVE=[]",
        "def _record_esm_state():",
        "    ops.reactions()",
        "    reaction=sum(float(ops.nodeReaction(n,DOF)) for n in ESM_SUPPORT_NODES)-BASE_REACTION_BEFORE_KN",
        "    roof=max(ESM_LOADS,key=lambda x:float(x['height_m']))",
        "    ESM_BASE_CURVE.append({'load_factor':float(ops.getLoadFactor(ESM_LOAD_PATTERN_TAG)), 'time_s':float(ops.getTime()), 'base_shear_kN':-reaction, 'roof_displacement_m':float(ops.nodeDisp(roof['master_node_tag'],DOF)-GRAVITY_BASELINE_U[roof['master_node_tag']])})",
        "ok=0",
        "for _step in range(10):",
        "    ok=ops.analyze(1)",
        "    if ok != 0: break",
        "    _record_esm_state()",
        "if ok != 0:",
        "    ops.wipeAnalysis(); ops.constraints('Transformation'); ops.numberer('RCM'); ops.system('UmfPack')",
        "    ops.test('NormDispIncr', 1.0e-7, 100, 0); ops.algorithm('NewtonLineSearch')",
        "    ops.analysis('Static')",
        "    # Earlier successful steps remain committed. Advance only to factor 1.",
        "    ok = 0",
        "    for _recovery_step in range(21):",
        "        remaining = 1.0 - float(ops.getLoadFactor(ESM_LOAD_PATTERN_TAG))",
        "        if remaining <= 1.0e-10: break",
        "        ops.integrator('LoadControl', min(0.05, remaining))",
        "        ok = ops.analyze(1)",
        "        if ok != 0: break",
        "        _record_esm_state()",
        "if ok != 0:",
        "    raise RuntimeError(f'NBC ESM static analysis did not converge; OpenSees return code={ok}. Review model, sections, gravity state, constraints, and solver log.')",
        "if abs(float(ops.getLoadFactor(ESM_LOAD_PATTERN_TAG)) - 1.0) > 1.0e-8:",
        "    raise RuntimeError('ESM did not reach the intended load factor 1.0; response output suppressed.')",
        "ops.reactions()",
        "ESM_RESPONSE = []",
        "KD = "+repr(result["deflection_scale_factor_kd"]),
        "R_MU = "+repr(result["ductility_factor_R_mu"]),
        "LIMIT_STATE = "+repr(result["limit_state"]),
        "DRIFT_LIMIT = "+repr(result["drift_limit_ratio"]),
        "print('ESM REMARK:',"+repr(result["esm_eligibility"])+")",
        "PREV_U = 0.0; PREV_H = 0.0",
        "for row in ESM_LOADS:",
        "    n = row['master_node_tag']; u_raw=ops.nodeDisp(n,DOF)-GRAVITY_BASELINE_U[n]",
        "    displacement_factor = (R_MU if LIMIT_STATE == 'ULS' else 1.0) * KD",
        "    u_design = u_raw * displacement_factor",
        "    story_h = row['height_m'] - PREV_H",
        "    elastic_drift = abs(u_raw-(PREV_U/((R_MU if LIMIT_STATE == 'ULS' else 1.0)*KD)))/story_h if story_h > 0 else float('nan')",
        "    drift_ratio = abs(u_design-PREV_U)/story_h if story_h > 0 else float('nan')",
        "    story_shear = sum(x['Fx_kN'] if ESM_DIRECTION == 'X' else x['Fy_kN'] for x in ESM_LOADS if x['height_m'] >= row['height_m'])",
        "    drift_m = abs(u_raw-(PREV_U/((R_MU if LIMIT_STATE == 'ULS' else 1.0)*KD)))",
        "    story_stiffness_kN_per_m = abs(story_shear)/drift_m if drift_m > 1.0e-15 else float('inf')",
        "    ESM_RESPONSE.append({**row, 'response_case':"+repr(suffix)+", 'U_elastic_m':u_raw, 'U_design_m':u_design, 'story_height_m':story_h, 'elastic_drift_ratio':elastic_drift, 'design_drift_ratio':drift_ratio, 'story_shear_kN':story_shear, 'story_stiffness_kN_per_m':story_stiffness_kN_per_m, 'drift_limit':DRIFT_LIMIT, 'drift_screen_pass':(drift_ratio <= DRIFT_LIMIT)})",
        "    PREV_U=u_design; PREV_H=row['height_m']",
        "BASE_REACTION_KN = sum(ops.nodeReaction(n, DOF) for n in ESM_SUPPORT_NODES) - BASE_REACTION_BEFORE_KN",
        "APPLIED_BASE_SHEAR_KN = sum(row['Fx_kN'] if ESM_DIRECTION == 'X' else row['Fy_kN'] for row in ESM_LOADS)",
        "BASE_SHEAR_ERROR_PCT = abs(abs(BASE_REACTION_KN)-abs(APPLIED_BASE_SHEAR_KN))/max(abs(APPLIED_BASE_SHEAR_KN),1.0e-12)*100.0",
        "with open(os.path.join(OUTPUT_DIR,'NBC105_ESM_OpenSees_Response_'+ESM_CASE+'.json'),'w',encoding='utf-8') as f: json.dump(ESM_RESPONSE,f,indent=2)",
        "with open(os.path.join(OUTPUT_DIR,'NBC105_ESM_OpenSees_Response_'+ESM_CASE+'.csv'),'w',newline='',encoding='utf-8-sig') as f:",
        "    import csv; w=csv.DictWriter(f,fieldnames=list(ESM_RESPONSE[0]) if ESM_RESPONSE else []); w.writeheader(); w.writerows(ESM_RESPONSE)",
        "ESM_BASE_CURVE_PATH=os.path.join(OUTPUT_DIR,'NBC105_ESM_BaseShear_RoofDisplacement_'+ESM_CASE+'.csv')",
        "with open(ESM_BASE_CURVE_PATH,'w',newline='',encoding='utf-8-sig') as f:",
        "    import csv; w=csv.DictWriter(f,fieldnames=['load_factor','time_s','base_shear_kN','roof_displacement_m']); w.writeheader(); w.writerows(ESM_BASE_CURVE)",
        "SUPPORT_REACTION_ROWS=[]",
        "for n in ESM_SUPPORT_NODES:",
        "    values=[float(ops.nodeReaction(n,j))-BASE_REACTIONS_BEFORE[n][j-1] for j in range(1,7)]",
        "    SUPPORT_REACTION_ROWS.append(dict(node_tag=n,Fx_kN=values[0],Fy_kN=values[1],Fz_kN=values[2],Mx_kNm=values[3],My_kNm=values[4],Mz_kNm=values[5]))",
        "with open(os.path.join(OUTPUT_DIR,'NBC105_ESM_Support_Reactions_'+ESM_CASE+'.csv'),'w',newline='',encoding='utf-8-sig') as f:",
        "    import csv; w=csv.DictWriter(f,fieldnames=list(SUPPORT_REACTION_ROWS[0]) if SUPPORT_REACTION_ROWS else ['node_tag']); w.writeheader(); w.writerows(SUPPORT_REACTION_ROWS)",
        "with open(os.path.join(OUTPUT_DIR,'Response_Recorder_Index_'+ESM_CASE+'.json'),'w',encoding='utf-8') as f: json.dump({'analysis':'equivalent_static','files':[{'name':os.path.basename(ESM_BASE_CURVE_PATH),'channels':['accepted load factor','base shear','roof displacement']},{'name':'NBC105_ESM_OpenSees_Response_'+ESM_CASE+'.csv','channels':['story displacement','story drift','story shear','secant story stiffness']},{'name':'NBC105_ESM_Support_Reactions_'+ESM_CASE+'.csv','channels':['incremental support forces and moments']}],'notes':['The base-shear/roof-displacement curve is the converged linear static load-control path, not a nonlinear pushover capacity curve.']},f,indent=2)",
        "try:",
        "    import matplotlib; matplotlib.use('Agg'); import matplotlib.pyplot as plt",
        "    fig, ax=plt.subplots(1,5,figsize=(20,6),constrained_layout=True); h=[r['height_m'] for r in ESM_RESPONSE]",
        "    ax[0].barh(h,[r['Fx_kN'] if ESM_DIRECTION=='X' else r['Fy_kN'] for r in ESM_RESPONSE]); ax[0].set(title='Floor force',xlabel='Fi (kN)',ylabel='Height (m)')",
        "    ax[1].step([r['story_shear_kN'] for r in ESM_RESPONSE],h,where='mid'); ax[1].set(title='Storey shear',xlabel='Vi (kN)')",
        "    ax[2].plot([1000*r['U_design_m'] for r in ESM_RESPONSE],h,'o-'); ax[2].set(title='Design displacement',xlabel='u (mm)')",
        "    ax[3].barh(h,[r['design_drift_ratio'] for r in ESM_RESPONSE]); ax[3].axvline(DRIFT_LIMIT,color='r',ls='--'); ax[3].set(title='Design drift',xlabel='Drift ratio')",
        "    ax[4].plot([r['story_stiffness_kN_per_m']/1000 for r in ESM_RESPONSE],h,'o-'); ax[4].set(title='Secant storey stiffness',xlabel='ki (MN/m)')",
        "    [a.grid(alpha=.25) for a in ax]; fig.suptitle('NBC 105:2025 '+ESM_CASE+' | V='+str(round(abs(APPLIED_BASE_SHEAR_KN),3))+' kN'); fig.savefig(os.path.join(OUTPUT_DIR,'NBC105_ESM_Response_Profiles_'+ESM_CASE+'.png'),dpi=180); plt.close(fig)",
        "except ImportError: print('Matplotlib not available: response plot skipped.')",
        "try:",
        "    from openpyxl import Workbook",
        "    wb=Workbook(); ws=wb.active; ws.title='ESM Summary'; ws.append(['Item','Value'])",
        "    for item,value in [('Case',ESM_CASE),('Limit state',LIMIT_STATE),('Design base shear (kN)',APPLIED_BASE_SHEAR_KN),('Base reaction (kN)',BASE_REACTION_KN),('Equilibrium error (%)',BASE_SHEAR_ERROR_PCT),('Drift limit',DRIFT_LIMIT),('kD',KD)]: ws.append([item,value])",
        "    wr=wb.create_sheet('OpenSees Response'); wr.append(list(ESM_RESPONSE[0]) if ESM_RESPONSE else []); [wr.append([r.get(k) for k in ESM_RESPONSE[0]]) for r in ESM_RESPONSE]",
        "    wb.save(os.path.join(OUTPUT_DIR,'NBC105_ESM_OpenSees_Response_'+ESM_CASE+'.xlsx'))",
        "except ImportError: print('openpyxl not available: response workbook skipped.')",
        "_processing_timer.finish('completed')",
        "print('OpenSees ESM elapsed (s):', round(time.perf_counter()-_run_started,3))",
        "print('Run log:', LOG_PATH)",
        "print('NBC ESM case completed:', ESM_CASE, '| stories:', len(ESM_RESPONSE), '| design V (kN):', "+repr(result["base_shear_kN"])+")",
        "print('Base reaction (kN):', BASE_REACTION_KN, '| applied shear (kN):', APPLIED_BASE_SHEAR_KN, '| equilibrium error (%):', BASE_SHEAR_ERROR_PCT)",
        "from OpenSees_Response_Checks_V132 import export_static_response",
        "export_static_response(ops,MODEL_BUILDER_SCRIPT,OUTPUT_DIR,ESM_RESPONSE,LIMIT_STATE,ALL_GRAVITY_U,(R_MU if LIMIT_STATE=='ULS' else 1.)*KD)",
        "print('Master-node displacement/drift results saved under', OUTPUT_DIR)",
        "# Interpret displacements using NBC 105:2025 §§5.5.1, 5.5.3, 6.5 and effective stiffness per §3.4.",
    ]
    out_py=out/f"OpenSees_NBC105_ESM_Load_And_Static_Analysis_{suffix}_V111.py"
    safe_write(out_py,"\n".join(script)+"\n")
    return tuple([out_json,out_csv,out_py]+[p for p in (out_png,out_xlsx) if p.is_file()])


class ESMGui:
    def __init__(self,root):
        self.root=root; root.title("NBC 105:2025 — Earthquake Analysis — Equivalent Static Method")
        self.export_dir=None  # initialized after locating the selected export
        export_dir=Path(sys.argv[1]) if len(sys.argv)>1 and Path(sys.argv[1]).is_dir() else Path(r"D:\Thesis Opensees\Etabs Nodes Export")
        self.export_dir=export_dir
        story_csv=export_dir/"NBC105_ESM_Story_Weights_Input_V108.csv"
        self.path=tk.StringVar(value=str(story_csv) if story_csv.is_file() else ""); self.outdir=tk.StringVar(value=str(export_dir if export_dir.is_dir() else HERE))
        self.state=tk.StringVar(value="Both"); self.soil=tk.StringVar(value="D")
        self.T=tk.StringVar(value=""); self.Z=tk.StringVar(value="")
        self.location=tk.StringVar(value=""); self.z_source=tk.StringVar(value="Annex C municipality table")
        self.I=tk.StringVar(value=""); self.vs30=tk.StringVar(value=""); self.n_spt=tk.StringVar(value=""); self.system=tk.StringVar(value="Select / determine from Table 5-2")
        self.rmu=tk.StringVar(value=""); self.omegau=tk.StringVar(value=""); self.omegas=tk.StringVar(value="")
        self.irregular=tk.StringVar(value="Unknown — confirm")
        self.use_workbook_periods=tk.BooleanVar(value=True); self.period_sources={}; self.awx=tk.StringVar(); self.awy=tk.StringVar(); self.period_factor=tk.StringVar(value="1.25")
        self.rayleigh_tx=tk.StringVar(); self.rayleigh_ty=tk.StringVar()
        self.etabs_tx=tk.StringVar(); self.etabs_ty=tk.StringVar()
        self.computed_x=tk.StringVar();self.computed_y=tk.StringVar()
        self.excel_tx=tk.StringVar();self.excel_ty=tk.StringVar()
        self.period_note=tk.StringVar(value="Load ETABS story export; enter directional Aw and optional ETABS modal periods.")
        self.etabs_modal_metadata=None
        self.direction=tk.StringVar(value="Both X/Y"); self.polarity=tk.StringVar(value="Both ±")
        self.centered_case=tk.BooleanVar(value=True); self.include_torsion=tk.BooleanVar(value=True); self.bdim=tk.StringVar(value="")
        self.eccsign=tk.StringVar(value="Both ±")
        self.model_script=tk.StringVar(value="")
        self.python_exe=tk.StringVar(value=sys.executable)
        self.last_runners=[]
        self.status=tk.StringVar(value="Select a verified story-weight CSV, enter project/code inputs, then run the in-house Python NBC calculation.")
        self._build()
        from OpenSees_Complete_Export_V118 import autoselect_inputs
        autoselect_inputs(self, export_dir, 'modal_model_script')
        self._load_project_inputs()
        self.direction.trace_add('write',lambda *_: self._direction_changed())
        self.root.after(250,lambda:self._read_etabs_periods(show_error=False))

    def _field(self,parent,label,var,row,col=0,width=18):
        ttk.Label(parent,text=label).grid(row=row,column=col,sticky="w",padx=5,pady=4)
        ttk.Entry(parent,textvariable=var,width=width).grid(row=row,column=col+1,sticky="ew",padx=5,pady=4)

    def _build(self):
        # Persistent actions remain visible; only the long input form scrolls.
        footer=ttk.Frame(self.root,padding=8);footer.pack(side='bottom',fill='x')
        ttk.Button(footer,text='1. Calculate + export',command=self._calculate).pack(side='left',padx=3)
        self.run_button=ttk.Button(footer,text='2. Run exported cases',command=self._run_generated,state='disabled');self.run_button.pack(side='left',padx=3)
        ttk.Button(footer,text='Open output folder',command=self._open_export_folder).pack(side='left',padx=3)
        ttk.Button(footer,text='Plots…',command=self._view_plots).pack(side='left',padx=3)
        ttk.Button(footer,text='Excel sources…',command=self._choose_workbooks).pack(side='left',padx=3)
        ttk.Button(footer,text='Preliminary Rayleigh run (X/Y)',command=self._run_rayleigh).pack(side='left',padx=3)
        for button in footer.winfo_children():button.pack_forget()
        for index,button in enumerate(footer.winfo_children()):
            button.grid(row=index//3,column=index%3,sticky='ew',padx=3,pady=3)
        for column in range(3):footer.columnconfigure(column,weight=1)
        from OpenSees_Research_UI_V123 import scrollable_form
        top=scrollable_form(self.root)
        ttk.Label(top,text="NBC 105:2025 EQUIVALENT STATIC ANALYSIS",font=("Segoe UI",16,"bold")).pack(anchor="w")
        ttk.Label(top,text="This earthquake-analysis module calculates NBC forces and can run a complete OpenSeesPy model builder followed by the lateral analysis. ETABS EDB extraction remains in the main converter.",wraplength=1000).pack(anchor="w",pady=(2,10))
        ttk.Label(top,text="ETABS linkage: story elevations and diaphragm tags are read from the selected export. The selected assembled ledger contains imported gravity components. Compare its story totals with ETABS before design use.",foreground="#174a7e",wraplength=1000).pack(anchor="w",pady=(0,6))
        filebox=ttk.LabelFrame(top,text="1. Verified story seismic-weight CSV (ETABS export or prepared input)",padding=8); filebox.pack(fill="x",pady=5)
        ttk.Entry(filebox,textvariable=self.path).pack(side="left",fill="x",expand=True,padx=4)
        ttk.Button(filebox,text="Browse…",command=self._browse).pack(side="left",padx=4)
        ttk.Label(filebox,text="The EDB is not read here. Load the story-weight CSV generated by the main ETABS converter, or a correctly formatted CSV prepared from verified data. Enter DL, storage LL and other LL in kN per storey; roof LL is excluded from Wi by Table 5-1. Count self-weight once.",wraplength=850).pack(anchor="w",pady=5)
        periodbox=ttk.LabelFrame(top,text="2. NBC §5.1 directional period: empirical estimate + Rayleigh period",padding=8); periodbox.pack(fill="x",pady=5)
        self._field(periodbox,"Wall index Aw-X (m²)",self.awx,0,0,12); self._field(periodbox,"Wall index Aw-Y (m²)",self.awy,0,2,12)
        self._field(periodbox,"ETABS / Excel reference T1-X (s)",self.etabs_tx,1,0,12); self._field(periodbox,"ETABS / Excel reference T1-Y (s)",self.etabs_ty,1,2,12)
        self._field(periodbox,"Rayleigh T1-X (s)",self.rayleigh_tx,2,0,12); self._field(periodbox,"Rayleigh T1-Y (s)",self.rayleigh_ty,2,2,12)
        self._field(periodbox,"Empirical multiplier (NBC 1.25)",self.period_factor,3,0,12)
        ttk.Button(periodbox,text="Adopt lesser period for selected direction",command=self._update_periods).grid(row=3,column=2,columnspan=2,sticky="w",padx=5)
        ttk.Button(periodbox,text="Open ETABS / refresh X-Y periods",command=self._read_etabs_periods).grid(row=3,column=4,columnspan=3,sticky="w",padx=5)
        ttk.Label(periodbox,textvariable=self.period_note,foreground="#174a7e",wraplength=1000).grid(row=4,column=0,columnspan=7,sticky="w",padx=5,pady=4)
        for row,axis,var in [(6,'X',self.computed_x),(7,'Y',self.computed_y)]:
            ttk.Label(periodbox,text=f'{axis}: empirical / amplified / adopted (s)').grid(row=row,column=0,columnspan=2,sticky='w')
            ttk.Entry(periodbox,textvariable=var,state='readonly',width=55).grid(row=row,column=2,columnspan=5,sticky='ew')
        ttk.Label(periodbox,text="Empirical Aw/T1 follows the supplied workbook and NBC §5.1.2 wall-system expression; include first-storey wall lengths and openings. NBC §5.1.3 amplifies empirical T1 by 1.25. Enter the separately calculated directional Rayleigh period (NBC §5.1.1); adopted design-action T1 is the lesser of Rayleigh and amplified empirical. ETABS eigenperiod is a comparison only and does not replace Rayleigh. Use the modal-analysis dialog to document modal results and Rayleigh inputs.",wraplength=1000).grid(row=5,column=0,columnspan=7,sticky="w",padx=5)
        for row,axis,var in [(8,'X',self.excel_tx),(9,'Y',self.excel_ty)]:
            ttk.Label(periodbox,text=f'{axis}: workbook adopted reference').grid(row=row,column=0,columnspan=2,sticky='w')
            ttk.Entry(periodbox,textvariable=var,state='readonly',width=55).grid(row=row,column=2,columnspan=5,sticky='ew')
        ttk.Label(periodbox,text='Workbook adopted references are not measured eigenperiods. Open ETABS to replace comparison fields with live modal results.',foreground='#8a3412').grid(row=10,column=0,columnspan=7,sticky='w')
        ttk.Checkbutton(periodbox,text="Use documented workbook adopted-period references when fresh Rayleigh results are unavailable",variable=self.use_workbook_periods).grid(row=11,column=0,columnspan=7,sticky="w")
        refs=ttk.LabelFrame(top,text="NBC 105:2025 references",padding=8); refs.pack(fill="x",pady=5)
        ttk.Label(refs,text="NBC 105:2025: §3.2.1 ESM eligibility | §4.1.2/Table 4-1 soil spectrum | §4.1.4 and Annex C/Table C-1 zoning Z | Table 4-4 importance I | §4.2 SLS spectrum | §5.1 period | Table 5-1 seismic weight | Table 5-2 system factors | §§6.1–6.4 coefficient, base shear, vertical distribution and eccentric application | §5.5.1 ULS/SLS deflections | §5.5.3 drift limits | §5.6 accidental eccentricity | §6.5 ESM deflection scale kd. Official PDF: https://ebps.kathmandu.gov.np/UploadFiles/NBC_105_2025_Sesmic_Design.pdf | DUDBC code index: https://www.dudbc.gov.np/pages/24231978/. Enter the exact municipality and verified Annex C/Table C-1 Z, or document map interpolation; the GUI does not guess Z. ETABS mapping: ETABS_to_OpenSees_Node_Mapping.csv and OpenSees_Element_Connectivity.json.",wraplength=1000).pack(anchor="w")
        calc=ttk.LabelFrame(top,text="3. Code/project parameters — verify against the approved design basis",padding=8); calc.pack(fill="x",pady=5)
        calc.columnconfigure(1,weight=1); calc.columnconfigure(3,weight=1)
        ttk.Label(calc,text="Limit state cases").grid(row=0,column=0,sticky="w",padx=5,pady=4); ttk.Combobox(calc,textvariable=self.state,values=["Both","ULS","SLS"],state="readonly",width=16).grid(row=0,column=1,sticky="w",padx=5)
        ttk.Label(calc,text="Soil type (Table 4-1)").grid(row=0,column=2,sticky="w",padx=5,pady=4); ttk.Combobox(calc,textvariable=self.soil,values=["A","B","C","D"],state="readonly",width=16).grid(row=0,column=3,sticky="w",padx=5)
        self._field(calc,"Adopted T1 (s), lesser value per §5.1",self.T,1,0); self._field(calc,"Zoning factor Z, §4.1.4",self.Z,1,2)
        self._field(calc,"Location / municipality",self.location,2,0)
        ttk.Label(calc,text="Z source").grid(row=2,column=2,sticky="w",padx=5,pady=4)
        ttk.Combobox(calc,textvariable=self.z_source,values=["Annex C municipality table","Figure 4-3 map interpolation","Geotechnical/project value"],state="readonly",width=34).grid(row=2,column=3,sticky="ew",padx=5)
        self._field(calc,"Importance factor I, Table 4-4",self.I,3,0)
        ttk.Label(calc,text="Structural system (Table 5-2)").grid(row=3,column=2,sticky="w",padx=5,pady=4)
        self.system_combo=ttk.Combobox(calc,textvariable=self.system,values=["Select / determine from Table 5-2"]+list(NBC105_ESM_SYSTEM_FACTORS),state="readonly",width=34)
        self.system_combo.grid(row=3,column=3,sticky="ew",padx=5)
        self.system_combo.bind("<<ComboboxSelected>>",self._apply_system_factors)
        self._field(calc,"Rμ (ULS; SLS uses spectrum per §4.2)",self.rmu,4,0); self._field(calc,"Ωu (ULS)",self.omegau,4,2)
        self._field(calc,"Ωs (SLS)",self.omegas,5,0)
        ttk.Label(calc,text="ULS regularity (§5.4)").grid(row=5,column=2,sticky="w",padx=5,pady=4)
        ttk.Combobox(calc,textvariable=self.irregular,values=["Unknown — confirm","Regular","Irregular"],state="readonly",width=30).grid(row=5,column=3,sticky="ew",padx=5)
        self._field(calc,"Measured Vs,30 (m/s; required if H>40 m)",self.vs30,6,0)
        self._field(calc,"N-SPT (blows/30 cm; alternative if H<=40 m)",self.n_spt,7,0)
        ttk.Label(calc,text="Choose a Nepal municipality and use its Annex C PGA (Z); otherwise select Figure 4-3 interpolation. Enter the verified Z explicitly. Site soil A–D comes from geotechnical classification (Table 4-2 / §4.1.3), not the municipality. Vs,30 is required for H>40 m; otherwise N-SPT may be used when Vs,30 is unavailable. A missing value or mismatch is flagged in the outputs; results remain exploratory until resolved.",foreground="#8a341c",wraplength=1000).grid(row=8,column=0,columnspan=4,sticky="w",padx=5,pady=5)
        tors=ttk.LabelFrame(top,text="4. Direction and accidental eccentricity",padding=8); tors.pack(fill="x",pady=5)
        ttk.Label(tors,text="Earthquake direction").grid(row=0,column=0,sticky="w",padx=5,pady=4); ttk.Combobox(tors,textvariable=self.direction,values=["Both X/Y","X","Y"],state="readonly",width=8).grid(row=0,column=1,sticky="w",padx=4)
        ttk.Label(tors,text="Polarity").grid(row=0,column=2,sticky="w",padx=5); ttk.Combobox(tors,textvariable=self.polarity,values=["Both ±","+1 (+ direction)","-1 (- direction)"],state="readonly",width=18).grid(row=0,column=3,sticky="w",padx=4)
        ttk.Checkbutton(tors,text="Include accidental eccentricity e = selected sign × 0.05b",variable=self.include_torsion).grid(row=1,column=0,columnspan=2,sticky="w",padx=5,pady=4)
        self._field(tors,"b = plan dimension perpendicular to loading (m)",self.bdim,1,2,14)
        ttk.Label(tors,text="Eccentricity sign").grid(row=2,column=0,sticky="w",padx=5); ttk.Combobox(tors,textvariable=self.eccsign,values=["Both ±","+1","-1"],state="readonly",width=10).grid(row=2,column=1,sticky="w",padx=4)
        ttk.Checkbutton(tors,text="Also export centered E0 cases for clause 5.4 torsion review",variable=self.centered_case).grid(row=4,column=0,columnspan=4,sticky="w")
        ttk.Label(tors,text="The default creates both force polarities and both accidental-eccentricity signs for the selected direction. Both X/Y calculates each axis with its own period and perpendicular plan dimension. Master nodes are assumed at the mass centre; verify inherent eccentricity separately.",wraplength=850).grid(row=3,column=0,columnspan=4,sticky="w",padx=5,pady=5)
        runbox=ttk.LabelFrame(top,text="5. OpenSeesPy model execution",padding=8); runbox.pack(fill="x",pady=5)
        ttk.Label(runbox,text="Complete model script (nodes, sections/elements, diaphragms, mass and gravity)").grid(row=0,column=0,sticky="w",padx=5,pady=3)
        ttk.Entry(runbox,textvariable=self.model_script).grid(row=0,column=1,sticky="ew",padx=5,pady=3)
        ttk.Button(runbox,text="Browse…",command=self._browse_model).grid(row=0,column=2,padx=4)
        ttk.Label(runbox,text="Python executable with OpenSeesPy").grid(row=1,column=0,sticky="w",padx=5,pady=3)
        ttk.Entry(runbox,textvariable=self.python_exe).grid(row=1,column=1,sticky="ew",padx=5,pady=3)
        ttk.Button(runbox,text="Browse…",command=self._browse_python).grid(row=1,column=2,padx=4)
        runbox.columnconfigure(1,weight=1)
        ttk.Label(runbox,text="The runner starts a fresh Python process and builds the model. Before NBC ESM it requires the model builder to attest that both IS 16700:2023 §7.2(d) initial P-Delta cases converged and were reviewed (1.2DL + 0.5IL ± 1.5EL/WL). It then freezes the reviewed gravity state, checks diaphragm masters, applies NBC floor forces/torsion, and reports reactions, drift and storey stiffness. This prerequisite does not replace NBC §3.6 combinations; a partial topology exporter is not a complete analysis model.",foreground="#8a341c",wraplength=1100).grid(row=2,column=0,columnspan=3,sticky="w",padx=5,pady=4)
        out=ttk.LabelFrame(top,text="6. Output folder (same folder as ETABS exports by default)",padding=8); out.pack(fill="x",pady=5)
        ttk.Entry(out,textvariable=self.outdir).pack(side="left",fill="x",expand=True,padx=4); ttk.Button(out,text="Browse…",command=self._browse_out).pack(side="left",padx=4)
        from OpenSees_Research_UI_V123 import copyable_log
        self.run_log=copyable_log(top,self.status)
        ttk.Label(top,text="Workflow: ETABS exports provide story weights, elevations and diaphragm tags. This dialog calculates NBC actions and can execute a complete user-selected OpenSeesPy model-builder script. It does not directly open an ETABS EDB; ETABS API periods can be read when ETABS is open and its modal case has been run.",foreground="#8a341c",wraplength=1100).pack(anchor="w",pady=5)
        if self.path.get(): self._load_story_metadata(self.path.get())

    def _view_plots(self):
        from OpenSees_Results_UI_V127 import esm_comparison
        esm_comparison(self.root,self.outdir.get())
        try:
            from matplotlib.figure import Figure
            from OpenSees_Research_UI_V123 import embed_figures
            import matplotlib.image as image
            paths=sorted(Path(self.outdir.get()).glob('NBC105_ESM_Force_Profile*.png'))
            if not paths:raise ValueError('Calculate cases first to create force-profile figures.')
            figures=[]
            for path in paths:
                fig=Figure(figsize=(10,6),layout='constrained');ax=fig.add_subplot();ax.imshow(image.imread(path));ax.axis('off');figures.append((path.stem.replace('NBC105_ESM_Force_Profile_',''),fig))
            window=tk.Toplevel(self.root);window.title('NBC force profiles | Matplotlib');window.geometry('1250x850');window._esm_output_folder=self.outdir.get();__import__('OpenSees_Case_Selection_V134').profiles(window,figures)
        except Exception as exc:messagebox.showerror('Result plots',str(exc))

    def _load_project_inputs(self):
        try:
            from OpenSees_Project_Inputs_V123 import populate
            data=populate(self,HERE,self.export_dir)
            from OpenSees_ETABS_Periods_V125 import populate as period_references
            refs=period_references(self,HERE,self.export_dir)
            self.etabs_modal_metadata={'source':'Excel references unless live ETABS replaces them','references':refs}
            for key,r in refs.items():
                if isinstance(r,dict) and 'value' in r:data['audit'].append(dict(field='period_reference_'+key,value=r['value'],source=r['source'],sheet=r['sheet'],cell=r['cell']))
            from openpyxl import Workbook
            out=Path(self.outdir.get());out.mkdir(exist_ok=True,parents=True)
            (out/'NBC105_Input_Provenance_V123.json').write_text(json.dumps(data,indent=2),encoding='utf-8')
            wb=Workbook();ws=wb.active;ws.title='Imported inputs';ws.append(['Field','Value','Workbook/result file','Sheet','Cell'])
            for row in data['audit']:ws.append([row[k] for k in ['field','value','source','sheet','cell']])
            notes=wb.create_sheet('Review notes')
            for note in data['notes']:notes.append([note])
            segments=data['values'].get('wall_geometry_segments',[])
            if segments:
                walls=wb.create_sheet('Wall Aw geometry');headers=list(segments[0]);walls.append(headers)
                for segment in segments:walls.append([segment[k] for k in headers])
                walls.freeze_panes='A2';walls.auto_filter.ref=walls.dimensions
            wb.save(out/'NBC105_Input_Provenance_V123.xlsx')
            self.period_note.set('Excel/project inputs loaded with cell-level provenance. Rayleigh values populate only from a matching preliminary run. Regularity remains unconfirmed.')
            self._show_empirical()
            self._direction_changed()
        except Exception as exc:self.status.set('Project input import needs review: '+str(exc))

    def _show_empirical(self):
        rows=read_story_csv(self.path.get());H=max(float(r['height_m']) for r in rows)
        values=[]
        for axis,var in [('X',self.awx),('Y',self.awy)]:
            empirical=.075*H**.75/math.sqrt(float(var.get()))
            amplified=float(self.period_factor.get())*empirical
            ray=(self.rayleigh_tx if axis=='X' else self.rayleigh_ty).get().strip()
            adopted=f'{min(amplified,float(ray)):.9g}' if ray else 'pending Rayleigh run'
            (self.computed_x if axis=='X' else self.computed_y).set(f'{empirical:.9g} / {amplified:.9g} / {adopted}')
            values.append(f'{axis}: empirical={empirical:.6g} s; amplified={amplified:.6g} s; adopted={adopted}')
        self.period_note.set(' | '.join(values)+' | Rayleigh and ETABS modal are separate analysis results.')

    def _direction_changed(self):
        data=getattr(self,'project_input_audit',{}).get('values',{})
        if self.direction.get()=='Both X/Y':
            self._show_empirical();return
        value=data.get('breadth_'+self.direction.get())
        if value is not None:self.bdim.set(str(value))
        ray=(self.rayleigh_tx if self.direction.get()=='X' else self.rayleigh_ty).get()
        if ray:self._update_periods()
        else:self.T.set('')

    def _choose_workbooks(self):
        names=filedialog.askopenfilenames(title='Select calculator and/or scaling workbooks',filetypes=[('Excel workbooks','*.xlsx')])
        if names:
            (self.export_dir/'NBC105_Workbook_Source.json').write_text(json.dumps({'workbooks':list(names)},indent=2),encoding='utf-8')
            self._load_project_inputs()

    def _run_rayleigh(self):
        try:
            awx=float(self.awx.get());awy=float(self.awy.get())
            if min(awx,awy)<=0:raise ValueError('Import and review positive directional wall indices first.')
            script=HERE/'OpenSees_Rayleigh_Preliminary_V123.py'
            command=[self.python_exe.get(),str(script),str(self.export_dir),str(awx),str(awy)]
            out=Path(self.outdir.get());out.mkdir(parents=True,exist_ok=True)
            self.status.set('Running empirical-shape reference loads and Rayleigh calculations for X and Y…')
            def worker():
                try:
                    proc=subprocess.run(command,capture_output=True,text=True,errors='replace', env={**dict(__import__('os').environ),'PYTHONIOENCODING':'utf-8','PYTHONUTF8':'1'}, encoding='utf-8')
                    (out/'NBC105_Rayleigh_Preliminary_V123.log').write_text(proc.stdout+'\n'+proc.stderr,encoding='utf-8')
                    self.root.after(0,lambda text=proc.stdout+'\n'+proc.stderr:self.status.set(text))
                    if proc.returncode:raise RuntimeError(proc.stderr[-1800:])
                    self.root.after(0,self._load_project_inputs)
                    self.root.after(0,lambda:self.status.set('Rayleigh X/Y complete. Imported calculated periods; Excel, PNG, PDF and JSON saved.'))
                except Exception as exc:
                    text=str(exc);self.root.after(0,lambda:self.status.set('Rayleigh run failed: '+text))
            threading.Thread(target=worker,daemon=True).start()
        except Exception as exc:messagebox.showerror('Rayleigh setup',str(exc))

    def _apply_system_factors(self,_event=None):
        factors=NBC105_ESM_SYSTEM_FACTORS.get(self.system.get())
        if factors:
            self.rmu.set(str(factors[0])); self.omegau.set(str(factors[1])); self.omegas.set(str(factors[2]))
    def _browse(self):
        p=filedialog.askopenfilename(title="Select a verified NBC 105 story-weight CSV",initialdir=self.outdir.get(),filetypes=[("CSV","*.csv"),("All files","*.*")])
        if p: self.path.set(p); self.outdir.set(str(Path(p).parent)); self._load_story_metadata(p)
    def _browse_out(self):
        p=filedialog.askdirectory(title="Select output folder",initialdir=self.outdir.get())
        if p: self.outdir.set(p)
    def _browse_model(self):
        p=filedialog.askopenfilename(title="Select complete OpenSeesPy model builder",initialdir=self.outdir.get(),filetypes=[("Python script","*.py"),("All files","*.*")])
        if p: self.model_script.set(p)
    def _browse_python(self):
        p=filedialog.askopenfilename(title="Select Python executable with OpenSeesPy",initialdir=str(Path(sys.executable).parent),filetypes=[("Python executable","python.exe"),("All files","*.*")])
        if p: self.python_exe.set(p)
    def _load_story_metadata(self,path):
        try:
            with open(path,newline="",encoding="utf-8-sig") as f: rows=[r for r in csv.DictReader(f) if str(r.get("story_number","")).strip()]
            blanks=sum(1 for r in rows for k in ("dead_load_kN","storage_live_load_kN","other_live_load_kN") if not str(r.get(k,"")).strip())
            if rows: self.period_note.set(f"Loaded {len(rows)} ETABS stories; roof elevation H={max(float(r['height_m']) for r in rows):.3f} m; diaphragm tags imported. {blanks} DL/LL input cells are blank and require verified ETABS totals.")
        except Exception as exc: self.period_note.set(f"ETABS export metadata could not be read: {exc}")
    def _adopted_axis_period(self,axis,height):
        aw=float((self.awx if axis=='X' else self.awy).get());empirical=float(self.period_factor.get())*.075*height**.75/math.sqrt(aw)
        ray=(self.rayleigh_tx if axis=='X' else self.rayleigh_ty).get().strip()
        if ray:
            self.period_sources[axis]='Calculated directional Rayleigh and amplified empirical minimum'
            return min(empirical,float(ray))
        reference=(self.excel_tx if axis=='X' else self.excel_ty).get().strip()
        if self.use_workbook_periods.get() and reference:
            numeric=getattr(self,'project_input_audit',{}).get('values',{}).get('workbook_adopted_t'+axis.lower())
            value=float(numeric if numeric is not None else reference.split()[0])
            if value<=0 or value>empirical+1e-5:raise ValueError(axis+' imported adopted reference exceeds the current amplified empirical bound; recalculate Rayleigh and resolve the workbook/model difference.')
            self.period_sources[axis]='Imported workbook adopted-period reference; no new Rayleigh calculation claimed'
            return value
        raise ValueError(axis+' requires a calculated Rayleigh period or a documented imported adopted-period reference.')

    def _update_periods(self):
        if self.direction.get()=='Both X/Y':
            self._show_empirical()
            self.period_note.set(self.period_note.get()+f' | Calculated Rayleigh X={self.rayleigh_tx.get() or "pending"} s, Y={self.rayleigh_ty.get() or "pending"} s. Export adopts each axis separately.')
            return
        try:
            with open(self.path.get(),newline="",encoding="utf-8-sig") as f: rows=[r for r in csv.DictReader(f) if str(r.get("story_number","")).strip()]
            H=max(float(r["height_m"]) for r in rows); d=self.direction.get().upper()
            aw=float((self.awx if d=="X" else self.awy).get()); factor=float(self.period_factor.get())
            if aw<=0 or H<=0 or factor<=0: raise ValueError("Enter positive directional Aw and verify the ETABS story heights.")
            empirical=0.075/math.sqrt(aw)*H**0.75; amplified=factor*empirical
            rayleigh=(self.rayleigh_tx if d=="X" else self.rayleigh_ty).get().strip()
            if not rayleigh:
                adopted=self._adopted_axis_period(d,H);self.T.set(f'{adopted:.9g}');self.period_note.set(d+': adopted T1='+str(adopted)+' s | '+self.period_sources[d]);return
            rayleigh=float(rayleigh)
            if rayleigh<=0: raise ValueError('Rayleigh period must be positive.')
            modal=(self.etabs_tx if d=="X" else self.etabs_ty).get().strip()
            adopted=min(amplified,rayleigh)
            self.T.set(f"{adopted:.12g}")
            self.period_note.set(f"{d}: H={H:.3f} m; Aw={aw:.4f} m²; empirical T1={empirical:.4f} s; amplified={amplified:.4f} s; Rayleigh T1={rayleigh:.4f} s; ETABS/Excel period reference={modal or 'not supplied'}; adopted T1=min(amplified empirical, Rayleigh)={adopted:.4f} s.")
        except Exception as exc: messagebox.showerror("Period calculation needs review",str(exc))
    def _read_etabs_periods(self, show_error=True):
        """Read analyzed ETABS modal periods and mass ratios from the active ETABS session."""
        try:
            import comtypes.client
            from OpenSees_ETABS_Periods_V125 import attach_or_open
            self.status.set('Opening/attaching to the source ETABS model…' if show_error else 'Checking live ETABS period references…')
            self.root.update_idletasks()
            etabs=attach_or_open(self.export_dir,open_if_needed=show_error)
            sap=etabs.SapModel
            # Select a single existing modal case before querying results.
            cases=sap.LoadCases.GetNameList()
            case_names=next((list(x) for x in cases if isinstance(x,(tuple,list))),[])
            modal_cases=[]
            for name in case_names:
                kind=sap.LoadCases.GetTypeOAPI(name)
                if int(kind[0])==3: modal_cases.append(name)
            if len(modal_cases)!=1: raise RuntimeError('Select one modal case in ETABS; found '+str(modal_cases))
            sap.Results.Setup.DeselectAllCasesAndCombosForOutput()
            sap.Results.Setup.SetCaseSelectedForOutput(modal_cases[0])
            raw=sap.Results.ModalParticipatingMassRatios()
            if not isinstance(raw,(tuple,list)) or len(raw)<9:
                raise RuntimeError("ETABS returned an unexpected ModalParticipatingMassRatios result. Verify CSI API version and run a modal case.")
            # CSI API order has LoadCase, StepType, StepNum, Period, UX, UY...
            # COM wrappers differ on whether the scalar return code is first or last;
            # anchor extraction on the LoadCase string SAFEARRAY instead of assuming.
            string_arrays=[]
            for idx,value in enumerate(raw):
                try:
                    seq=list(value)
                    if seq and all(isinstance(v,str) for v in seq): string_arrays.append((idx,seq))
                except Exception: pass
            if not string_arrays: raise RuntimeError("ETABS returned no modal result rows for the selected case. Run ETABS analysis with a modal load case and confirm its results are available; OpenSees modal results do not populate the ETABS Results API.")
            case_idx,cases=string_arrays[0]; period_idx=case_idx+3
            periods=list(raw[period_idx]); ux=list(raw[period_idx+1]); uy=list(raw[period_idx+2])
            scalar_codes=[]
            for value in (raw[0],raw[-1]):
                try:
                    if isinstance(value,(int,float)): scalar_codes.append(int(value))
                except Exception: pass
            ret=scalar_codes[-1] if scalar_codes else 0
            if ret!=0: raise RuntimeError(f"ETABS ModalParticipatingMassRatios returned code {ret}; run/verify modal analysis first.")
            n=min(len(periods),len(ux),len(uy),len(cases)); modes=[]
            for i in range(n):
                p=float(periods[i]); rx=float(ux[i]); ry=float(uy[i])
                if math.isfinite(p) and p>0: modes.append((p,rx,ry,str(cases[i]),i+1))
            if not modes: raise RuntimeError("No positive modal periods were returned. Run a modal analysis in ETABS.")
            mx=max(modes,key=lambda a:a[1]); my=max(modes,key=lambda a:a[2])
            self.etabs_tx.set(f"{mx[0]:.8g}"); self.etabs_ty.set(f"{my[0]:.8g}")
            self.etabs_modal_metadata={"source":"Active ETABS API Results.ModalParticipatingMassRatios","X_dominant":{"mode_number":mx[4],"period_s":mx[0],"UX_ratio":mx[1],"load_case":mx[3]},"Y_dominant":{"mode_number":my[4],"period_s":my[0],"UY_ratio":my[2],"load_case":my[3]}}
            from OpenSees_ETABS_Periods_V125 import save_live
            save_live(self.export_dir,mx[0],my[0],self.etabs_modal_metadata)
            self.period_note.set(f"ETABS autofill complete: T_X={mx[0]:.5g}s from UX-dominant mode {mx[4]} ({mx[1]:.1%} UX); T_Y={my[0]:.5g}s from UY-dominant mode {my[4]} ({my[2]:.1%} UY). These ETABS eigenperiods are comparison values only; the NBC adopted ESM period is still the lesser of 1.25×empirical and the directional Rayleigh period. Verify mode shapes when torsion couples with translation.")
        except Exception as exc:
            if show_error: messagebox.showerror("ETABS modal results unavailable",f"ETABS comparison periods could not be refreshed. This does not invalidate the separate workbook-adopted or Rayleigh periods already shown. To refresh ETABS references, run a modal case in ETABS and retry.\n\n{exc}")
            else: self.status.set('Excel/source inputs loaded. ETABS modal refresh unavailable: '+str(exc))
    def _open_export_folder(self):
        try: os.startfile(self.outdir.get())
        except Exception as e: messagebox.showerror("Could not open export folder",str(e))
    def _open_csv(self):
        if not self.path.get(): self._browse(); return
        try: os.startfile(self.path.get())
        except Exception as e: messagebox.showerror("Could not open CSV",str(e))
    def _calculate(self):
        if self.direction.get()=='Both X/Y':
            if not all((self.rayleigh_tx if axis=='X' else self.rayleigh_ty).get().strip() or (self.use_workbook_periods.get() and (self.excel_tx if axis=='X' else self.excel_ty).get().strip()) for axis in ('X','Y')):
                self.last_runners=[];self.run_button.configure(state='disabled')
                self.status.set('Both X/Y export needs calculated Rayleigh periods in both axes. Excel reference periods are already shown separately. Run Preliminary Rayleigh (X/Y), then Calculate + export.')
                messagebox.showerror('Directional periods needed',self.status.get());return
            all_runners=[]
            axis_counts=[]
            for axis in ('X','Y'):
                self.last_runners=[]
                self.direction.set(axis);self._calculate()
                all_runners.extend(self.last_runners)
                axis_counts.append(len(self.last_runners))
            self.last_runners=all_runners;self.direction.set('Both X/Y')
            if not all(axis_counts):
                self.last_runners=[];self.run_button.configure(state='disabled')
                self.status.set('Both-axis export incomplete. Resolve the axis-specific errors in the run log; no combined run is enabled.');return
            self.status.set(f'Combined X/Y export: {len(all_runners)} runners. Each direction uses its own adopted period.')
            return
        started=time.perf_counter()
        try:
            if not self.path.get(): raise ValueError("Choose the verified story-weight CSV first; this dialog does not read an EDB directly.")
            self._update_periods()
            if not self.T.get().strip() or not self.Z.get().strip(): raise ValueError("Enter the adopted T1 and seismic zoning factor Z from project/geotechnical/code data.")
            self._adopted_axis_period(self.direction.get(),max(float(r["height_m"]) for r in read_story_csv(self.path.get())))
            if not self.location.get().strip(): raise ValueError("Enter the project municipality/location used to establish Z.")
            if self.system.get()=="Select / determine from Table 5-2": raise ValueError("Choose the verified structural system in NBC Table 5-2 so the Rμ and Ω factors are populated.")
            rows=read_story_csv(self.path.get())
            # Signed stability checks use a documented ULS coefficient reference.
            # This does not make high-rise ULS ESM a permitted design method.
            references={}
            height=max(float(r['height_m']) for r in rows)
            for axis,awvar,rayvar in [('X',self.awx,self.rayleigh_tx),('Y',self.awy,self.rayleigh_ty)]:
                adopted=self._adopted_axis_period(axis,height)
                reference=calculate_esm(rows,period_s=adopted,soil_type=self.soil.get(),zoning_factor_z=float(self.Z.get()),importance_factor_i=float(self.I.get()),limit_state='ULS',r_mu=float(self.rmu.get()),omega_u=float(self.omegau.get()),vs30_m_s=self.vs30.get(),n_spt_blow_30cm=self.n_spt.get(),diagnostic_reference=True)
                references[axis]=dict(period_s=adopted,stories=reference['stories'],basis='NBC ULS coefficient and floor distribution used only as a diagnostic EL reference; high-rise design requires the applicable dynamic method.',code_eligible=reference['code_eligible'],site_subsoil_status=reference['site_subsoil_status'])
            safe_write(Path(self.model_script.get()).parent/'OpenSees_Initial_PDelta_Input_V124.json',json.dumps(references,indent=2))
            H=max(float(r["height_m"]) for r in rows)
            selected=self.state.get(); states=["ULS","SLS"] if selected=="Both" else [selected]
            irr={"Regular":False,"Irregular":True,"Unknown — confirm":None}[self.irregular.get()]
            if self.include_torsion.get():
                if not self.bdim.get().strip(): raise ValueError("Enter the plan dimension b perpendicular to the selected earthquake direction.")
                bdim=float(self.bdim.get())
                if bdim<=0: raise ValueError("The perpendicular plan dimension b must be positive.")
                ecc_signs=[1,-1] if self.eccsign.get()=="Both ±" else [int(self.eccsign.get())]
            else: bdim=0.0; ecc_signs=[0]
            if self.include_torsion.get() and self.centered_case.get() and 0 not in ecc_signs:ecc_signs.append(0)
            polarities=[1,-1] if self.polarity.get()=="Both ±" else [1 if self.polarity.get().startswith("+1") else -1]
            made=[]; summaries=[]; failures=[]; self.last_runners=[]
            for state in states:
                try:
                    result=calculate_esm(rows,period_s=self._adopted_axis_period(self.direction.get(),H),soil_type=self.soil.get(),zoning_factor_z=float(self.Z.get()),importance_factor_i=float(self.I.get()),limit_state=state,r_mu=float(self.rmu.get() or 1.0),omega_u=float(self.omegau.get() or 1.0),omega_s=float(self.omegas.get() or 1.0),irregular=irr,location_name=self.location.get(),zoning_source=self.z_source.get(),vs30_m_s=self.vs30.get(),n_spt_blow_30cm=self.n_spt.get(),diagnostic_reference=True)
                    if not result["code_eligible"]:summaries.append("ULS COMPARISON ONLY: "+result["esm_eligibility"])
                    if not result["site_subsoil_verified"]:summaries.append("SITE CLASSIFICATION NOT VERIFIED: "+result["site_subsoil_status"])
                    result["period_source"]=self.period_sources.get(self.direction.get(),"Directional Rayleigh/empirical calculation")
                    result["input_provenance"]=getattr(self,"project_input_audit",{})
                    result["period_selection"]={"direction":self.direction.get(),"empirical_Aw_X_m2":self.awx.get() or None,"empirical_Aw_Y_m2":self.awy.get() or None,"empirical_multiplier":self.period_factor.get(),"Rayleigh_T1_X_s":self.rayleigh_tx.get() or None,"Rayleigh_T1_Y_s":self.rayleigh_ty.get() or None,"ETABS_or_Excel_reference_T1_X_s":self.etabs_tx.get() or None,"ETABS_or_Excel_reference_T1_Y_s":self.etabs_ty.get() or None,"ETABS_modal_metadata":self.etabs_modal_metadata,"selected_period_s":self.T.get(),"selection_rule":"NBC §5.1: adopt the lesser of 1.25 times the empirical period and the directional Rayleigh period. ETABS eigenperiod is a comparison only."}
                    for polarity in polarities:
                        for ecc_sign in ecc_signs:
                            breadth=(bdim,ecc_sign)
                            output_paths=write_outputs(self.path.get(),rows,result,self.direction.get(),polarity,self.include_torsion.get(),breadth,self.outdir.get(),self.model_script.get())
                            made.extend(output_paths); self.last_runners.append(str(output_paths[2]))
                            summaries.append(f"{state} {self.direction.get()}{'+' if polarity>0 else '-'} E{('+' if ecc_sign>0 else '-') if self.include_torsion.get() and ecc_sign else '0'}: V={result['base_shear_kN']:.3f} kN, k={result['period_exponent_k']:.4f}")
                except Exception as state_exc:
                    failures.append(f"{state}: {state_exc}")
            if not made: raise ValueError("No case was calculated. " + " | ".join(failures))
            elapsed=time.perf_counter()-started
            files=", ".join(Path(p).name for p in made)
            note=("\nBlocked cases: " + " | ".join(failures)) if failures else ""
            self.status.set(f"Completed in {elapsed:.2f} s | Location={self.location.get()} | soil={self.soil.get()} | Z={self.Z.get()} ({self.z_source.get()})\n"+"\n".join(summaries)+note+f"\nCreated: {files}")
            self.run_button.configure(state="normal" if Path(self.model_script.get()).is_file() else "disabled")
            messagebox.showinfo("NBC 105 ESM calculation",f"Created {len(self.last_runners)} case runner(s) and documented force profiles.\n" + ("Blocked: "+" | ".join(failures) if failures else "Review inputs before running the exported cases.") + "\n\nUse Plots or Open output folder to inspect the results.")
        except Exception as exc:
            self.status.set("Blocked: "+str(exc)); messagebox.showerror("ESM calculation blocked",str(exc))

    def _run_generated(self):
        if not self.last_runners:
            messagebox.showerror("No generated cases","Calculate the selected ESM case(s) first."); return
        if not Path(self.model_script.get()).is_file():
            messagebox.showerror("Model builder required","Select a complete OpenSeesPy model-building script first."); return
        if not Path(self.python_exe.get()).is_file():
            messagebox.showerror("Python executable required","Select the Python executable with OpenSeesPy installed."); return
        self.run_button.configure(state="disabled")
        self.status.set("Running OpenSeesPy ESM case(s)… per-case logs are being written to the output folder.")
        def worker():
            outcomes=[]
            for script in self.last_runners:
                try:
                    proc=subprocess.Popen([self.python_exe.get(),'-u',script],cwd=self.outdir.get(),stdout=subprocess.PIPE,stderr=subprocess.STDOUT,text=True, env={**dict(__import__('os').environ),'PYTHONIOENCODING':'utf-8','PYTHONUTF8':'1'}, encoding='utf-8')
                    chunks=[]
                    for line in proc.stdout:
                        chunks.append(line);self.root.after(0,lambda text=line:self.status.set(text.rstrip()))
                    proc.wait();proc.stdout=''.join(chunks);proc.stderr='' 
                    log_path=Path(self.outdir.get())/("GUI_Run_"+Path(script).stem+".log")
                    safe_write(log_path,"COMMAND: "+repr([self.python_exe.get(),script])+"\nRETURN CODE: "+str(proc.returncode)+"\n--- STDOUT ---\n"+proc.stdout+"\n--- STDERR ---\n"+proc.stderr)
                    outcomes.append((Path(script).name,proc.returncode,str(log_path)))
                except Exception as exc:
                    outcomes.append((Path(script).name,-999,str(exc)))
            def done():
                self.run_button.configure(state="normal")
                lines=[f"{name}: {'SUCCESS' if code==0 else 'FAILED'} | {log}" for name,code,log in outcomes]
                self.status.set("OpenSeesPy earthquake cases finished.\n"+"\n".join(lines))
                if all(code==0 for _,code,_ in outcomes): messagebox.showinfo("ESM analysis complete",self.status.get())
                else: messagebox.showerror("ESM analysis needs review",self.status.get())
            self.root.after(0,done)
        threading.Thread(target=worker,daemon=True).start()


def main():
    root=tk.Tk(); ESMGui(root); root.mainloop()
if __name__=="__main__": main()



