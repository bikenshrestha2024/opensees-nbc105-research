# -*- coding: utf-8 -*-
"""Small GUI to run the NBC 105:2025 eigen/modal-mass audit on a model script."""
from __future__ import annotations

__release__ = '134'
from OpenSees_Timing_V130 import timed_outputs

import json
import os
import subprocess
import sys
import threading
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, ttk


def _runner_source(model, output, nmodes, solver, period_inputs, mode_x, mode_y, limit_state='ULS'):
    helper_dir = str(Path(__file__).resolve().parent)
    return f'''import json, runpy, sys, traceback, os
from pathlib import Path
MODEL = Path({str(model)!r})
OUT = Path({str(output)!r})
sys.path.insert(0, str(MODEL.parent))
sys.path.insert(0, {helper_dir!r})
try:
    if not MODEL.is_file(): raise FileNotFoundError(MODEL)
    # This script must build the final model and assign the reviewed mass.
    # Include the intended gravity preload in the same script if P-Delta effects
    # are to be present in the eigenanalysis.
    os.environ['OPENSEES_LIMIT_STATE']={limit_state!r}
    runpy.run_path(str(MODEL), run_name='opensees_modal_model')
    import openseespy.opensees as ops
    from OpenSees_Modal_Analysis_NBC105_V111 import run_nbc105_modal_audit
    run_nbc105_modal_audit(ops, nmodes={int(nmodes)}, output_dir=OUT,
                           solver={str(solver)!r},
                           period_inputs={period_inputs!r},
                           fundamental_mode_x={mode_x!r}, fundamental_mode_y={mode_y!r})
    import hashlib
    data=MODEL.parent/'OpenSees_Complete_Model_Data_V118.json'
    target=OUT/'NBC105_Modal_V111.json'
    if data.is_file():
        result=json.loads(target.read_text(encoding='utf-8-sig'))
        result['model_sha256']=hashlib.sha256(data.read_bytes()).hexdigest()
        result['limit_state']={limit_state!r}
        from OpenSees_Model_Audit_V130 import implementation_digest, modal_numeric_digest
        result['implementation_sha256']=implementation_digest(Path({helper_dir!r}))
        result['modal_numeric_sha256']=modal_numeric_digest(Path({helper_dir!r}))
        target.write_text(json.dumps(result,indent=2),encoding='utf-8')
except Exception as exc:
    (OUT / 'NBC105_Modal_V111_error.txt').write_text(
        str(exc) + '\\n' + traceback.format_exc(), encoding='utf-8')
    traceback.print_exc(); sys.exit(1)
'''


@timed_outputs('Modal plot and workbook export','output')
def _export_plot_and_excel(result, output):
    """Create optional Matplotlib and Excel views from the audit JSON."""
    output = Path(output)
    rows = result["rows"]
    warnings = []
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        modes = [r["mode"] for r in rows]
        fig, axes = plt.subplots(1, 2, figsize=(12, 4.8))
        fig.suptitle(str(result.get('limit_state','Unspecified'))+' stiffness case — modal audit')
        axes[0].plot(modes, [r["period_s"] for r in rows], marker="o", color="#345995")
        pcheck = result.get("NBC_5_1_period_check")
        if pcheck:
            for axis, color in (("X", "#d1495b"), ("Y", "#00798c")):
                val = pcheck.get(f"amplified_empirical_{axis}_s")
                if val is not None:
                    axes[0].axhline(val, color=color, linestyle="--", alpha=.8, label=f"NBC amplified empirical {axis}")
                val = pcheck.get(f"rayleigh_{axis}_s")
                if val is not None:
                    axes[0].axhline(val, color=color, linestyle=":", alpha=.8, label=f"Rayleigh {axis}")
        axes[0].set(title="Computed modal periods", xlabel="Mode", ylabel="Period (s)")
        axes[0].grid(True, alpha=.3)
        if pcheck: axes[0].legend(fontsize=7)
        axes[1].plot(modes, [r["UX_cumulative_percent"] for r in rows], marker="o", label="UX", color="#d1495b")
        axes[1].plot(modes, [r["UY_cumulative_percent"] for r in rows], marker="o", label="UY", color="#00798c")
        axes[1].axhline(90.0, color="black", linestyle="--", linewidth=1, label="NBC 90% target")
        axes[1].set(title="Cumulative modal mass", xlabel="Mode", ylabel="Cumulative participation (%)", ylim=(0, 105))
        axes[1].grid(True, alpha=.3); axes[1].legend()
        fig.suptitle("NBC 105:2025 modal audit — eigen results, not response-spectrum design actions")
        fig.tight_layout()
        fig.savefig(output / "NBC105_Modal_V111.png", dpi=300)
        fig.savefig(output / "NBC105_Modal_V111.pdf")
        plt.close(fig)
    except Exception as exc:
        warnings.append(f"Matplotlib plot not created: {exc}")
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook(); ws = wb.active; ws.title = "Modal Results"
        headers = list(rows[0].keys()) if rows else []
        ws.append(headers)
        for row in rows: ws.append([row.get(h) for h in headers])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="244062")
        ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            letter = col[0].column_letter
            ws.column_dimensions[letter].width = min(38, max(14, max(len(str(c.value or "")) for c in col) + 2))
        summary = wb.create_sheet("NBC Summary")
        for row in [
            ("Standard", result["standard"]), ("Clauses", ", ".join(result["clauses"])),
            ("Requested modes", result["requested_modes"]), ("Returned modes", result["returned_modes"]),
            ("Cumulative UX (%)", result["cumulative_UX_percent"]), ("Cumulative UY (%)", result["cumulative_UY_percent"]),
            ("UX 90% reached", result["UX_90_percent_reached"]), ("UY 90% reached", result["UY_90_percent_reached"]),
            ("First mode reaching target", str(result["first_mode_reaching_target"])),
            ("Modes >=33 Hz", str(result["modes_at_or_above_33_Hz"])),
            ("Missing-mass correction", result["missing_mass_correction"]),
            ("Modal combination", result["modal_combination"]), ("Base-shear scaling", result["base_shear_scaling"]),
            ("Interpretation", "Eigen/modal-mass audit only; not a complete NBC response-spectrum analysis."),
        ]: summary.append(row)
        if pcheck:
            summary.append(("NBC 5.1 system family", pcheck.get("system_type")))
            summary.append(("NBC 5.1 height H (m)", pcheck.get("height_m")))
            for axis in ("X", "Y"):
                for label, key in (("Wall Aw (m²)", f"wall_Aw_{axis}_m2"),
                                   ("Empirical period (s)", f"empirical_{axis}_s"),
                                   ("Amplified empirical period (s)", f"amplified_empirical_{axis}_s"),
                                   ("Rayleigh period (s)", f"rayleigh_{axis}_s"),
                                   ("Adopted design period (s)", f"adopted_design_period_{axis}_s"),
                                   ("Modal/adopted ratio", f"modal_to_adopted_period_ratio_{axis}")):
                    summary.append((f"{axis} — {label}", pcheck.get(key)))
                summary.append((f"{axis} — status", pcheck.get(f"adopted_period_status_{axis}")))
        for cell in summary[1]: cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="244062")
        summary.column_dimensions["A"].width = 34; summary.column_dimensions["B"].width = 110
        wb.save(output / "NBC105_Modal_V111.xlsx")
    except Exception as exc:
        warnings.append(f"Excel workbook not created: {exc}")
    return warnings


class ModalAuditGUI:
    def __init__(self, parent=None):
        self._owns_root = parent is None
        self.root = tk.Toplevel(parent) if parent is not None else tk.Tk()
        self.root.title("NBC 105:2025 — OpenSees Modal Audit")
        self.root.geometry("1120x780")
        self.root.minsize(900, 680)
        self.model = tk.StringVar()
        self.python = tk.StringVar(value=sys.executable)
        self.output = tk.StringVar(value=str(Path.cwd()))
        self.nmodes = tk.StringVar(value="40")
        self.solver = tk.StringVar(value="PythonSparse")
        self.limit_state=tk.StringVar(value='ULS')
        self.height = tk.StringVar()
        self.period_system = tk.StringVar(value="concrete shear-wall building")
        self.project_system=tk.StringVar(value='Project structural system: load the reviewed Excel source; confirm the period formula separately.')
        self.aw_x = tk.StringVar()
        self.aw_y = tk.StringVar()
        self.tr_x = tk.StringVar()
        self.tr_y = tk.StringVar()
        self.mode_x = tk.StringVar()
        self.mode_y = tk.StringVar()
        self.confirm = tk.BooleanVar(value=False)
        self.status = tk.StringVar(value="Select an OpenSees model script to run the modal audit.")
        self._build()
        self.root.after_idle(lambda:__import__('OpenSees_Complete_Export_V118').autoselect_inputs(self,Path(self.model.get()).parent if self.model.get() else r'D:/Thesis Opensees/Etabs Nodes Export','modal_model_script'))
        self.model.trace_add('write',lambda *args:self._sync_project_system())
        from OpenSees_Research_UI_V123 import scrollable_form


    def _build(self):
        from OpenSees_Research_UI_V123 import scrollable_form
        outer = scrollable_form(self.root,padding=14)
        tools=ttk.Frame(outer);tools.pack(fill='x',pady=4)
        ttk.Button(tools,text='Open output folder',command=lambda:os.startfile(self.output.get())).pack(side='left',padx=4)
        ttk.Button(tools,text='View modal figure',command=self._view_plot).pack(side='left',padx=4)
        ttk.Label(outer, text="NBC 105:2025 MODAL ANALYSIS AUDIT", font=("Segoe UI", 16, "bold")).pack(anchor="w")
        tk.Label(outer, text="Eigen periods + frequency + directional modal mass participation + NBC checks (§7.3–7.5)", wraplength=850, anchor="w", justify="left").pack(anchor="w", pady=(4, 12))
        form = ttk.LabelFrame(outer, text="Run inputs", padding=10); form.pack(fill="x")
        self._path_row(form, "OpenSees model script", self.model, self._pick_model)
        self._path_row(form, "Python executable with OpenSeesPy", self.python, self._pick_python)
        self._path_row(form, "Output folder", self.output, self._pick_output)
        row = ttk.Frame(form); row.pack(fill="x", pady=4)
        ttk.Label(row, text="Requested modes").pack(side="left")
        ttk.Entry(row, textvariable=self.nmodes, width=8).pack(side="left", padx=(8, 18))
        ttk.Label(row, text="Stiffness case / eigen solver").pack(side="left")
        ttk.Combobox(row,textvariable=self.limit_state,values=('ULS','SLS','Both'),state='readonly',width=10).pack(side='left',padx=6)
        ttk.Combobox(row, textvariable=self.solver, values=("PythonSparse", "-genBandArpack", "-fullGenLapack"), state="readonly", width=22).pack(side="left", padx=8)
        ttk.Label(form,text="PythonSparse is recommended for building-scale models. Failed/partial native solves fall back to PythonSparse; dense fullGenLapack is skipped above 500 nodes.",foreground="#8a3412",wraplength=900).pack(anchor="w")
        ttk.Label(form,textvariable=self.project_system).pack(anchor='w')
        period = ttk.LabelFrame(form, text="NBC §5.1 period comparison (SI units)", padding=8); period.pack(fill="x", pady=(8, 2))
        prow = ttk.Frame(period); prow.pack(fill="x", pady=2)
        ttk.Label(prow, text="Seismic height H (m)").pack(side="left")
        ttk.Entry(prow, textvariable=self.height, width=10).pack(side="left", padx=(6, 14))
        ttk.Label(prow, text="NBC §5.1 period formula (not the Table 5-2 system factor)").pack(side="left")
        ttk.Combobox(prow, textvariable=self.period_system, values=("concrete shear-wall building", "concrete moment-resisting frame", "steel moment-resisting frame", "steel eccentrically braced frame", "other structural system"), state="readonly", width=42).pack(side="left", padx=6)
        prow2 = ttk.Frame(period); prow2.pack(fill="x", pady=2)
        ttk.Label(prow2, text="Wall Aw X (m²)").pack(side="left")
        ttk.Entry(prow2, textvariable=self.aw_x, width=10).pack(side="left", padx=(6, 14))
        ttk.Label(prow2, text="Wall Aw Y (m²)").pack(side="left")
        ttk.Entry(prow2, textvariable=self.aw_y, width=10).pack(side="left", padx=(6, 14))
        tk.Label(prow2, text="RC dual systems with shear walls: use the reviewed wall-period expression here; system factors remain separate. For wall systems, enter directional NBC §5.1.2 Aw from first-storey effective wall segments; review openings.", wraplength=360, anchor="w", justify="left").pack(side="left")
        prow3 = ttk.Frame(period); prow3.pack(fill="x", pady=2)
        ttk.Label(prow3, text="Rayleigh T X (s)").pack(side="left")
        ttk.Entry(prow3, textvariable=self.tr_x, width=10).pack(side="left", padx=(6, 14))
        ttk.Label(prow3, text="Rayleigh T Y (s)").pack(side="left")
        ttk.Entry(prow3, textvariable=self.tr_y, width=10).pack(side="left", padx=(6, 14))
        ttk.Label(prow3, text="Optional modal mode IDs reviewed as fundamental X / Y").pack(side="left")
        ttk.Entry(prow3, textvariable=self.mode_x, width=5).pack(side="left", padx=4)
        ttk.Entry(prow3, textvariable=self.mode_y, width=5).pack(side="left", padx=4)
        tk.Label(period, text="NBC §5.1.3: amplified empirical period = 1.25 × empirical; adopted design period = the lesser of this and Rayleigh period. ETABS eigenperiod is a comparison diagnostic; it does not replace NBC's adopted period.", wraplength=850, foreground="#9a4b24", anchor="w", justify="left").pack(anchor="w", pady=(4, 0))
        tk.Label(form, text="Select a complete OpenSees model builder with the intended geometry, constraints, seismic mass and gravity state. This audit runs it in a separate process and does not retain its domain for later analyses.", wraplength=870, foreground="#9a4b24", anchor="w", justify="left").pack(anchor="w", pady=(8, 4))
        tk.Checkbutton(form, text="I reviewed the model script, mass source, constraints and gravity preload for this modal run.", variable=self.confirm, wraplength=840, anchor="w", justify="left").pack(anchor="w", pady=5)
        self.run_button=ttk.Button(outer, text="Run selected SLS / ULS modal audit", command=self._run);self.run_button.pack(anchor="w", pady=10)
        from OpenSees_Research_UI_V123 import copyable_log
        from components.shared.case_status import panel
        self.case_board=panel(outer)
        self.run_log=copyable_log(outer,self.status,height=6)
        box = ttk.LabelFrame(outer, text="NBC interpretation", padding=10); box.pack(fill="both", expand=True, pady=(10, 0))
        notes = (
            "• NBC 105:2025 §7.3: target at least 90% cumulative seismic mass in each horizontal direction.\n"
            "• Modal combination uses modes below 33 Hz; effects of modes at/above 33 Hz require missing-mass correction.\n"
            "• Inspect mode shapes and classify modes belonging to the horizontal load-resisting system.\n"
            "• Apply CQC/SRSS and closely-spaced-mode provisions under §7.4; compare MRSM and ESM base shear and apply §7.5 scaling when required.\n\n"
            "This audit computes eigen/modal-mass results and flags required follow-up. It does not calculate response-spectrum member forces, missing-mass correction, modal combination, or base-shear scaling."
        )
        tk.Label(box, text=notes, justify="left", wraplength=850, anchor="nw").pack(anchor="nw")

    def _sync_project_system(self):
        try:
            from OpenSees_Project_Inputs_V123 import import_project_inputs
            data=import_project_inputs(Path(__file__).parent,Path(self.model.get()).parent)
            system=str(data['values'].get('system_excel','')).strip()
            if not system:return
            text=system.lower()
            if 'wall' in text:self.period_system.set('concrete shear-wall building')
            elif 'concrete' in text or 'rc ' in text:self.period_system.set('concrete moment-resisting frame')
            elif 'steel' in text and 'eccentr' in text:self.period_system.set('steel eccentrically braced frame')
            elif 'steel' in text and 'moment' in text:self.period_system.set('steel moment-resisting frame')
            else:self.period_system.set('other structural system')
            self.project_system.set('Excel structural system: '+system+' | Corresponding period formula shown below; review applicability. ESM system factors remain separate.')
        except Exception as exc:
            self.project_system.set('Structural-system import unavailable: '+str(exc))

    def _path_row(self, parent, label, var, command):
        row = ttk.Frame(parent); row.pack(fill="x", pady=3)
        ttk.Label(row, text=label, width=31).pack(side="left")
        ttk.Entry(row, textvariable=var).pack(side="left", fill="x", expand=True, padx=6)
        ttk.Button(row, text="Browse…", command=command).pack(side="left")

    def _pick_model(self):
        p = filedialog.askopenfilename(parent=self.root, filetypes=[("Python scripts", "*.py"), ("All files", "*.*")])
        if p: self.model.set(p)

    def _view_plot(self):
        from OpenSees_Results_UI_V127 import modal_results
        try:modal_results(self.root,self.output.get())
        except Exception as exc:messagebox.showerror('Modal results',str(exc),parent=self.root)

    def _pick_python(self):
        p = filedialog.askopenfilename(parent=self.root, filetypes=[("Python", "python.exe"), ("All files", "*.*")])
        if p: self.python.set(p)

    def _pick_output(self):
        p = filedialog.askdirectory(parent=self.root)
        if p: self.output.set(p)

    def _run(self):
        try:
            model = Path(self.model.get()).expanduser().resolve()
            python = Path(self.python.get()).expanduser().resolve()
            output = Path(self.output.get()).expanduser().resolve()
            n = int(self.nmodes.get())
            if not model.is_file(): raise ValueError("Select an existing model script.")
            if not python.is_file(): raise ValueError("Select the Python executable that has OpenSeesPy installed.")
            if n < 1: raise ValueError("Requested mode count must be positive.")
            if not self.confirm.get(): raise ValueError("Review the model inputs and check the acknowledgement before running.")
            output.mkdir(parents=True, exist_ok=True)
            period_inputs = None
            if self.height.get().strip():
                period_inputs = {"height_m": float(self.height.get()), "system_type": self.period_system.get()}
                if self.period_system.get() == "concrete shear-wall building":
                    period_inputs["wall_aw_x_m2"] = float(self.aw_x.get()) if self.aw_x.get().strip() else None
                    period_inputs["wall_aw_y_m2"] = float(self.aw_y.get()) if self.aw_y.get().strip() else None
                period_inputs["rayleigh_x_s"] = float(self.tr_x.get()) if self.tr_x.get().strip() else None
                period_inputs["rayleigh_y_s"] = float(self.tr_y.get()) if self.tr_y.get().strip() else None
            mode_x = int(self.mode_x.get()) if self.mode_x.get().strip() else None
            mode_y = int(self.mode_y.get()) if self.mode_y.get().strip() else None
        except Exception as exc:
            messagebox.showerror("Modal audit inputs", str(exc), parent=self.root); return
        states=('ULS','SLS') if self.limit_state.get()=='Both' else (self.limit_state.get(),)
        from components.shared.case_status import CaseLedger
        ledger=CaseLedger(output/"Modal_Batch_Status_V151.json",states);self.case_board["ledger"]=ledger;self.case_board["cancelled"]=False
        base_output=output
        solver=self.solver.get()
        self.run_button.configure(state='disabled')
        def worker():
            messages=[]
            try:
                for state in states:
                    if self.case_board.get("cancelled"):break
                    ledger.set(state,"RUNNING","Stage 1/3: model and gravity preload")
                    output=base_output/('Modal_'+state);output.mkdir(parents=True,exist_ok=True)
                    self.root.after(0,lambda state=state,output=output:self.status.set(f'Running {state} stiffness case: stage 1/3: rebuild + gravity preload; stage 2/3: solve {n} modes with {solver}; stage 3/3: mass participation, period comparison and export. Results: {output}'))
                    source=_runner_source(model,output,n,solver,period_inputs,mode_x,mode_y,state)
                    runner=output/'_run_nbc105_modal_v111.py'
                    runner.write_text(source,encoding='utf-8')
                    proc = subprocess.Popen([str(python), '-u', str(runner)], cwd=str(model.parent), stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, encoding="utf-8", errors="replace", env={**dict(__import__('os').environ),'PYTHONIOENCODING':'utf-8','PYTHONUTF8':'1'})
                    self.case_board["process"]=proc
                    log = output / "NBC105_Modal_V111_run.log"
                    with log.open('w',encoding='utf-8') as stream:
                        for line in proc.stdout:
                            stream.write(line);stream.flush()
                            if "[GRAVITY] PASS" in line:ledger.set(state,"RUNNING","Stage 2/3: eigenanalysis; native solver may be quiet")
                            if "[NBC105 MODAL]" in line:ledger.set(state,"RUNNING","Stage 3/3: reports and mass participation")
                            self.root.after(0,lambda text=line:self.status.set(text.rstrip()))
                    proc.wait()
                    if proc.returncode:
                        msg = f"Run failed (exit {proc.returncode}). See {log}"
                    else:
                        result = json.loads((output / "NBC105_Modal_V111.json").read_text(encoding="utf-8"))
                        result['limit_state']=state
                        (output/'NBC105_Modal_V111.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
                        export_warnings = _export_plot_and_excel(result, output)
                        import shutil
                        shutil.copy2(output/'NBC105_Modal_V111.json',base_output/f'NBC105_Modal_{state}_V127.json')
                        # Keep legacy latest-result readers working, with explicit case metadata.
                        shutil.copy2(output/'NBC105_Modal_V111.json',base_output/'NBC105_Modal_V111.json')
                        msg = (f"Finished {state}: {len(result.get('rows',[]))} eigenmodes. UX cumulative={result['cumulative_UX_percent']:.3f}%; UY cumulative={result['cumulative_UY_percent']:.3f}%. Modal mass coverage is not a section design PASS. "
                               f"Review CSV/JSON/TXT outputs and log in {output}.")
                        if export_warnings:
                            msg += "\n" + "\n".join(export_warnings)
                    ledger.set(state,"FAILED" if proc.returncode else "COMPLETED",msg)
                    messages.append(msg)
                summary='\n'.join(messages)
                self.root.after(0,lambda summary=summary:self.status.set(summary))
                self.root.after(0,lambda summary=summary:messagebox.showinfo('Modal audit batch finished',summary,parent=self.root))
            except Exception as exc:
                error_text=str(exc)
                self.root.after(0,lambda error_text=error_text:self.status.set('Error: '+error_text))
            finally:
                ledger.finish_pending("Batch interrupted before completion; inspect log")
                self.root.after(0,lambda:self.run_button.configure(state='normal'))
        threading.Thread(target=worker, daemon=True).start()

    def run(self):
        if self._owns_root:
            self.root.mainloop()


if __name__ == "__main__":
    ModalAuditGUI().run()
