"""Explicit Excel → GUI mapping. Workbook formulas are not evaluated by openpyxl.

Imported data remain attributed project inputs; no import certifies code
compliance or replaces the independent ETABS/OpenSees reconciliation.
"""

__release__ = '134'
from pathlib import Path
import json
import math

CALCULATOR='NBC105_2025_Seismic_Calculator_ETABS_Direct Interaction - Copy.xlsx'
SCALING='ULTIMATE LOAD  - NEPAL - 3 - BELT TRUSS, OUTRIGGER CAP TRUSS_NBC105_RS_Scaling.xlsx'


SERVICE_SCALING='SERVICE LOAD  - NEPAL - 3 - BELT TRUSS, OUTRIGGER CAP TRUSS_NBC105_RS_Scaling_2.xlsx'

def _wall_aw_from_lines(height_m, lines):
    height=float(height_m)
    if not math.isfinite(height) or height<=0:raise ValueError('Seismic height must be finite and positive.')
    aw={'X':0.,'Y':0.};segments=[]
    for (axis,plane,thickness),intervals in lines.items():
        t=float(thickness)
        if axis not in aw:raise ValueError(f'Unsupported wall direction: {axis}')
        if not math.isfinite(t) or t<=0:raise ValueError('Wall thickness must be finite and positive.')
        merged=[]
        for start,end in sorted(intervals):
            start=float(start);end=float(end)
            if not math.isfinite(start) or not math.isfinite(end) or end<=start:
                raise ValueError('Wall segment coordinates must be finite with positive length.')
            if merged and start<=merged[-1][1]+1e-8:
                merged[-1]=(merged[-1][0],max(end,merged[-1][1]))
            else:
                merged.append((start,end))
        for start,end in merged:
            length=end-start
            area=t*length
            contribution=area*(.2+min(length/height,.9)**2)
            aw[axis]+=contribution
            segments.append(dict(axis=axis,plane_m=plane,start_m=start,end_m=end,
                                 thickness_m=t,wall_length_m=length,
                                 wall_area_m2=area,Aw_contribution_m2=contribution))
    return aw,segments

def source_workbooks(root, export_dir,limit_state='ULS'):
    root=Path(root);export_dir=Path(export_dir)
    search=[export_dir/'Project_Input_Workbooks',root/'Project_Input_Workbooks',Path(r'D:/Thesis/Main Python Programs/Excel Sheets')]
    selected=export_dir/'NBC105_Workbook_Source.json'
    if selected.is_file():
        names=json.loads(selected.read_text(encoding='utf-8-sig'))
        return [Path(p) for p in names.get('workbooks',[]) if Path(p).is_file()]
    # Named files only: never take the newest workbook from another scenario.
    return [next((d/name for d in search if (d/name).is_file()),None) for name in (CALCULATOR,SERVICE_SCALING if str(limit_state).upper()=='SLS' else SCALING)]


def import_project_inputs(root, export_dir,limit_state='ULS'):
    from openpyxl import load_workbook
    state=str(limit_state).strip().upper()
    if state not in ('ULS','SLS'):raise ValueError('limit_state must be ULS or SLS')
    values={};audit=[];notes=[]
    def put(key,value,path,sheet,cell):
        if value is None or str(value).startswith('='):return
        values[key]=value;audit.append(dict(field=key,value=value,source=str(path),sheet=sheet,cell=cell))
    for path in source_workbooks(root,export_dir,limit_state):
        if path is None:continue
        wb=load_workbook(path,data_only=True,read_only=True)
        if 'NBC Calculator' in wb:
            ws=wb['NBC Calculator']
            for key,cell in {'soil':'B9','Z':'B10','I':'B11','period_factor':'B15','system_excel':'E10'}.items():put(key,ws[cell].value,path,ws.title,cell)
            system=ws['E10'].value
            if 'NBC Parameters' in wb:
                for row in wb['NBC Parameters'].iter_rows(min_row=16,max_row=36,max_col=4):
                    if row[0].value==system:
                        for key,c in zip(('rmu','omegau','omegas'),row[1:]):put(key,c.value,path,'NBC Parameters',c.coordinate)
            # B16/B17 are ambiguously labelled ETABS/Rayleigh in the workbook.
            # Import neither as Rayleigh without its force/displacement evidence.
            if ws['B16'].value or ws['B17'].value:notes.append('Workbook B16/B17 need a method/source classification; not assigned as Rayleigh periods.')
        if 'SCALING SUMMARY' in wb:
            ws=wb['SCALING SUMMARY']
            for key,cell in {'awx':'B14','awy':'B15','workbook_adopted_tx':'C14','workbook_adopted_ty':'C15','reference_base_shear_X':'E7','reference_base_shear_Y':'E8','source_edb':'B4'}.items():put(key,ws[cell].value,path,ws.title,cell)
        wb.close()
    folder=Path(export_dir)
    topology=folder/'OpenSees_Element_Connectivity.json'
    if topology.is_file():
        topo=json.loads(topology.read_text(encoding='utf-8-sig'));xyz=[n['coordinates_m'] for n in topo.get('node_map',[])]
        if xyz:
            values['breadth_X']=max(p[1] for p in xyz)-min(p[1] for p in xyz)
            values['breadth_Y']=max(p[0] for p in xyz)-min(p[0] for p in xyz)
            notes.append('Plan breadth is computed from the selected source geometry; confirm per-floor accidental eccentricity for setbacks.')
            # Recalculate effective wall areas from current first-story
            # geometry. Each continuous segment keeps its own Lwi; openings
            # and physical gaps split segments instead of using gross bay length.
            try:
                from OpenSees_Complete_Export_V118 import _area_property
                assignments=json.loads((folder/'OpenSees_Section_Assignment_Audit.json').read_text(encoding='utf-8-sig'))
                area_rows={str(r.get('etabs_area_object')):r for r in assignments['area_assignments']}
                coords={n['opensees_node_tag']:n['coordinates_m'] for n in topo['node_map']}
                base=min(p[2] for p in xyz);roof=max(p[2] for p in xyz);H=roof-base
                levels=sorted({g['master_xyz'][2] for g in topo['rigid_diaphragm_groups'] if g['master_xyz'][2]>base+1e-6})
                if not math.isfinite(H) or H<=0:raise ValueError('Source model has no positive base-to-roof height.')
                first=levels[0];lines={}
                for e in topo['elements']:
                    if e['category']!='SHEAR WALL':continue
                    q=[coords[n] for n in e['node_tags']]
                    zmin=min(p[2] for p in q);zmax=max(p[2] for p in q)
                    if zmax<=base+1e-6 or zmin>=first-1e-6:continue
                    if zmin<base-1e-6 or zmax>first+1e-6:
                        raise ValueError('A wall shell crosses the first-storey bounds; review its floor/mesh assignment.')
                    dx=max(p[0] for p in q)-min(p[0] for p in q);dy=max(p[1] for p in q)-min(p[1] for p in q)
                    if dx>1e-6 and dy<1e-6:axis='X';along=0;across=1
                    elif dy>1e-6 and dx<1e-6:axis='Y';along=1;across=0
                    else:raise ValueError('Non-axis-aligned wall: directional projected properties need separate review.')
                    thickness=_area_property(area_rows[str(e['etabs_object'])],{})['thickness']
                    lines.setdefault((axis,round(q[0][across],6),thickness),[]).append((min(p[along] for p in q),max(p[along] for p in q)))
                aw,segments=_wall_aw_from_lines(H,lines)
                if not all(v>0 for v in aw.values()):raise ValueError('Both directions need recognized first-story walls.')
                for axis in ('X','Y'):
                    key='aw'+axis.lower()
                    geometry_key='geometry_aw_'+axis.lower()
                    values[geometry_key]=aw[axis]
                    audit.append(dict(field=geometry_key,value=aw[axis],source=str(topology),
                                      sheet='first-story wall geometry',
                                      cell='Σ Awi·[0.2+min(Lwi/H,0.9)^2], NBC 105:2025 §5.1.2'))
                    if key in values:
                        values['workbook_'+key]=values[key]
                        try:
                            workbook_aw=float(values[key])
                            if not math.isfinite(workbook_aw) or workbook_aw<=0:
                                notes.append(f'Workbook {key.upper()} must be a finite positive area before use; review the source cell.')
                            elif not math.isclose(workbook_aw,aw[axis],rel_tol=1e-6,abs_tol=1e-9):
                                notes.append(f'Workbook {key.upper()} ({workbook_aw:g} m²) differs from the independent exported-geometry estimate ({aw[axis]:g} m²). The reviewed workbook input is preserved; reconcile wall grouping, openings, and Lwi before adopting either value.')
                        except (TypeError,ValueError):
                            notes.append(f'Workbook {key.upper()} is not numeric; review the source cell before period calculations.')
                    else:
                        values[key]=aw[axis]
                        audit.append(dict(field=key,value=aw[axis],source=str(topology),
                                          sheet='first-story wall geometry',
                                          cell='automatic fallback to geometry_aw_'+axis.lower()))
                values['wall_geometry_segments']=segments
                notes.append('The automatic geometry estimate merges shell mesh subdivisions and treats separated wall segments independently; it is an audit/fallback, not an interpretation of how to group perforated walls for NBC Aw. A reviewed workbook Aw is preserved for period calculations. Confirm wall/opening grouping, directional Lwi, and whether seismic height H excludes qualifying basements under NBC 105:2025 §5.1.2.')
            except Exception as exc:
                notes.append('Live wall Aw reconstruction unavailable; imported workbook Aw remains a review input: '+str(exc))
    ray_name='NBC105_Rayleigh_Preliminary_V123'+('_SLS' if state=='SLS' else '')+'.json'
    ray=folder/'OpenSees_Analysis_Results'/ray_name
    data=folder/'OpenSees_Complete_Model_Data_V118.json'
    if ray.is_file() and data.is_file():
        import hashlib
        result=json.loads(ray.read_text(encoding='utf-8-sig'))
        ledger=folder/'NBC105_ESM_Story_Weights_Assembled_V118.csv'
        valid=(result.get('method_version',0)>=125 and result.get('limit_state')==state
               and result.get('model_sha256')==hashlib.sha256(data.read_bytes()).hexdigest()
               and ledger.is_file()
               and result.get('story_weights_sha256')==hashlib.sha256(ledger.read_bytes()).hexdigest())
        if valid:
            for axis in ('X','Y'):
                axis_result=result[axis]
                period=axis_result.get('period_s')
                aw=values.get('aw'+axis.lower())
                if (not isinstance(period,(int,float)) or not math.isfinite(period) or period<=0
                        or not isinstance(axis_result.get('Aw_m2'),(int,float))
                        or not math.isfinite(axis_result['Aw_m2']) or axis_result['Aw_m2']<=0
                        or (aw is not None and not math.isclose(axis_result['Aw_m2'],float(aw),rel_tol=1e-8,abs_tol=1e-9))):
                    valid=False
                    break
            if valid:
                for axis in ('X','Y'):put('rayleigh_t'+axis.lower(),result[axis]['period_s'],ray,axis,'period_s')
        if not valid:notes.append(f'Saved {state} Rayleigh results are stale, incomplete, or inconsistent with the current model, seismic weights, or wall Aw; rerun the preliminary analysis.')
    notes.append('Regularity is not inferred from Excel or model symmetry. Confirm NBC §5.4 for the asymmetric building.')
    return dict(values=values,audit=audit,notes=notes)


def populate(gui, root, folder):
    state=getattr(gui,'ls',getattr(gui,'limit_state',getattr(gui,'state',None)))
    state=state.get() if state is not None else 'ULS'
    data=import_project_inputs(root,folder,state);v=data['values']
    for key in ('soil','Z','I','awx','awy','rmu','omegau','omegas','period_factor','rayleigh_tx','rayleigh_ty'):
        var=getattr(gui,key,None)
        if var is not None and key in v:var.set(str(v[key]))
    # The project's explicitly stated location, not a municipality inferred
    # from a filename or an unverified zoning lookup.
    if hasattr(gui,'location') and not gui.location.get():gui.location.set('Kirtipur')
    if 'Dual' in str(v.get('system_excel','')) and 'Shear wall' in str(v.get('system_excel','')):
        if hasattr(gui,'system'):gui.system.set('RC dual system with shear wall')
    for target,source in [('aw_x','awx'),('aw_y','awy'),('tr_x','rayleigh_tx'),('tr_y','rayleigh_ty')]:
        if hasattr(gui,target) and source in v:getattr(gui,target).set(str(v[source]))
    if hasattr(gui,'bdim'):
        gui.bdim.set(str(v.get('breadth_'+gui.direction.get(),'')))
    gui.project_input_audit=data
    return data


def connect_etabs():
    """Use CSI Helper attachment first; tolerate ROT registration differences."""
    import comtypes.client
    errors=[]
    try:
        helper=comtypes.client.CreateObject('ETABSv1.Helper')
        try:
            from comtypes.gen import ETABSv1
            helper=helper.QueryInterface(ETABSv1.cHelper)
        except ImportError:pass
        etabs=helper.GetObject('CSI.ETABS.API.ETABSObject')
        if etabs is not None:return etabs
    except Exception as exc:errors.append(str(exc))
    try:return comtypes.client.GetActiveObject('CSI.ETABS.API.ETABSObject')
    except Exception as exc:raise RuntimeError('ETABS attachment failed. Use the same Windows account and elevation for ETABS and Python. Cached workbook inputs remain available. '+str(exc)+'; '+ '; '.join(errors)) from exc
