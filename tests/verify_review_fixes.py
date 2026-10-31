"""Regression checks for the reviewed defects; run with OpenSeesPy/openpyxl.

These small fixtures check command behavior, not the project building.
"""
import ast
import csv
import json
from pathlib import Path
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import openseespy.opensees as ops
from openpyxl import Workbook, load_workbook
from OpenSees_Area_Element_Integration_V116 import (
    apply_etabs_area_uniform_loads, calculate_area_seismic_mass_contribution,
)
from OpenSees_Runtime_Utilities_V117 import add_path_excitation, excel_cell


class ReviewFixes(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(dir=ROOT / 'tests', prefix='fixture_')
        self.folder = Path(self.tmp.name)
        ops.wipe(); ops.model('basic', '-ndm', 3, '-ndf', 6)
        nodes = []
        for tag, xyz in enumerate(((0,0,0),(2,0,0),(2,2,0),(0,2,0)), 1):
            ops.node(tag, *xyz); ops.fix(tag, 1,1,1,1,1,1)
            nodes.append(dict(etabs_point_name=str(tag), opensees_node_tag=tag, coordinates_m=xyz))
        self.topology = self.folder / 'topology.json'
        self.topology.write_text(json.dumps(dict(node_map=nodes, elements=[dict(etabs_object='A',dimension='area_polygon',category='FLOOR')])))
        self.loads = self.folder / 'loads.csv'
        self.write_loads()

    def tearDown(self):
        ops.wipe(); self.tmp.cleanup()

    def write_loads(self, legacy_error=False):
        rows = [dict(object='A',pattern='Dead',coordinate_system='Global',direction_code=6,value=-5,area_m2=4,total_force_raw_units=20 if legacy_error else -20,node_names=['1','2','3','4']),
                dict(object='A',pattern='Live',coordinate_system='Global',direction_code=10,value=2,area_m2=4,total_force_raw_units=-8,node_names=['1','2','3','4'])]
        with self.loads.open('w',newline='') as f:
            w=csv.DictWriter(f,fieldnames=list(rows[0])); w.writeheader(); w.writerows(rows)

    def test_separate_patterns_actual_opensees(self):
        apply_etabs_area_uniform_loads(ops,load_assignment_csv=self.loads,topology_json=self.topology,
            pattern_specs={'Dead':dict(pattern_tag=901,time_series_tag=901,factor=1), 'Live':dict(pattern_tag=902,time_series_tag=902,factor=1)},raw_force_to_kN=1)
        # Remove Live and evaluate reactions: only the -20 kN Dead load remains.
        ops.remove('loadPattern',902); ops.setTime(1); ops.reactions()
        self.assertAlmostEqual(sum(ops.nodeReaction(n,3) for n in range(1,5)),20)

    def test_global_z_and_gravity_seismic_mass(self):
        m=calculate_area_seismic_mass_contribution(load_assignment_csv=self.loads,topology_json=self.topology,pattern_factors={'Dead':1,'Live':.3})
        self.assertAlmostEqual(sum(m.values())*9.80665,22.4)

    def test_legacy_sign_error_rejected_before_patterns(self):
        self.write_loads(legacy_error=True)
        with self.assertRaisesRegex(ValueError,'re-export'):
            apply_etabs_area_uniform_loads(ops,load_assignment_csv=self.loads,topology_json=self.topology,
                pattern_specs={'Dead':dict(pattern_tag=901,time_series_tag=901,factor=1)},raw_force_to_kN=1)
        self.assertEqual(list(ops.getPatterns()),[])

    def test_excitation_preserves_existing_series_and_patterns(self):
        record=self.folder/'motion.dat'; record.write_text('0\n0.1\n0\n')
        ops.timeSeries('Linear',1); ops.pattern('Plain',1,1)
        # Force a time-series-only collision on a build without enumeration.
        ops.timeSeries('Linear',910000)
        a=add_path_excitation(ops,direction=1,dt=.01,file_path=record,factor=9.80665)
        b=add_path_excitation(ops,direction=2,dt=.01,file_path=record,factor=9.80665)
        self.assertNotEqual(a[0],910000); self.assertNotEqual(a,b)
        self.assertIn(1,ops.getPatterns())
        self.assertIn(a[1],ops.getPatterns()); self.assertIn(b[1],ops.getPatterns())

    def test_nested_workbook_values(self):
        values=[(1.,2.),[dict(mode=1,zeta=.05)],dict(source='fixture'),12.5]
        wb=Workbook(); wb.active.append([excel_cell(x) for x in values])
        path=self.folder/'result.xlsx'; wb.save(path)
        loaded=load_workbook(path); cells=list(loaded.active.values)[0]
        self.assertEqual(json.loads(cells[0]),[1.,2.]); self.assertEqual(json.loads(cells[1]),values[1]); self.assertEqual(cells[3],12.5)
        loaded.close()

    def test_generated_esm_recovery_keeps_target(self):
        source=ast.parse((ROOT/'components/equivalent_static/calculation_and_gui.py').read_text(encoding='utf-8-sig'))
        block=None
        for node in ast.walk(source):
            if isinstance(node,ast.List):
                lines=[x.value if isinstance(x,ast.Constant) and isinstance(x.value,str) else None for x in node.elts]
                if 'ok = ops.analyze(10)' in lines:
                    i=lines.index('ok = ops.analyze(10)'); j=lines.index('ops.reactions()',i)
                    block='\n'.join(lines[i:j]); break
        self.assertIsNotNone(block)
        class PartialFailure:
            def __init__(self, accepted): self.factor=0.; self.step=.1; self.accepted=accepted
            def analyze(self,n):
                if n==10: self.factor=self.accepted*.1; return -3
                self.factor+=self.step; return 0
            def getLoadFactor(self,tag): return self.factor
            def integrator(self,kind,step): self.step=step
            def __getattr__(self,name): return lambda *a: None
        for accepted in (0,1,5,9):
            fake=PartialFailure(accepted)
            exec(compile(block,'esm_recovery','exec'),dict(ops=fake,ESM_LOAD_PATTERN_TAG=98001))
            self.assertAlmostEqual(fake.factor,1.)

    def test_converter_signed_area_pressure(self):
        source=ast.parse((ROOT/'components/ui/application.py').read_text(encoding='utf-8-sig'))
        expr=next(n.value for n in ast.walk(source) if isinstance(n,ast.Assign)
                  and any(isinstance(t,ast.Name) and t.id=='total' for t in n.targets)
                  and any(isinstance(x,ast.Name) and x.id=='area_dir' for x in ast.walk(n.value)))
        code=compile(ast.Expression(expr),'signed_pressure','eval')
        for direction,pressure,expected in ((6,-5,-20),(6,5,20),(10,5,-20),(10,-5,20)):
            self.assertEqual(eval(code,dict(area_dir=direction,a={'value':pressure},area=4)),expected)

    def test_roof_plot_uses_highest_storey_in_both_modules(self):
        stories=[dict(story_name='L1'),dict(story_name='Roof')]
        for filename in ('OpenSees_Linear_Time_History_NBC105_V115.py','OpenSees_Nonlinear_Time_History_NBC105_V116.py'):
            source=ast.parse((ROOT/('components/nonlinear_time_history/runner_and_gui.py' if 'Nonlinear' in filename else 'components/linear_time_history/runner_and_gui.py')).read_text(encoding='utf-8-sig'))
            expr=next(n.value for n in ast.walk(source) if isinstance(n,ast.Assign)
                      and any(isinstance(t,ast.Name) and t.id=='roof_key' for t in n.targets))
            self.assertEqual(eval(compile(ast.Expression(expr),'roof_selection','eval'),dict(stories=stories)),'Roof_UX_m')


if __name__=='__main__': unittest.main(verbosity=2)
