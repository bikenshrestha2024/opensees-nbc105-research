"""Workbook mapping must not manufacture modal/Rayleigh results."""
import json
from pathlib import Path
import sys
import tempfile
import unittest
from openpyxl import Workbook

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from OpenSees_Project_Inputs_V123 import CALCULATOR, import_project_inputs


class InputMappingTests(unittest.TestCase):
    def test_formula_blanks_and_ambiguous_periods_are_not_modal_results(self):
        with tempfile.TemporaryDirectory(dir=ROOT/'tests',prefix='input_fixture_') as tmp:
            root=Path(tmp);books=root/'Project_Input_Workbooks';books.mkdir()
            wb=Workbook();ws=wb.active;ws.title='NBC Calculator'
            for cell,value in {'B9':'C','B10':.35,'B11':1,'B12':'=1+2.5','B15':1.25,'B16':2.0,'B17':1.9,'E10':'RC Shear wall — Dual System'}.items():ws[cell]=value
            factors=wb.create_sheet('NBC Parameters');factors.append(['unused'])
            for cell,value in {'A34':'RC Shear wall — Dual System','B34':3.5,'C34':1.4,'D34':1.2}.items():factors[cell]=value
            wb.save(books/CALCULATOR)
            # Explicit source selection prevents external project fallbacks.
            (root/'NBC105_Workbook_Source.json').write_text(json.dumps({'workbooks':[str(books/CALCULATOR)]}),encoding='utf-8')
            result=import_project_inputs(root,root)
            self.assertEqual(result['values']['rmu'],3.5)
            self.assertEqual(result['values']['soil'],'C')
            self.assertNotIn('rayleigh_tx',result['values'])
            self.assertNotIn('etabs_tx',result['values'])
            self.assertTrue(any('B16/B17' in note for note in result['notes']))
            row=next(row for row in result['audit'] if row['field']=='rmu')
            self.assertEqual((row['sheet'],row['cell']),('NBC Parameters','B34'))

    def test_new_story_ledger_includes_interfloor_weights(self):
        # Guard against returning to the old diaphragm-slave-only accounting.
        text=(ROOT/'components/model/assembly.py').read_text(encoding='utf-8')
        self.assertIn('story_ledger_mass_error_percent',text)
        self.assertIn('for node,vec in values.items()',text)
        self.assertIn('fixed-base masses excluded',text)


if __name__=='__main__':unittest.main(verbosity=2)
