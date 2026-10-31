from pathlib import Path
import json,sys,ast,math
p=Path('outputs/OpenSees_NBC105_V141').resolve();sys.path.insert(0,str(p));folder=Path('work/V141_Project_Rebuild').resolve();d=json.loads((folder/'OpenSees_Complete_Model_Data_V118.json').read_text());cfg=json.loads((folder/'OpenSees_Complete_Model_Config_V118.json').read_text())
# Check new generation path receives current payload, not a stale on-disk model.
from OpenSees_Storey_Gravity_V131 import export_storey_gravity
export_storey_gravity(folder,d,cfg)
from openpyxl import load_workbook
path=folder/'Gravity_Storey_Load_Breakdown_V131.xlsx';w=load_workbook(path,read_only=True)
counts={ws.title:sum(cell.data_type=='f' for row in ws.iter_rows() for cell in row) for ws in w};print(counts);assert counts['Surface load formulas']>0 and counts['Loaded node path']>0 and counts['Source detail']>0 and counts['Storey totals']>0
# Independent arithmetic: integrated source cell pressures must match assembled area-load provenance.
a=json.loads((folder/'OpenSees_Load_Conversion_Audit.json').read_text());by={}
for cell in d['shells']:
 if cell['category']=='FLOOR':by[str(cell['object'])]=by.get(str(cell['object']),0)+cell['area']
expected=sum(float(item['value'])*by.get(str(item['object']),0)*cfg['gravity_factors'][item['pattern']]*(1 if int(item['direction_code'])==10 else -1) for item in a['applied_area_loads'])
actual=sum(item['downward_kN']*cfg['gravity_factors'][item['pattern']] for item in d['load_provenance'] if item['component']=='area_load' and str(item['source']).removeprefix('AREA:') in by)
assert math.isclose(expected,actual,abs_tol=1e-5);print('Applied floor-pressure reconciliation',expected,actual)
from OpenSees_Mapping_Audit_V125 import parameter_sources,source_parameter_inventory,validation_inventory
assert len(parameter_sources())>100;assert source_parameter_inventory(folder);assert validation_inventory();print('material source rows',len(parameter_sources()))
for f in p.rglob('*.py'):ast.parse(f.read_text(encoding='utf-8-sig'))
print('All syntax pass')
# Standalone JS syntax checker input.
for n in ast.walk(ast.parse((p/'components/model/id_audit.py').read_text(encoding='utf8'))):
 if isinstance(n,ast.Constant) and isinstance(n.value,str) and n.value.startswith('<script id="shared-id-selector-v121">') and len(n.value)>1000:Path('work/v141_selector.js').write_text(n.value.split('>',1)[1].rsplit('</script>',1)[0].replace('__DATA__','{}'))
Path('work/V141_Formula_Validation.json').write_text(json.dumps(dict(formula_counts=counts,floor_load_formula_kN=expected,floor_load_provenance_kN=actual,note='Formula structures and independent arithmetic checked; Excel recalc is requested on opening. No cached Excel recalculation claimed.'),indent=2))
